# -*- coding: utf-8 -*-
"""YSF HICH BOM kurucu v2.12 (bom_kur_py.txt, talimat v26.8) — sabit betik; oturumda yamalanmaz, yalnız kullanıcı talebiyle sürüm değişir.
  python3 bom_kur.py proje.json cikti.xlsx
  python3 bom_kur.py --yenile duzenli.xlsx proje.json
  python3 bom_kur.py --donustur eski.xlsx proje.json
Sayfalar: BOM · FIYAT · KURALLAR · TEKLIF · OZET · LISTE
v2.12 değişiklikleri (kullanıcı talebi, 261014):
  - KORPUS_YEDEK varsayılanı 0,20 → 0,10 (aksi belirtilmedikçe modüler kapasitede %10 yedek). Proje farklı oran isterse
    json tablolar.param.KORPUS_YEDEK ile ezilir. Başka değişiklik yok.
v2.11 değişiklikleri (kullanıcı onayı, 260930 — bara, tarak, normlar):
  - Bara taşıma kapasitesi kullanıcı tablosundan (ELEKTRIK_PANOSU_BAKIR_VE_ALUMINYUM_BARALARI.xlsx › LISTE answer):
    tek şerit · 2 paralel · 3 paralel (= tek × BARA_PAR3 2,25). Yalnız fiyat listesindeki EKT baraları (м). BARA_DERATING = 1
    (tablo 25 °C, doğal havalandırma; ek düşürme yok — KURALLAR'da değiştirilebilir). Seçim: önce tek şerit, yetmezse 2, sonra 3
    paralel; aynı şerit sayısında en küçük uygun (liste fiyatı kapasiteyle artar). Tablo aşılırsa BARA TABLOSU AŞILDI + TEKNIK: KONTROL.
    N = faz (aynı şerit sayısı), PE = faz toplam kesitinin yarısı (tek şerit yetmezse 2 şerit).
  - Bara kararı (KURALLAR › PANO Y "BARA TİPİ" FAZ/SNI): IN_ANA ≥ BARA_ESIK_A · montaj DİKİLİ · pano adında ВРУ/ГРЩ/АВР ·
    CIKIS rolünde mccb/ppn/rubilnik → FAZ. json pano.bara_tip ile ezilebilir. ППН IN_ANA'ya girer.
  - Bara boyu (faz başına) = iç genişlik + korpus yüksekliği × BARA_DIKEY_ORAN (0,70); korpus adedi (bölme) ile çarpılır.
    Baralı panoda korpus G/Y yoksa HATA.
  - ШНИ baralı panoda da modüler çıkışlar için kalır. N delik: 1P/2P/4P MCB ve RCBO/УЗО (3P MCB N almaz); PE delik: güç çıkışları
    (mccb/ppn/rubilnik hariç — PE barasına bağlanır).
  - Tarak (CHINT PIN): aynı sırada (BOM T sayı) ≥ 2 adet oto/dif/uzo varsa sıra başına tarak; 3 fazlı pano 3P (IN_ANA ≤ 63 A → 63 A,
    üstü 100 A), tek fazlı 1P 100 A; boy = modül × 17,5 mm → TARAK_BOY_MM parça. Taraklı MCB faz teli almaz; RCBO/УЗО yalnız N teli.
  - İzolatör yalnız faz + N (4) × (ROUNDUP(boy / aralık) + 2).
  - Normlar: korpus 8 / 30 / 90 dk, lamba ve buton 8 dk (kullanıcı 260927). Legrand katsayı 1,00. Tarak normu boş → KONTROL.
  - Liste fiyatı boş, markası bilinen (pys değil) FIYAT satırı birlesik.csv'den doldurulur (eski dosyadan --yenile).
  - --yenile: eski bara tablosu (2 paralel sütunu yok) ve v2.10 varsayılan normları v2.11 değerleriyle değiştirilir (KAYNAK: KONTROL).
v2.10 değişiklikleri (kullanıcı onayı, 260930 — sayım ve veri bütünlüğü):
  - Veri yolları: birlesik.csv / kaynaklar.csv / ebat.csv / tekled_logo.png sırasıyla $BOM_VERI → /mnt/project → /mnt/project/claude
    → betik klasörü. birlesik.csv yoksa HATA. İki farklı ebat.csv varsa ilki kullanılır, KORPUS: KONTROL.
  - mutabakat[] zorunlu (talimat §5.1/§7): pano · grup · artikul[] · sema · cetvel · bom · count_a · count_b · lokator.
    HATA (Excel kurulmaz): mutabakatsız pano, sema ≠ bom, sema ≠ max(count_a, count_b), boş lokatör, satır toplamı ≠ bom,
    mutabakat grubunda olmayan cihaz satırı. KONTROL: count_a ≠ count_b, cetvel ≠ şema. Satırda "projede yok" gerekçesi olan
    (kaynak EK) kalemler muaf. kaynak ELLE (--yenile ile kullanıcı değişikliği) → farklar HATA değil SAYIM: KONTROL.
  - KONTROL satırları kodlu; tanımsız tip HATA; yasak desenler HATA; SIRA_MM referansı; faz lambası kontrolü; EBAT çakışması; lok alanı.
v2.9: FIYAT'ta kullanılmayan sarf artikülleri altta gri, gruplu (YEDEK SARF); görseller betiğe gömülü.
v2.8: EBAT tablosu (ebat.csv) ile cihaz/korpus ölçüleri.
v2.7: TEKLIF TEKLED КП şablonu. v2.6/2.5/2.4/2.3/2.2: TEKLIF ve FIYAT/KURALLAR altyapısı.
"""
import sys, json, re, csv, os, math, datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.formula import ArrayFormula

_BD = os.path.dirname(os.path.abspath(__file__))
KOKLER = [k for k in (os.environ.get('BOM_VERI'), '/mnt/project', '/mnt/project/claude', _BD) if k]
def _yollar(ad): return [os.path.join(k, ad) for k in KOKLER]
def _bul(ad): return next((y for y in _yollar(ad) if os.path.exists(y)), None)
BIRLESIK = _bul('birlesik.csv')
KAYNAKLAR = _bul('kaynaklar.csv') or ''
EBAT = _yollar('ebat.csv')
LOGO = _bul('tekled_logo.png') or ''
GORSEL = _yollar('teklif_gorseller_b64.txt')
KODLAR = ('SAYIM', 'OKUMA', 'ARTIKUL', 'TEKNIK', 'KORPUS', 'FIYAT', 'KUR', 'KAYNAK', 'TEKLIF')
kodlu = lambda k: any(str(k or '').startswith(c + ':') for c in KODLAR)
GIRIS_PANO = re.compile(r'ВРУ|ВРЩ|ГРЩ|ШВУ|АВР|(?<![А-ЯA-Z])ВУ(?![А-ЯA-Z])', re.I)
BARA_PANO = re.compile(r'ВРУ|ГРЩ|АВР', re.I)
TARAK_TIP = ('oto', 'dif', 'uzo')
# kategori, tip (boş = genel), IN alt sınır A, sıra yüksekliği mm — talimat §6.2 başlangıç sıra yükseklikleri
SIRA_MM = [('MODULER', None, 0, 150), ('MCCB', None, 0, 250), ('MCCB', None, 160.01, 300), ('MCCB', None, 320.01, 400),
           ('MCCB', 'ats', 0, 150), ('MCCB', 'ats', 63.01, 250), ('KONTAKTOR', None, 0, 200), ('KONTAKTOR', None, 95.01, 300),
           ('OLCU', None, 0, 300), ('KLEMENS', None, 0, 120), ('BARA', None, 0, 150)]
# TEKLIF görselleri betiğe gömülü (base64 PNG: tekled logo, teksan mühür, KZ rozet). Yukarıdaki dosya varsa onu ezer.
GORSEL_GOMULU = {
    'tekled': 'iVBORw0KGgoAAAANSUhEUgAAAfQAAABxCAMAAAAUAqFnAAAAwFBMVEX+/v7RIjejp6rMGjGeo6bu6Om0trjZ2du7GDDZd4TbpK7OVGXIRFbLzM3ILEHRZnTtxMrbhZHz09bgiJPuuL67jJHDOE27Ok353eW8JTrkqbDsusK4l5q7R1e5c328vsDheoicnqCznKO+wMHbmqa6K0G6WGe3Z3O4d4K4MT/hm6jTcXzITWKZmpyyoZ3EUV7SkZjd3uDf4OG+TWG8U1/DCibduL7glJydsK+3DCzEMj/PbIDgITvswL/54t4AAADl11TkAAAAQHRSTlP///////////////////////////////////////////////////////////////////////////////////8AwnuxRAAACP5JREFUeNrtnX1/mjoUxwNJQK+i6FTEh3Zdte3artvutrv7/P5f1gBtVXKAkICSkN9nf6wqEfiacHJyzglCRkZGRkZGRkZGRkZGRkZGRkZGRmfVH79VquVRy5sK2qNvzflscz7H9U2K2t2JuoxE7ibQTHnRku1RSkue5bdOJNypSneHltcVNBzM39ob/pVqDHe6xdf3GLKNYty5St8m9xNJ6YYKQHc/2URWN0c/N+cTxwF2JMfzvnKfMO1hq0LhYaUt49nhQrpMc/h94eUtBxg8S+b+uDYjIeh2BTqGzn8Uicm7VAPo68NFPAlAH/dB5usxUgU6KXNg3OkdDu5a93T4FPDMR7VBJ5eEvtftV8Whz2Wg/wOeQbhAekOPu7u+0AsMuXenR+z/6oyQ9tALsDcceiABfXFE+kC9M0R1QrcbAj0Xu77Q/QCfMt8d06MtgW7b17Rt0DMM95mPWgOd2G67oL9+e+qxPliieqGTBkGP5LUK+uun8ckjPdwg1J6eHnd2h7YH+k/wu3MO0BT6aWuqQKdC0BdbnH9yLYIOUFfaI5cF3Z9hdn5u4f64fuikedBZ6js0uInQsahzJm247/+a+aid0BnqYmgwPg90S+SZDn1x/MoCoVYO7yz16A5BBPcvgk/G+E3c4OG9C34v7H1VD7rYr+D3FPRwNhhE/wBFrwJ3L3p5Bn8+OiK8P4GO5fq5kCH3Mzydq+1/vkN0Fug38sxJfk9/eVmtPryJt9HTi5n7Y3+cJb8HLG3lfD5668icWvqyWvqoLPQJOEBleV+rho5oBUJ50FdBMAs+7hQkWn/uXU1Xdn68hVPiGobsnf5yoWA+Puh7wx2nHxNjdB7oFYuFPkqermlZwfp+mjf6E08K+qjJ0OfPoBkymCBNoP85DcErjBQGw1XOM9/VFvoQDpt4RLpAJzD0Pfnt3dSRH+AB6O8aDH1klfS+Nh36NQs9zzqOuvvdyskw9D09oT9iq6T39SLQXY9bTL/Nh5709qusyR3VEXo63hkXe18vAt3hndcBz+ci6DH2qLNDT/YbT0Po41mpsInLQb+VmcTnD+87P1TwBfbkUe2gZ8Q753tfLwL9mtQF/XXJYjuyCRGetikEvdsBbwHnZEOPnv7a2/F2JWHAqwN9BBtxT6h50K9rhb4PRFiJz9WVgb4IwWvv0XZCj6iD8bFaQV/O4EW6MWoidFI/dAvfQ8acTtDheGer2Puq4TP9zQ05BaZ8rj7Qswz3DWokdOccPd3CawC6pw/0Lhb1vl4G+ll6OrZYPzxxtIE+AuO5gMoDDYFeYmndEYceWbFAV6eaQF9YbGxnxLxfCpo6q2z80AcrsUmbAtDjsAngPvB5X7WFnvSDh9xQLEWhv49envdlvK96Q78T88Q2vqdnxNyWDvXQcHiPRjtmqY7LkmswdJz09K6c91Vv6OFKyHyXge4vY00ElRxZEA37Hm0sOe9r5dDrjoYtAz2atAldlQR02hvIaPdTXefGvb/7Dy4Tx+99rRp67ckOpXo6YMnVDf05SYRJcg1O/mMB/8HY2n8av306jvDMS3bAVi+j2sQEGejJnbg6f0+vu7pURkpdKU+c3tDv1YSeN7xbPTBAyrpkTydNgY7VhV7wTF9k/FR8A92CJ+rKD++R9Z4RL/OZGug76CQdU6v68B4/vJ/Amgpcoe7mma4q9IZ55BrW0/WDjhMzPaNIYLgwPf0SU7b+DyyrH0FxFejJALMVYHkSVWuBHu/scCO1r8Ppzg5y0M/ukUOb7khaXzjqvW/CCtxyFUGnXgWiVUG3WN87T2ikGpEzWN4Br+WCy9/nX3A5G3Q07AhnqzYdukS4lIX77Hr6rTbQ6WdpE1476Bg23j1toKdKRB7Wkxdt7ungyqqrD3S0HJyWBC3tj9UR+lowW1mZXLZNaEFVpfrzFkO/EsxbVSdrtQvGxHKb8OeE7jnXDpeuHeGs1YzE1Wu9oNNhB0ssrp+5/Aifo0Yuw+VKNFdZpUoUfQwW49g0DvpZEhjhXGWqGfSdCc/eke2iadDPk6o8AioV8ZWiUKq61GQALrNymfC6QY/m6ERslq5cHTm4HMXzvEXD+3696TtYL5RqCB1MfeDLX9XrmQ7mpmtYaGhnzPVE42M1Gd5xtluGKfWvC/SMdEbLemxUT68RevxmZ/0C1pvkLQmsGnTAC4+5qs/oU34E455tg3N8V1foKCsqetwkjxy3BAy5QdcBy8oS7trf6kHf7ZqOy/pjm+l7p+ygMA1ze/lbNWAifj0KQs/KXh4queDCQt/mFPm31tPMvR349/NQEXrG/oL5IRXqQIeH96iPdz5+n8bb9xDZnXtUhL6LisblKpIoAp3Yo/Dt4vCedpLqO/v2EK+qJRYcEZ+uyUKn1Sbol9lrNYDjY5eqQyfx8J56qG/jTbqG05f8PbpKbNYkFffefe7Jqj8U3FU5I1BwXAY6v419bHlVoNzN+D48RLp61cPDdLqKZuWvwAk8tpMy27LJFSXoSCc7iG6lXT4qGkpN4c1POBoTkmQHGcWH5+/ACM7pSO7UvxzzS6c1WX1B6KUrh0rlI9FqmuFJa8rY2oX7BBWAXpCf3i317YlNv1EbOgdvIm2dXBq6yFbaRyY845fGg0UN0I+aaeL+6WUtUnWho2UATddx4NcK3W4e9NKzEIWhZ5rw8zZBJ075mafK0FkTPmcndV2HdwehdkHP2G63020PdA+1DnpWIM1jO4Z34riofdCzduIEQip0hH4r6EiW2D+9CdDR5HjXB3xodqw9dNFuDj8VVerpCD1Cq6yAP1Y76MLI5Yd3jCWhpzxymKkuVaQRX5UKnaAT2/EQuhz0CkuKPYlBp0MwFr4z0hV6hNxFSFno+KRM6JPQ8I7Q/Bk+j4WO0ONOLh0HcOln+klt2H/hkmJFyqhNcpripj70eF22AuIJdEbc1ntffjkdz/4/6umM+KCjyRZu+5i6e1PNKtvNxaA7nltRrE931k9pzbuBKR0yx5bX3eE6RkH6zYD39/cInsmse9LTJUJdDs18teUDZ46hexzt3XoR7ibEdhkZGRkZGRkZGRkZGRkZGRkZFesXBuExwh39AyYAAAAASUVORK5CYII=',
    'teksan': 'iVBORw0KGgoAAAANSUhEUgAAAIwAAABgCAMAAAD4mGr1AAAAwFBMVEX+/v6Xl5jq6Oeop6e2U1G3trbX19bIx8f///+zTEqLiYcAAABtamh4dnTr2djEdXPZp6bmxsbMiIe9ZWPRlZTivLzLubiAfnyenqCgnp3Aa2rBwL6CgH6hoJ67YF7j4N++vsBiXlxkYF7WoJ7jwL/KgX/Cv8Df3+D///8AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACd6FhOAAAAQHRSTlP+//7//////1X//wD/////////////////////////////////////IgAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAQWlQgwAAClhJREFUeNrNWwmboroSJRVS0TyCqIjaavcs9y3//xe+qoQlLEJwnHsn30wviuRQdWpNdZL8OWvzlrsAr38cDHyh1koKt6RUGuEfAgOID2NpZe2yZ/NQLwN6GQxoac4E5W6EVErTUiQhw8Cska/heREMSsObXkdSAK3EmQAJBX8TGJT2e2akdvvlLBK3GsaAvhIesx7OC2Dwmn23wiFBJQwtIR5SXoX7WWLOeJTJvp/XwlkNBqT9YSV6hRBfbtgZNSBZljRCubcFSUf/VjCauHLlvRSZcjXFUwfyQYIDvpZh/yYwJBYjtIcyY8FkacJdJu0q4WzWYcks8YCh6AU6EJtYKCjcJ36LZKSpSFMiym5RCglOOBLeD0YrRGQ7ieQBUZh1pc6ZwHeD0edaRfEkIOD0EeJxJJpoMJUTN7Jw1hgfqwpj0cSC0VdDt9UC13pIQY6A0Xy9Dww6566Fio2iqlEmCOHRwLvAgMkIBjEgCgiguHew4UFo9D2T7wJzzQzdcBkLcDi/azQi8IhwZTQ2u70HDBjLol7A8lVJozQHAAglw76S6KayM74DDFJykHwJOa8cZN8PaCWFSWl7TCc0SSJ+iHeAERmxUcwwEKQmblNgpD1NrVFUGIZXleA5U5O5q25DSwQYndmcyIvPbFezGiULThiKGGfKKdDYc082/HGV2fE9tPQxNRKMs6Rc4cQzVfSi+mGAZHc+k+TUnZlKWQ7lNf2YDpLelhNeSnNZoaLBKGJCItXQbOj+lNfR81pSgboby9qxFfsVpaeiOIWSCUXjhWsciARDgrmRzrEHBOlheFciCf03xBYlLL/ESQws1HsDuq0BQ3KkTWRwL5LSjeWgjKCMgiRjFZk9ytsiEBc7ZV/jivWk48DkfGVlXGRByv6VBKKposQbNJVvmrRIz4akmbj6liUhw8wsJM0mhjFeMEjMpMoNlIUbVQLsZ4VBZQD0iiqAJXEJhIOBnhbAgCBTqhlD/lWhPiNa+mrslW6jhV5MKf7btyvty/IWP8SDwTP5hqusc05NiqLwbZWy5B60riJ0g/KiB9moW61wHGkwBozKiL51gk8OlpRD/6yqBDm2ON2QJCQOzScUTkCaTYSW6kCA5qHOirJPgxCfY/FePTRKNMsLx5MmAozTUmvXVDTdFdvUmlTP7a2G5uPXRUOPNPNgbpSgQVuG5XKhXKIoQKt3Ta2Vb0PSXGpdYUiaeTCS8jNtIg0XVbODHoEJSOxeuehvtXQ06BbtZoEymsJbJDlkJ/9OLRiIoO9zsW2+iYY0m/k0nCkzn+BBXu0dlkuHRXQ1pB7ja8wHvl0Gn5gFoy2517miDU7l4Xjcd0yQ3EsLxcA7Xy5hnhD4XAyEictgBOVJM8ZzOqbp9gDN47pWHuiQM6wT5TWoez4X23ZFLbklMIr5O5duHrbpNj0lvkfSCiREL52GdI82YaDGxu3IV8BAvi+L+qUTQdnuYJiWDC1be6zt+73srhEOvzkLRlr2cf38rtwd02PeCiZNq85ecSJD9kLw7FDjQN0JB5c4E3hOkkhR7hPYbbdeM8yYRjC16CncwDgYOIzeqvQ4u2uFo2MrSgZCEkmLJClJGtvC3YN/TIueCXN/GoMQqptta0V52oakaYWjoqqDGghJhA2nShsD2n+QYD7yJp0MfIbUIVcl5EWZe1fs5dEjTdMpwMWova/K3afDkTII2jr/JAhMGij5pbJLbkXo9eoA5Y2pYAFi5218SbCuosxZIg0Q98OHI403ZxZMQ+WBz2g9rhOHYqsrQ9o8s73NE8UUZL77Ywgk/TgUe6hJ85O+9wTjq7pvsgOkO2PKd8zzgDYj0kyDIaoePlJvMSyEGsjRA3H9NPqVSJMzY1wkGB77BI5D+03JB/CVHW00MV3PS+aLgOzc7qm3mLKVSN59cu88XV7Qt/TwvCftLFrXDqTw7ggbx/8kZ910qjkQVdNGL85iquOxD8RdyaQ5lju+bJfPpHe6MaYk+evoFapYIjjfEoE+EP6+dVoeAnGrZAYfU3fx7vQ08WUwDVE9aZ5KpAfG6T9t13FXFody/6xSrhjrrjw67Meiu31bQjXBAS61MUN5LGMbjAFT+Wn3Xy4MjdOoU3Eg5XnS7E87/wCHvM1T5LfAw3IDuPUs+R7CCDfbufJMPe4cEQ6TQJhTrJvO0+QHRrMl8J3XI0o0h7na10zD3j5KtQDmRECYqhWFvhErPRDnh+ky2tllMQVT7egk+lFAvwRpnB4OjpaBXpQzPVgPJi9Ojqo5P/TnfhwiGcjnrjatYluHJ/Cq2qZlPowHk+cuVGiAsTqywQiH1D90m1cOgQSexmENVNXFg4ucrq+oaOe0OrbbWTjShL+z3g614EJP8+mNGtwVW68q4CNKyiJG58tNEwxntTQEc/rs8iUnhH8PgLh1CBIZEp73kmXubX+YX1FhrpoEjJtyOhqMI026f9qC6+S3bY3Oq+p4mHBMuZL6P02ABqMSnc3Vp5txvp+GpJmwcfI0ny1pagdenvKJchcr7v1JKxv6JmL2PGMz+dDPEhzyelyzOfkd93O1KEdJbkJSbQF1UkfpFFqLK8D8lfZJ0wH5i3MLjl8Fk4bWz6cVjk5uDw8i5xJQyUYw2sweno7ymSFpGq/30QRSllvBAWwqRiLcNBiDKGoQPmFwTQQqe8z8Sd4mWSANNElO2oT0D0jyn9Uexi1ehbJSj68Hd4UVgDZ5EqTmkFy/y3WnKoXLZQJj/whzCw4IE/HXpQY38rtUGAh9NeYmzUNjWHaR7nRmcR2YsDJznuZz2yY5nFvs81FWwYHAdaCkUALpP0+vDBtYfPaydDQ4AjMkjfv9KRBM8Krym6kTW8PRslI49E7uUEVlSz2wTbJAGv7dA4Gpw6IvdedQLB1FzVlqORUiSWA8iLB0PL+Zdq9BCXI6TAKhVEpLUwFy+uTPuwmbnkydNHfNHmrxvHUM5jT0NNAe7ITnIdIYpdzB+9Ug1EX0s1awSuCRLZ9Dj8FwQjxO+3kEAsKjND5Y4RkNbgRRgglzbWnFOrrjC2Cg3JVjLyKzs/Yn1EHnhayIDXh+F83DAmhsxDzHZvGsrH7xkd1ZNibszGm9PCDojuYxbnAhekoERI1GtZ2kmLY9O7tE2XvUEEX8/AyhYVErI+KPutANgsksZhBiFZhm4gpF7EwXOGujEHWOHC5ZM3MFyrqpHGVipov47JBrJwnRB0LrBsCqsxt1Q2nkAhzwY2Jk0jb+RGjl0CCPurmRNymuM8c9rmHoLlsxi7Z+gpFH3e6uWpRTrVYXPCtG+uWg2FXjpusHTd3MpoNzoyLpIbnV6jDRV3SFtu9Kobx7Kf5OMGwk/MRYV23kia9+BNd3yeo3GIrQ68ZeX5oHZjjZWbg5U27iKeUryZuv86lqE3Y9lJcnpcFtZ8Rogh24v2izdnT57wDjHp8HxgkQaeam2yHysxOawlfu+WvT9epq/Gi9m7HnMXtrldb44g1/9Y8g/HA0T2y7U4wbfv3Czd7zFxnOrn/9Ju8C85b1h4HZ/Dnrf8m//qD1f5kXeBtEmjtJAAAAAElFTkSuQmCC',
    'kz': 'iVBORw0KGgoAAAANSUhEUgAAA4QAAABNCAMAAADJn1roAAAAwFBMVEVf2OJjxddsbGyg2uUeHh7b8fAqKio2rblVw9cgICA5sc+ZmZlRqtv+8gsDf3/X7/MAAP9/f/9NwdTY6Onu70hNrrj//wCg3Obt7C0yzdH+81/++KAHef8webVVVar/AAB//3+V2ujs8Zg/v386tMoA/wAA/39Pz7+fv9/p7lsAAAAAAABKus4AAAAAAAAAAAAAAAAAAAAAAAD8/f0A//9PwtdJuc1Juc1Jus9Jus9Lu9D+/v5JuM5It8z+7AFTqawqpCG/AAAAQHRSTlMW7hD+YPYkC5uGERAN+wKRAQJlII4MAQ2kCP7+AgUDAQKXmQT/AQIQCF8A8f3RLVCwbo/+AftxT6+P0QMWLv4FKInawgAAJpRJREFUeNrtfQl7G7mVLUlRpCwpiiyP4/SkO7O8mbeghNVkLVxa//9fPdwNQC2UZMf+vkxS6DZF1oJCFXF4L+5y7mKxXB4m2/KwfFJqp24Pz9Ia9VFBe2zytuf94aTOzxfbvlNHNdVcpbWuLLz12vFfaFZZ+FO5eAz81XyC1wZOqqr4FxseHY8yeLhXoeLz+eB4fsArObmmhe48vpFRON5vKktb4Y3HnjxfBv4G7E8HGQrvxA0GL+njycbj5S0NxMNGiwfAyRYuc/Qy+rnNjdtisd8nvOzLP/HNjXpcqcVhn0H4AhDcbdt9gbLmXjWvgLBWq6kLmwpbnKPG8RuFbyKaLP6Nc1bjG5q1Ht5onN48jXGnxj2xhQhC6RJAk+BndCXz3qdTBIQ6vqkSCCschJGB0MBjz3EvoM3yWAINE4cgneCLo4HQ+KHbgKMx+Orh5E+h8vOsm9sQhJfxcwOScDGUhKuHZt+TdC9vgPBlGoTBxIZTOliLQHHaGIsgtLRLwzE2gTaC0BVdVN4g2Hx8g4fHk50hwBUg9FUBwtixi1eYAKEmECLcgrbGmuMqSTunNwQ+FGkwYMAqdG0VQY3GHS8PHyKMYw+I6NAHoaqqMM+6uV0E4fLqGtrVUjagJDw9FyBcACw/H0oQnnfb7wShYImEjIepDtthMovCmSbsce1gqy5nMExtAJsl4dfrtAdCbTMIoVcX31yShAiuoIvLxGsaFq+kCYcKB4y/Gs7zZUmApvshFdhoB2cUIDQoIOc2tx4Il4yo5fVXaddLkYS7gSSMIPyo6n0PY1FgfAMIZYFoBE+GECeLM5aEQxBG4WIHktAICP3rIAwuuB4I474pSQhXwN+AMQgdfjaxI9wrkjP2H/CC2hm4FI0OX+mMeDNVXxJ6guXc5jYhCa++lu0qq6Of+2vC47+XVph9+9tCvQ7CvmHGrq3t46UAoQ4iCccgVAJCkz+/D4TaJ0xlEF6ShASeUJme1kvruIjTbOOB4+N22uArXOrlvRmEA0noQpbbc5ubSMKBGCxQGEG4UotSHX3Zqf/VlRA7q1/fBGFpmEnGzSlJOAahH4HQJeNoAULzCghNtZ4CoZkEoSW4haonCQlRAMIsIqGHOCLaYMiGOwJhfKV7MnzGMa5yZxDObUoSjjAYVVJRR089dXSnzj2rzF+Pj+pb1FEt1s64ONLomSD7/QCEzntvASLxry9ACH4N8VkUIAzxKNMHIW7TqN+aBGpeljnyQYD1E0EIVwmojmrNEC6uK44UOJtFpPEo+aLktGIv9VMgBE0WfwdgMLBqXOmw0vOicG4TIBxjEGVhBOHAMHP/66JwT+wPnyNKIwjP3wRCyyCMUi3Oy0lJSKZ+dlGUkjBCxbOR32bDjHgmMgjJGYGWzniauAUiXK3XKMPg4oFAqMW1YZODo3CNwG6CTRwiw9mTPPeyhmRXxAiEcSeBsOIO4/krp+dpN7eROlqsB69LFOKasC8J1VOJsCbuVyQJ9+BezC/vkIQB3eylJHQJhMHa9RogYq21BQjXhcvcix0zojIevhlIwrjNV2wt0b6EZmDPukZU41XWa48uCq15TVhcN4pegh5cl3paxa41rVAJhK4wiWYQWowggN8ZGCBIQhhzmEE4t5EkFAxeXy2X++XySoC4HIPw4VQKwvZ2BSCMKmrXQuv4pesDdXERhCRLpteEx+OEYQZsPCveGjII+fDJNSG6+ZyA0JHe6jUfjSB0sV9eE/KqsbcmdIbNnkFGdCSAAQBxV0RvNQFCAqAlw8wRAOi1BPPMbW4ZhM/L0hZTWkqvCYQlnh5L98ThjIJw3MplY7ftHZMkzAQISTi+ah1VBQhpnr/lJzSAO52ibtLiUNHicOQnRFD1rKPak7aJK9UURAd9a9kQtJ2yjkZx7L2AEKWgjlvCbJmZ21ASMuSWhfRiQ80fwQdY+gl3hVhEwci9/LfaQXuEl4eXMqotAnUxeeEMwqSOmrddFBg+IyB06h0uCl/1Qt/4SI4SdSM/IR0SSmEFl0MVVmLeKNA1OFrqwaE6HFMQqkLXIYLQsfYrIORIuzlwbW59EN4uh3IQUSggHEjCp9Iq87KbEoQ9W2lcEU5jMINQcWRXoKiZCxEzYh3V3w5C760EbA4lYRhLQkMhnwMQBlzPeRvF2JoHrF2QrmXtx9YYAWHsynob3yVJaMAw5Gfz6NwGkvB6AoOskf4RIFUIv6f7MmBtwTkV/RZRuyxcGF92uwsgRLGnOVAMIy7ZgfCaJKR1FZlX9btAiG94UVeAkASVkTVhjh0d2040B3BzssV6Q4vRylIgDkSOc7gddu+Tt8JwpCmDEB0aNllO5zY3AeGnnjK6XwqCrgmEO3VfqKM3BQbryQXh7rjoioPOSWMdS8IqhIq8bRV6AdCWyWFrUcw4B0uo+DfYDMJ4rONlGa3zCIQuBBemQUjuBZ73hSQMDlXLkSQ8bjAfI14+sNaIfcKqjzDn0MUJaqhEqZkUqKbB8+EMf/Lio0/qqCsiuec2twTCa3bNAwSvcuDoFamjD6U62j3ntKd2u5gC4cdSGV12l5RRFIIVR1Zbduhp8reLnxDtGMUSymIEWJznVuwlMKGd+Am1KsJwuLO4PNO8gGMcJ0DDNQ3bT3upTIDfIK4MBDAlIAbqCcBuwcdoDI3ABc4npDwonXyKHM9tIezN0jqUIk2Nni0zcxuD8KqMXUMU7gmEx1IdLVN5m2lBqAofxr49XVJGYSoaY/LbKILu+APvg23yt3cevBzV3abfleF3adP0RYdXz8eZtNuUPapyjMWnsrOibzO+mhkdObe5TaijV9kaE9ueRSEaZhaTcTCHl4dJZbSXYnGezuedbsfjNw38OH93c/tHAeExS8Lktb9KIIxz/Xy4EI02IeQWZQDbvvv46/yA5za3b5GEqIKmFeLV2DDTy2GacFDs1G+FMnpYqN38gOc2t29ZE+ZA7iUpp0M/YdmmwmWOu8KP+NzMGJzb3N4Pwus+CK8SCI9lxEw/T/BliLHV40txbHuan+7c5vZ+dZQsosuvQxCqi5Jwf/hFPfb7elB13p34Eec2t7m9RxImW0xeFC7ZMHO6lCjYDBic+mmF9eOsjM5tbt8iCb9eLZdXX8cgfFCnwwUQtovVrmeVKV2ENw8zBuc2t2+ShCNui+Qn3F5Mmu8rnItdP15tMT/cuc3t+0F4lUC4umSYgSi228L+OXARvryJQUsNCFuK7HlPgWz0amgT/7GWA0+MpCTCiXwQbpIu4z+feyzCVYx0z1fMx8MpRlFkuKVr4BYj/ays4X0mDdHYPFbL+20aZv6Qbicd6+nCeASmCMvl/IZvPF9FnpP1uT+bn2J+LU+C19UxP1wZAt8cPguj0j3lR2/yw+k9Obktw3u9POI0POp5lR4LbeZvT56kyReZwTdSR6dBeNkwA8Hep9IJsepnET68ceHMuxIKihiO2OR0dYyLpk2mYLjnJL5MRYMHOQpHBU4aJ7GkyuoyfU/O1Hwex6gKf77n2FE4ISRGfj57TXHcKWFJ48j5TerZluQ0RtINgxDOBMoTCY5ZbTBFUa8lz9DGAzZ8+5RNhcmI3I3jvrFsgEsEOHghzstgcmO6XMWk4dgx021YvB/KmQQCLLonxaG80KmTB03PVEnH/BxN+ejyHTODgQrujqJo5dEajsq3zDLJBFny/cztsiRckruCUpn60Kuz0rkv9NFFj3GteVsZjd9O/C3mPD1DCeqQNKFTijpmVUCGBPDWG6Sep3R1FwJ/tYHmERDpY16fRWqZOH+8JdYXjakaKWrbWU+pE0Bjj50aFkiVthLArWmKAkVNkVXFrMACQixLQWn9XkZJn4jZhk+h63E1Ge7Nx9kPB8ULeg1D1kKmYwpWDgIhsn87kj6OiW8cXkyiy2kwkyCksHIaQyLkcUxFAOigXxhO6+JOnZ64SrwAP0d4tJp+nIhnh56hSSDkDuFZ42b+pkx6uB6YhbyZmXbekIRsl2FnfU8dPZwKN8TN9mNaEJ4PReD27cfd2yAM/O0E4UTCTDyDhVVIBgrtjBAqWQQO8U9YnnIuqE0Vjmk6WlVwYONh696ERZKJ2NfKhSzmkuwFvn38yccUQR0yCNd9EAYBIRWDkTFZEousuVG5CsvQEp4ovt8g899qm5jFA9a8cC5VrghCJZXIMwzLft5Y4ZV7IHSUUMIpWPxU8Flpug+LMtcUINxo6ZSvsq5Kmg/8a1PetdZHS1tMSsC+IxDqOyg8pfKjTZKQkzVFH/DVDL5XJeFVTmXaqc+DkNEyljRpnTtVCMjl28pozqyXX3/M4pMvnHkChcwMk9M9oZFTAkn+GJz/8Ss+Mv6sHM8/+2UyoWF9yiMc7jJVWg+EoKtaERjhoiTUmt4lNoCCE8AkTkZDI3Ui1fBcTg5OQggT/EcgLCRhH4S0QyiEnU7Uw31J6BwzDxi6O7jd+G+1wUEbQzolDqISyMPN0lWOltcHhDAtVMaiU274Ru2A+QfUUQahUAnhg2XuEHi4nn8MZ0n4KgivnksQlpJw350efymy60Uf3e16LsJXlNGcdKT7IIxfPoHF80wIGlL3vClASHQVNC0SCIPg0iuZVinnnklfKp4PdBhJQiKBGoMQarh4yYwvhGUfhBGf+A7lthyECh4rnhmE2mZAMVPpUZWaIEJEv1sSMuc3b8XkYR3GktC5oFWP5oMfACjJlfG0nhNqOenUDqHu+QfzSLomPX3rCVLMUMdPUCShuiwJIyS158XGDMLX1NGrMrO+tybcg+jbdj2aGejkYXc69LIIL8fKFOS/fRBql+o3kHYWVyEuI4QkYRUSCNfHIQhVH4ROiCtEqcr885oPKspCJBBGnVgzk2GvNk1fHXWORu5ZlhglPRrSL5M6arI6incQl5/rY7r9BMIkCVc9SYggDFMghAfAhpk4mARCkyQhFb1g/AN8TBpmHHViMSBJmEEoScm6YDQ3OOReDQ9WR+2AD4gMMwUIDdMWCIuWr2YQvi0Jr68y09PQOgq0TS9qLPReyizCfaNeyWC6BEIkcUkgRK5rL4sgIJWB/cZVaS1B6ugIS6K+GrLYaJenZglCt0EzHeTGs2BOINQau8sgDCEIubdJINSeRp6Fdx6KENrAYAyKyqgTxkZDC4E1v0Id9cCRgSoCHpmso8agOhoMZgUPQJjV0cpmEIIvQOu8DEwPF0pqyHPQXBrKYo1IOIR+7woQDiShVyMQwo1qWleaAoTJMBPo0WLFSM/qqENqyFkdfQWE11dXqTYhEz0dP5Z+QsyUf7jvSpKL2MevqrTKdNvF4qVo949DEJqBOhpXKCaX92MQxm+cOVwqEoJsJOcvT/sjE0r0JCHP9ERF2JeEvhhE5UofRwJhnId4HPGDh3zIOrlUYGmkDY28KkGI6mjBx5HGHDLTYVypBVUYZvjWdarsmzjCq8SXyPsuqaMwGF14EvjWDFuB+ArxDoX7Q4vzgf0IekoSlitPkYSuBKHUWeZhjgwzfMdFBQDSA7xY1WbwJRAuCkm4/68h21o/YgajRR96Yu8MjKOfS1h+fvV6yW+XJWHipu+po8la4XSFkpDL2Kc14TFJwpE66ioqYn9RHa2otr3NpekFhOg987lGNlfltWpDdehpJeg0dhoqzQpdloQ8TBE5+AHHI4eKhudVQcuPI0IQ4pGheIe9FDTiGS9sHYUFXwIhnqRV4jFOVyDDZMYPSUK8Jz2xJhQQymCPpWFGTiW007O8Gxpm6NEavoQtHi5wfPmZcq4A4fY/+q4Jbl+TOtorEvqCpPd5y/7pfx971bL3bTNsPZf+hCSMM8dR8c0sCYPNP8roi2YXRegbZtT0mhCtKvr1NaGZto5K0RdZgiafR2GYcYb5GR0zrokk1GKiteW6yReGGZiAVhXLMU2kbtpJlW9TrgnN24YZZ7UvQCiGGdAq0OICvHH8qyemThc36pAcOn111L2ijrrBmtCPDTPumNXR0jBj0sP1QvI1gy+B8FSWRVsOyX/Vx1HN+p26bQtX4e3uQbWX694Psy2Ym3BsmFF9SZjNKjQjkq9CQPjhyAErlwwzuCp6xTDjpdMShF4MopmHmKd1YZhxPpd28TIZFbu/i4o1Rowi2TCjsK5UqY6qwgfwLYaZig0zDp7CGISAahpiij3jaW/JOpodOtOGmT4IJ9RRU5ii3jLMkAUnPVwkbZ1BWIDwcDUhCq8KSTiM2F5tm54ZZvE6CJ+byaL1EyAM2UWB+hd8xVZmM+uZyRLp2VGd1VGfZo2WMLEpw4w1IpMmQOiofBmBMFtHXdWThGI4QWd9yOpo4CmZ3Opc4Z6QQBfzuvi16FXtVqVH4RU/4cAwg0rnyE8IxKi6T2PMCuCkn3AEwsw/OWGYMWIUDTr/tA0NM37gJ/SVq4pfyXlNWIDwvt1/HYtCAeHxY78q0wvaQkv1s7vffR8I9QiE7AmmMElTyJCEl6BF/YFVCnkiRGW0ZCZIklCJR+pu4Ky3JJMmQShTk2MDQtZvC0lYSaRAHM4xUNns0lkvE5lA6KoiYqbwIYYeCN8fMeP7znpUb0eSEOiTHaxCMwiTkcWjJ97liJlKIF+sAciRYqvCdpz0kGPQn/jXRid//hsRM8TbfFQpdnW2jvbU0bI8IXPNMC7/qP5tUIsC4PSoTu1zGTXzoyShT2IMI9bo91JcDmlflUPErMiKVIJloI6iUrURfU6XYWtx32YChFYi6OxgTViCUMqeAQizK8QnEFalOoq+l6Sk5kiSN0B4URLK3RoGIVYrHUrC5B7MVTUYUPG2jxQ7GopfghwLJ1fR4U6U0mHY2krC1qiPHgjdJ5GE/YgZrMC6ZhDqanZR9EHY7pcjfXT5Naujp8MwQGZX0FhATfo3QFhPcoROSEJivSfToISAxo+ebITBekf2fucd5UwECITWgLWAAdk5bM1TXLSBwGOdue+1D1Q3FOKm3ZHqs5TWUYkjqcS9ZqHStRPZyCB0PPJUhJCCyvHnPkAhNjHMYF3gXsQMLwnH6ijUB0/RQ6EfwA1pQyYVXnP4AIIsl0swFyBk2OZfOO0ShG3KorAqB3BTp8VVBgHcmC4SHwUeZysuAE7P0NJDtxpeaL+jfLJCEmpT/MLNIOxJwn0pCnsFCv8TImYWYxCWQWr7w/2bhpl/vygJrSQEMTBTapKAUPJ2THbpBeLP53SglMMUCsNMJRT7vqxEZgI7yOg0k/JtxGpheUxUkZ6NsclPmEFoi9gZuhEe05SfMKgcHWsyp3/OIkqeSwyfXaleKpOTnK9sLFZlKpPrFTVNIORgzyrtkSu7ZEiC/IoqGaVtevQMDtNLZbIplUkPXKyDkgXiFjTyaCudDDP8cCWQfgZfD4TL4aKQDKbX/zIqjfZCkrCs1/vcHL9HHd2s6duN3+v6Lm31lCm6XvP6HzZhIezV2n+QRF37YY2am/8AG3G+rT9Ini+ubfyHD5KD6j8Uqal8GJwItXPV2scOfLri5gOPiTNO4cp+Tdfd4DDgMEvGinXKocXJ+WHNl1unfGK5Dh+HveV9uLLzaWB07HoNIITL0/U+bKiXD2AnWotG4en54AU/qLIfzrfl68TNaQ9n6HqbPkTpSoOg52/LR188rPTByE2l78PI0OFU/IsvHn4Z4t8PMOr0YHkk683mgy2HPDdURzPVIS8KeUn4lz+rP30c+QkV+C16+uin7wFh2d5BaT8+JCSD4t0/CyX+8e+n31dPMXMt4m+XhIU+Whal+NcIwp3ajtRR9dCjstjf/q0g/K5WgHD+Gv+u2sbOZae+QxIux2lMURv9w5+RBj/XQ0tw6nkl9t03OOt/XMsFmOZv8e9MYK/MZn4K3ywJsyi8ykvCv/zhX+CIU53aKdH9/tbkrXX5fqK9J8V3bnObQTjg3kZtNILwTxdOm0lF5za3HwdCLIG9HzJbxPaHCMJ/i3h72EnrwTC1B8jqfaX9z39I1s/GvLn9dEnYK03IS8LfSR2dG/jVnJsfw9x+riTsgZCXhL+/oo7+czWjq9neN7efDsJlBqEsCWdJWMDw79QE68P83fzjqKPir1/mJeFPBiHOa37JU9y8DYNPqjzabD6l8yb7McMuzGS/RnaY4l/5bmqEr0ETOzSD4/NnM9lH/8L5nEvxCEV+urkwJDOx0Ux8NOVjNBfuzrx113P7GyShLATTu7/8LiDc4X8KLCw79bh7fERTzQO4K3DXA+7G90c8EDN/k0FH0amjhtnVFSaaVZojheNbh7GQ+NExK4PGjIZ4LAVMruM7NJPY+KaqAqaR4kH4udIUyUnBZUGnMEuescz1xQqmY04WXzHxhsdEPI9UheqIqe3YleUcpA1/smoFJ8tS0eqwottylMzH8ZO6pIex+DlQBCazYDhdzGm6I8ojqZB4w9AAN0fcxxmJTMlhnU5RpNbRXgvPCbnHC8bjtKDVkuSHgWO6TM/A/GJIdNZcJOBIGSypo0CjqpDzJ507++V/BAhvl0MQXg9ASP6I3Xv9Ert3bMGJoJ3DvEFfQeobcW5rot7NtSgQoZh5ox1x3HqZycSLwXwsmFcX4BAgMK0omY64W1JtCJfqNljuFRM3MDUpxL9O26JT3P0BI6gp1pom44azfCMIcbg8v+PkX/FfCBDHxF/jcDyh4LRw1Dfk1lFGFJWySCCEu3RIY4i3EPgxBUrH08wA5TRTLGm5jwTv9DCpXAfAvaAoqEoQBhk7ZwIyCB0/xCPvr4S1lVLM6A5SfpWfI9R+hCQ8lCAsl4S/s4viF3V7e16o27Pafj7efnn5bXH+fD7ff/7t/vbL4ryFffH/++1pu4U/6tTcxg2nbdwV993eKti9G4Owl8rkMdkFORlEclDWEeUH0beOv81a208rDsRnvFJYf2bSxPQLSkUIkqYokjBwcYg7i0dRjyrTQPhqvcJpKRiVROGUJGhxalNGnmTtcyoiYpjTd1h8bDY5hcFyqjINU2pa+BKEGFldJWbthBuSmBrEq97QzWnhhtSU2WEwblrqA5TjSUprkoSYFZjyLF0BQpUCAgPhWMsPSCIPL5Mc7QzCHwFCjjlb9peE17+LJDx1zc35qblt29OpPdd13bWH+vB8aNpzUy+eIoi7+67ZnrsmHri4efrcHppFt9i2z123b5vuc3d+etodxxbHHgg3icMJM+uJBj9wfmwiy0ZOE9iyrgoGbpNyi6YZuEUQwFxcEazkACNoSMmvHtJqBYQ2UfFbmbIbophiqifGagKTQflWhQxCzIpAtg4oc6EEhEZKY+iSdqwPQvwowksSlY0cgndNOfhWCDuEjiMOIOd9TamjRmgEXgGhZgr9Cnhj6MdrBuHPNcwsR0tCBCGkLbX1+dDWXVefn8/1oeu6c1e3T83zuetON+f2Jm6v7+/rQwTn+Tnuam6aQxOP6dqnc9duu6Z52v76+LokXKUEVKHBB8UqTuoxDX7cstqUtSiYZ8mZXi0KX0kerEg5rIcCmhYRtHgBYeaBV8TNkEAYtPCsEesNgtBiJQofF66bRAoIgpiSC4eSEF+RtwpWWSMQhtdBGDLIaZnsCYSWHwTeIBc8Q6nJGYXOIMMMcM7pzGDTWxPGn4spdVQVTGucncg6wQzCnwbC+5sChNlL+K8MwpU6HQ6fz213/7SMiKrr57p+UidVn+q2i7hD2Rh3n9T9Yd9+btr6pt4umjYC9VxHmRgP6A6HertjEIY4IUfkv2wd6JH/QjEFly0q8K0fQYJxRnaFvKN2yLZm+pKwKomeSOLp1CvIvKQXJhCGAoQGlkDYp2GSQGJYMkSVf7zLyywqTgHdT4BQB4bOFAjL7NYpEPriN0tl2hfqmlQGL+xXoQAhltSIPUyqo6hMi4CcloR8k7FTSoR3YxCuZhD+eEl4nR2GfyAQ/mmnbg+HKP+iaLup2+c2grCrVZRwi/ihO9QRbFHkPbdndW4PHRzY1hGDz21TNwDCQzw+HvHLx8ImOkWDvxrWorBaaPAT++bxA/7UUz8fjpd5R7UTM2VJeUicaanCRRQSK5pDJQhTp4oWWCwJ19oyCFfII+qzYULWhJYLaAzUUcMgtCrz9hcglF0XQGhVj5PmEgiF/SqBEMDrcDyeuACsGYCwiAIaSULSSDQNokJiDKovU4AwGHM3S8KfAMIyZi2B8NR2Xfz/uWsOLbyJmIuLRABfRN3nmw53HyIw466ubdtD20TR17bdU9c2DWxrmigJ3wChl2qyIgnjHoZQYotxhSVET9eiEAZuR4woeigJ4xTmKRZ0qkuEvENQc8JJVV0GIWp1VCtCZzJRm5BjgstqGXQcfzkQhBGL5GXBqQ01KaZAWBVkjpMgBB0B2GucnwChvgxCeAoVclxDSU/2LIS+YcYFc3FN2ONLhXrCd1EJD1UJQp34PGYQ/lDDzPVzb0nI6ui5a/btwyJqnlEEdocGDTP3XV2ruPI73UTR99yppj7dRNl3aFUTd7fxz3O3uImScbG7ac9PT9vFYwbhiIE7eKq5WaqjzuVaFKHgaxnWophk4JZCCbosjQYTG1aejv2DQslG5cVS3YgChKANkiSs/CqtCf0GzCPEslhUAYZyS3FETgvHTRJXaXU6loRYMiJMgRCY4otKFGMQurE6ysRUaIxB4CPLb/Z3lJIQbupVdVSKfGBH8ceiD0LwpWghaZ3bD5KEV1+v94n/HpeE4idc3NaNut3en7aLL6fzuelOp606Ybu9XXTnc3va/XK6hY+ne7U93Z8+n39Z4O7fFr/cLuLm2626LAm90DKV6mjmoTXWV1Q02oSSBv81EGpr0KQ+UEc1zB+aYki3pnIV6aCjxpaK+xEI0RbpaJKrDMJ1prwFZrHEnmZQAyRWRsP6XxBOo1wssAdCIlk1UyCs+McnxGGZKRBWUyBMDxNqeYBYDtnfkU2tWCbNVK8ZZtLKD4uqSke9NaGt3AzCHwNCiZi5Xj73l4RlAPdlP33zsmg+v/96wBg9koSeYsdKdRTpQZMJJTFwu3FptDAujSac9iN1FMJCitUfkxGGck2I84xACGCjNaF2mdtee9yWKPVDQj8GCOixYSb1n0AYBIRaF1z4fXXUmM0x15x4lyTEHyR+mATIkOrUl5qvwfWguSgJq+SEJRAGMA2PQUj6/QzCHygJl88DbVRA+BFIt5FdJgWjxfMe1MNuodQ9dfNxB593L49qt1vBTvX4MR73+ACnPE6FrU3Voqiyi4LNKqu7qBAfN1yvDInYWR1dH5kGP0lCo0YM3ENJ6Kj2yaeNjICsGyUIdQJhwKJC4NAACSyScKXRMR7/3a1gbKzrEp9mnMh6bJjRyIZtSxBqI3Vb4q5JELryVyXXR7pjEFLNP3FR5LJQSStA1ZR1fC0FupM6SrzhrxlmcuSQp2CDKRBWMwh/KAi5ZQeFgPCojvEX+Rf1y/Ze7RZ/BShuAY6/xbNfHtRiC4CL72+BgQtDt+KW7fbE4nOxnY49vliLIiBzNhc2C8wmWoAw16yvdDnxrO77CQsQCiWULPqslhGQOKmmQKgBHCQJNUcFkCQkO5FVFDVZ1JbYYIjPQBIKZWrlRco6lS5rdclTOjLMFCBkP6EWPyFok+ysx+Vo6PkJCYRKilE4PQQh3tur6miQ0oUYKarVDMKfCMKXHkvT/msfhOitvwFfX10f2ib+2XWL26dT/dK1zf2pPZxV03RPzdMvdXta3CyaBgw3dXvomptT05zAXnp+fy0K9swzuXvCV1ZHdeKQlQoQxYzY9FwUaAJkD8FKQGhQEjLzNJleBpJQOsWINZKEuDjLa0JaqkkRsRSb7eMvkFEjSci1LVLlUZVUWyztEEdvtXsbhBupwSbhQYAkg1YmcLUkBdyWIFRGpTXhsQy/4eXma4YZDdVFQ2G8nQChnteEP0MSLrODIoGwrusWTKPd+Xw41OdDvbg5td1NczifW7CDtu2iW57b+ul2/wRb2qbtlk3z3NT16dDW59N7JSFMK+ORi11qUUAcsrFkHaQ3EMtIVpeoFhooKG1wwnhMp7CpdgWWp/WVs/GNP/I0wtkGCmaUKBBV6bELhytFrlctfjbFBYIdBYpmEHJ9WkvFpkOKaCkCVKX4NS8HpVKTONKzvPaS12HNRRAG7OsId2up3pI1nqv6wmJNw+3Fe4S9pi8JE3biza3RUJVA6ItlJg43aGuMgNBwmkTpQRn4CT19VXMWxY80zDyPloSsjn7smm3TRSQeIFytXrTd4nBo21p1DfjnVbOvb8EX2Db14UZFMdipKBPPdfz8fGi7Q3eaNOsUIEwhI0HKXLtU3Ciwt8EI23oq1BwyDb5NLoGiXHZ+l+njWaUlNxeYCFMXTHdPBRNwBleSc2GqoqgSf8IomrJzW1adKapjZxOt9hwQ7bhcp6sYahGlqVaTGEZ0AqE8AM1uF58K7bqszBqdqmwLCF1eaKfjkiSke3DZYp346xU9fXaEki+jQHOKnUgpYDMIf5yfcLwkZBCuuvOi7qJYi6BbdqCV1hAW2qgW4kNvVBNF4yECNG4+HF6ihKxvz+ebqMCCa7/tuuZ2GoRYP9bC95i0sdXagWCDTRgzjEcEFywFQeIb2KIDFjRxUIfZQQ6R5fM4KjTAdpoc3um04qIL4WuUfs5wSWrIHQIxAml6nky3cGVHAdmB8xpXOBp1XPF2EhkusFHF8vVQQXRcIRovdhfCke5X+ubagj7QHcY3Ik8slZIPRcp86ivKKrxGekjxbbwNS6p23BuAQt+Klcf5fNPpechGKtSahqxzMzQwxYONJ8hRsJXzPCmWlqpMuTnF9wero/vCQZHV0S4CCjzwTcRbRF9E4QnCZkAdrcE7H9d99TMsBKPAxJAaCOHGiG8I/m7O7yZILEw4b7Nqu7+dgTtbbN6eSXKR1bde5E5t3qSZN9+g1JWdHQHym8ldc/sfBcLlPjWh3uYGkpDSKA6H5w7DQrs6LgLrUx0lXXO7iLhTzbk+nL8AOuP7tll0IBqf6vihO59qiDr9OQzcruoj6XtAWL0fhD+1me+UJ6uZ6fofAoT3Tzddh3Gf8eU//gjtP//M7f/xMQ91c1+ftrf392q7+PjX+y//N279FH93t1+O6ssXtYi/wQ8ncCf+H3iJh91joW11v91+Pp1+0sSVgLHvX5TYv72LH3QvM1/TPzUIJ7d+Ub9GIP1KH3bo+esfOGQCZhIa4KBJyufM0j23ub0PhLuXxW6xe3mJL4svxZ5MOfpI/3YRYFSMIr08Pma47QpIzm1uc3t3+/9XLR3ofmfzgwAAAABJRU5ErkJggg==',
}
SERTIFIKAT_VARS = ('У компании есть сертификат отечественного производителя «Сертификат CT-KZ», международные сертификаты '
                   'и сертификат ISO 9001-2015.')
MARKA_GOSTER = {'chint': 'CHINT', 'abb': 'ABB', 'abb acik': 'ABB', 'schneider': 'Schneider Electric', 'tekled': 'TEKLED electric',
                'ekt': 'EKT', 'iek': 'IEK', 'legrand': 'Legrand', 'harmony': 'Harmony', 'saiman': 'SAIMAN', 'teksan f': 'TEKSAN',
                'kaztek': 'KAZTEK', 'viko': 'VIKO', 'cnc': 'CNC', 'ekf': 'EKF'}
_GP = None
def gorsel_paket():
    global _GP
    if _GP is None:
        _GP = dict(GORSEL_GOMULU)
        for g in GORSEL:
            if os.path.exists(g):
                try:
                    _GP.update(json.load(open(g, encoding='utf-8'))); break
                except Exception:
                    pass
    return _GP
OTO_SATIR = 40

MARKA = [('chint', 'KZT', 0.60, 'EVET', 'liste ×0.60'), ('teksan f', 'KZT', 0.60, 'EVET', 'KAZTEK/TEKSAN ×0.60'),
         ('kaztek', 'KZT', 0.60, 'EVET', 'KAZTEK/TEKSAN ×0.60'), ('tekled', 'KZT', 0.60, 'EVET', '260901 uygulaması ×0.60'),
         ('ekt', 'KZT', 1.00, 'EVET', 'onaylı 260927'), ('iek', 'KZT', 1.00, 'EVET', 'KONTROL: katsayı teyitsiz'),
         ('harmony', 'KZT', 1.00, 'EVET', 'KONTROL: katsayı ve liste tarihi teyitsiz'), ('saiman', 'KZT', 1.00, 'EVET', 'onaylı 260927'),
         ('abb', 'EUR', 0.405, 'EVET', '1,35 × 0,30'), ('abb acik', 'EUR', 0.3105, 'EVET', '1,35 × 0,23'),
         ('schneider', 'TRY', 0.60, 'EVET', '×0.60'), ('legrand', 'TRY', 1.00, 'EVET', 'onaylı 260927'),
         ('pys', 'KZT', 1.00, 'EVET', 'piyasa, net'), ('cnc', 'KZT', 0.55, 'HAYIR', 'yalnız referans'),
         ('ekf', 'KZT', 0.67, 'HAYIR', 'yalnız referans'), ('klemsan', 'EUR', 1.00, 'HAYIR', '2022 verisi')]
CSV_MARKA = {'CHINT': 'chint', 'KAZTEK': 'teksan f', 'EKT': 'ekt', 'HARMONY': 'harmony', 'SAIMAN': 'saiman', 'ABB': 'abb',
             'SCHNEIDER': 'schneider', 'LEGRAND': 'legrand', 'CNC': 'cnc', 'EKF': 'ekf', 'KLEMSAN': 'klemsan'}
SINIF_TIP = {'kontaktor': 'contactor', 'role': 'rele', 'buton': 'button', 'termik': 'rele', 'fu': 'aksesuar', 'tms': 'mccb',
             'opn': 'aksesuar', 'blok': 'aksesuar', 'altlik': 'aksesuar', 'olcu': 'indicator'}
KUR = ['KZT', 'EUR', 'USD', 'TRY', 'RUB']
# tip, grup, tel/adet (None = 2×kutup+N), tel m, norm dk/adet, sıra, pabuç eşiği A (None=YÜKSÜK), sıra kategorisi, yan boşluk mm (boş=KONTROL)
TIPS = [('korpus', 'KORPUS', 0, 0, None, 1, None, '', None),
        ('rubilnik', 'GUC', 3, 0.5, 10, 2, 0, 'MCCB', None), ('bh', 'GUC', 3, 0.3, 4, 2, None, 'MODULER', None),
        ('mccb', 'GUC', 3, 0.5, 15.3, 3, 0, 'MCCB', None), ('ppn', 'GUC', 0, 0, 17.9, 3, 0, 'MCCB', None),
        ('ats', 'GUC', None, 1.0, 194.6, 3, 0, 'MCCB', None),
        ('motor', 'KUMANDA', 4, 0.8, 2, 3, None, '', None), ('shunt', 'KUMANDA', 2, 0.8, 2, 3, None, '', None),
        ('kilit', 'AKSESUAR', 0, 0, 8, 3, None, 'KONTAKTOR', None),
        ('sayac', 'OLCU', 10, 1.2, 8.1, 4, None, 'OLCU', None), ('tt', 'OLCU', 0, 0, 5.9, 5, None, 'OLCU', None),
        ('oto', 'GUC', None, 0.1, 4, 6, None, 'MODULER', None), ('rn', 'KUMANDA', 4, 0.8, 0.4, 6, None, 'MODULER', None),
        ('fotorele', 'KUMANDA', 4, 0.8, 0.4, 6, None, 'MODULER', None),
        ('uzo', 'GUC', None, 0.15, 4, 7, None, 'MODULER', None), ('dif', 'GUC', None, 0.15, 4, 8, None, 'MODULER', None),
        ('contactor', 'GUC', 3, 0.6, 5.9, 9, 96, 'KONTAKTOR', None), ('kontak', 'KUMANDA', 2, 0.8, 2, 9, None, 'KONTAKTOR', None),
        ('condencator', 'GUC', 0, 0, 4.2, 9, None, '', None),
        ('krm', 'KUMANDA', 6, 0.8, 0.0, 10, None, 'KAPI', None), ('avr_rele', 'KUMANDA', 8, 0.8, 21.8, 10, None, 'MODULER', None),
        ('rele', 'KUMANDA', 4, 0.8, 21.8, 10, None, 'MODULER', None), ('timer', 'KUMANDA', 4, 0.8, 21.8, 10, None, 'MODULER', None),
        ('lamb', 'KUMANDA', 2, 0.8, 8, 10, None, 'KAPI', None), ('indicator', 'KUMANDA', 2, 0.8, 21.8, 10, None, 'KAPI', None),
        ('button', 'KUMANDA', 2, 0.8, 8, 10, None, 'KAPI', None), ('perekluchatel', 'KUMANDA', 4, 0.8, 21.8, 10, None, 'KAPI', None),
        ('rozetka', 'KUMANDA', 3, 0.8, 21.8, 10, None, 'MODULER', None), ('stable', 'KUMANDA', 4, 0.8, 21.8, 10, None, 'MODULER', None),
        ('aksesuar', 'AKSESUAR', 0, 0, None, 10, None, '', None),
        ('bara', 'SARF', 0, 0, 30, 11, None, '', None), ('izolator', 'SARF', 0, 0, 1, 11, None, '', None),
        ('nbara', 'SARF', 0, 0, 2, 11, None, '', None), ('kanal', 'SARF', 0, 0, 5, 11, None, '', None),
        ('kablo', 'SARF', 0, 0, 1, 11, None, '', None), ('klemens', 'SARF', 0, 0, 2, 11, None, '', None),
        ('yuksuk', 'SARF', 0, 0, 0.1, 11, None, '', None), ('pabuc', 'SARF', 0, 0, 2, 11, None, '', None),
        ('bag', 'SARF', 0, 0, 0.0, 11, None, '', None), ('tarak', 'SARF', 0, 0, None, 11, None, '', None)]
KORPUS_NORM = [('KORPUS_KUCUK', 8, '< 700 mm'), ('KORPUS_ORTA', 30, '700–1499 mm'), ('KORPUS_AYAKLI', 90, '≥ 1500 mm')]
ESKI_NORM = {'KORPUS_KUCUK': 21.3, 'KORPUS_ORTA': 245.2, 'KORPUS_AYAKLI': 244.7, 'lamb': 21.8, 'button': 21.8}
PARAM = [('DK_UCRET', 100, 'KZT / dk'), ('ISCILIK_KAT', 1.0, 'norm çarpanı'), ('MONTAJ_MALZ_KAT', 0.4, 'монтажный материал = işçilik tutarı ×'),
         ('GUN_SAAT', None, 'saat / iş günü (OZET gün hesabı)'),
         ('KABLO_FIRE', 1.10, ''), ('KLEMENS_YEDEK', 1.10, ''), ('YUKSUK_FIRE', 1.10, ''), ('PABUC_FIRE', 1.00, ''),
         ('BAG_ADET_M', 0.5, 'kablo metresi başına bağ'), ('KUMANDA_KESIT', 1.5, 'mm²'), ('OLCU_KESIT', 2.5, 'mm²'),
         ('KABLO_MAX_IN', 250, 'A; üstü bara/doğrudan'), ('KLEMENS_MAX_KESIT', 35, 'mm²; üstü klemenssiz'),
         ('BARA_ESIK_A', 125, 'IN_ANA ≥ eşik → faz barası (ayrıca DİKİLİ · ВРУ/ГРЩ/АВР · CIKIS mccb/ppn/rubilnik)'),
         ('BARA_DERATING', 1.0, 'bara tablosu kullanıcı değerleri (25 °C, doğal havalandırma) — ek düşürme yok'),
         ('BARA_PAR3', 2.25, '3 paralel şerit kapasitesi = tek şerit ×'), ('BARA_DIKEY_ORAN', 0.70, 'dikey bara = korpus yüksekliği ×'),
         ('BARA_FIRE', 1.10, ''), ('TARAK_BOY_MM', 1000, 'tarak parça boyu mm'),
         ('KORPUS_IC_PAY_MM', 100, 'iç ölçü = dış − pay'), ('KANAL_FIRE', 1.10, ''), ('IZOLATOR_ARALIK_M', 0.5, ''),
         ('KAPI_TOPRAK_M', 0.5, 'kapı başına 6 mm² kablo'), ('MODUL_MM', 17.5, ''), ('KORPUS_YEDEK', 0.10, 'modül yedeği (aksi belirtilmedikçe %10)'),
         ('BAKIM_PAY_MM', None, 'YERLEŞİM: sıralar arası bakım/kablo payı mm — boşsa KONTROL'),
         ('UST_ALT_PAY_MM', None, 'YERLEŞİM: üst + alt toplam pay mm — boşsa KONTROL'),
         ('MONTAJ_PLAKA_PAY_MM', None, 'YERLEŞİM: korpus derinliğinden montaj plakası/ray payı mm — boşsa KONTROL'),
         ('KAPI_ARKA_PAY_MM', None, 'YERLEŞİM: kapı cihazı arkasında bırakılacak boşluk mm — boşsa KONTROL'),
         ('KANAL_DOLULUK_MAX', None, 'YERLEŞİM: kanal doluluk üst sınırı (oran) — boşsa KONTROL')]
KABLO_TABLO = [(0, 1.5), (10.01, 2.5), (16.01, 4), (25.01, 6), (40.01, 10), (50.01, 16), (63.01, 25), (100.01, 35), (125.01, 50), (160.01, 70), (200.01, 95)]
KLEMENS_TABLO = [(0, 2.5, 5.2), (2.51, 4, 6.2), (4.01, 6, 8.2), (6.01, 10, 10.2), (10.01, 16, 12.2), (16.01, 35, 16.0)]
# kesit, kablo ПВ3, НШВИ, JG, klemens L/N, klemens PE, uç kapağı
SARF_TABLO = [(0.75, '151.01.PV3.0,75', '311100813', None, None, None, None), (1, '151.01.PV3.01', '311100815', None, None, None, None),
              (1.5, '151.01.PV3.01,5', '311100817', None, None, None, None),
              (2.5, '151.01.PV3.02,5', '311100820', None, '130300223', '5130300225', None),
              (4, '151.01.PV3.04', '311100823', None, '130300041', '5130300225', '130300142'),
              (6, '151.01.PV3.06', '311100826', '260100138', '130300042', '130300227', '130300142'),
              (10, '151.01.PV3.10', '311100828', '260100153', '130300039', '130300227', '130300143'),
              (16, '151.01.PV3.16', '311100829', '260100165', '130300193', '130300228', None),
              (25, '151.01.PV3.25', None, '260100167', None, None, None),
              (35, '151.01.PV3.35', None, '260100158', '130300040', '130300365', None),
              (50, '151.01.PV3.50', None, '260100168', None, None, None), (70, '151.01.PV3.70', None, '260100169', None, None, None),
              (95, '151.01.PV3.95', None, '260100140', None, None, None)]
SARF_TEK = {'SARF_STOPOR': 'YZN11DF-003-K03', 'SARF_BAG': '240100162', 'SARF_IZOLATOR': '32101DEK',
            'SARF_TARAK_1P': '811004', 'SARF_TARAK_3P63': '811002', 'SARF_TARAK_3P100': '811006'}
SARF_TEK_AD = {'SARF_STOPOR': 'klemens stoperi', 'SARF_BAG': 'kablo bağı', 'SARF_IZOLATOR': 'bara izolatörü',
               'SARF_TARAK_1P': 'tarak PIN 1P (tek fazlı pano)', 'SARF_TARAK_3P63': 'tarak PIN 3P 63 A (IN_ANA ≤ 63 A)',
               'SARF_TARAK_3P100': 'tarak PIN 3P 100 A (IN_ANA > 63 A)'}
SARF_TEK_TIP = {'SARF_STOPOR': 'aksesuar', 'SARF_BAG': 'bag', 'SARF_IZOLATOR': 'izolator', 'SARF_TARAK_1P': 'tarak',
                'SARF_TARAK_3P63': 'tarak', 'SARF_TARAK_3P100': 'tarak'}
SNI = [('130200364', 6), ('130200365', 8), ('130200361', 10), ('130200362', 12), ('130200363', 14)]
# bara (EKT, м): artikül, kesit mm², tek şerit A, 2 paralel A — kullanıcı tablosu 260930; tek şerit akımına göre artan (fiyat da artan)
BARA_CU = [('130300473', 60, 245, 417), ('130300321', 75, 300, 510), ('130200207', 80, 320, 544), ('130300322', 90, 350, 595),
           ('130300554', 120, 420, 714), ('130300812', 150, 450, 765), ('130300323', 160, 480, 816), ('130300324', 200, 540, 918),
           ('130300810', 200, 540, 918), ('130200208', 250, 600, 1020), ('130300811', 360, 750, 1275), ('130301396', 480, 830, 1411)]
BARA_AL = [('130200188', 60, 160, 272), ('130200213', 80, 205, 349), ('130200152', 90, 228, 388), ('130200143', 120, 280, 476),
           ('130200197', 150, 320, 544), ('130200153', 160, 325, 553), ('130200009', 200, 365, 621), ('130200058', 250, 400, 680),
           ('130200330', 360, 485, 825), ('130200069', 480, 530, 901), ('130200057', 640, 620, 1054), ('130200052', 1000, 720, 1224)]
BUK_TABLO = [(k[0], None, None) for k in SARF_TABLO]
ICERIK = [('MCCB', 'MCCB', 'giriş koruması / MCCB / rubilnik / ATS'), ('MODULER', 'MODULER', 'NXB, NXBLE, NXHB, röleler'),
          ('OLCU', 'OLCU', 'sayaç + AT, mühürlü bölüm'), ('AVR', 'KONTAKTOR', 'АВР kontaktörleri + mekanik kilit'),
          ('KONTAKTOR', 'KONTAKTOR,MODULER', 'yol verici: kontaktör + termik + kumanda sigortası'),
          ('KLEMENS', '', 'klemens sırası'), ('BARA', '', 'bara bölgesi')]
KATEGORI = ['MCCB', 'MODULER', 'OLCU', 'KONTAKTOR']
KORUMA_TIP = ('mccb', 'oto', 'bh', 'rubilnik', 'ats', 'dif', 'uzo', 'ppn')

F10 = Font(name='Calibri', size=10, bold=True); F10n = Font(name='Calibri', size=10)
TOTF = Font(name='Tahoma', size=9, bold=True)
SARI = PatternFill('solid', fgColor='FFFF00'); PANOF = PatternFill('solid', fgColor='FFD966'); GRI = PatternFill('solid', fgColor='D9D9D9')
ACC = '_-* #,##0_-;\\-* #,##0_-;_-* \\-??_-;_-@_-'
ince = Side(style='thin'); KEN = Border(left=ince, right=ince, top=ince, bottom=ince)


def birim_norm(b):
    b = str(b or '').strip().lower().strip('.')
    if b in ('шт', 'adet', 'ad', 'nan', ''): return 'шт'
    if b in ('м', 'мт', 'm', 'метр'): return 'м'
    if b.startswith('упак') or b in ('уп', 'pk'): return 'упак'
    if b.startswith('компл'): return 'шт'
    return b


from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
KAZ = str.maketrans('әғқңөұүһіӘҒҚҢӨҰҮҺІ', 'агкноуухиАГКНОУУХИ')
def temiz(x):
    if not isinstance(x, str) or x in ('nan',): return None if x == 'nan' else x
    return ILLEGAL_CHARACTERS_RE.sub(' ', x).translate(KAZ)


def num(x):
    try:
        v = float(str(x).replace(',', '.'))
        return None if math.isnan(v) else v
    except Exception:
        return None


def parse_ozellik(ad, tip):
    s = str(ad or ''); o = {}
    m = re.search(r'(?<!\d)([1-4])\s*[PРpр](?![a-zа-я])(\s*\+\s*N)?', s)
    if m:
        o['kutup'] = int(m.group(1))
        if m.group(2): o['_n'] = 1
    m = re.search(r'(?<![\dkкKК.,/-])(\d{1,4})\s*[AАaа](?![a-zа-яC])', s.replace('кА', ' kA').replace('kA', ' kA'))
    if not m:
        m = re.search(r'\b[BCDK](\d{1,3})\b', s)
    if m and tip not in ('korpus', 'kablo', 'klemens', 'yuksuk', 'pabuc', 'bara', 'nbara', 'kanal', 'bag', 'izolator', 'aksesuar', 'tarak'): o['in'] = int(m.group(1))
    m = re.search(r'\((\d+)\s*(?:шт|pcs)|\((\d+)(?:/\d+)?\)\s*(?:NEW|!|IEK|UNIT|КВТ|TDM|$)|(\d+)шт/упак', s)
    if m and ('упак' in s.lower() or tip in ('yuksuk', 'bag')):
        o['paket'] = int(next(g for g in m.groups() if g))
    if tip == 'korpus':
        m = re.search(r'(\d{3,4})\s*[xXхХ×*]\s*(\d{3,4})(?:\s*[xXхХ×*]\s*(\d{2,4}))?', s)
        if m: o['y'], o['g'] = int(m.group(1)), int(m.group(2)); o['d'] = int(m.group(3)) if m.group(3) else None
        m = re.search(r'ЩР[вВнН]?-?\s?(\d{2})|S/U\s*(\d+)|(\d+)\s*modul|\b[xXхХ](\d{1,2})\b', s, re.I)
        if m: o['kap'] = int(next(g for g in m.groups() if g))
    if tip == 'kanal':
        m = re.search(r'(\d{2,3})\s*[xXхХ×]\s*(\d{2,3})', s)
        if m: o['g'], o['y'] = int(m.group(1)), int(m.group(2))
    if tip in ('kablo', 'klemens', 'yuksuk', 'pabuc'):
        m = re.search(r'(\d+(?:[.,]\d+)?)\s*(?:мм|mm|PEN|-\d)', s) or re.search(r'\*\s*(\d+(?:,\d+)?)', s)
        if m: o['kesit'] = float(m.group(1).replace(',', '.'))
    if tip in ('contactor', 'mccb', 'rubilnik', 'ats') and 'kutup' not in o:
        o['kutup'] = 3
    if tip in ('oto', 'dif', 'uzo', 'bh', 'rubilnik', 'rn', 'fotorele'):
        k = o.get('kutup', 1) + o.get('_n', 0)
        o['modul'] = k * (1.5 if re.search(r'NXB-125|ВА47-100|NXC-(4|5|6|8|9|1\d)\d', s) else 1)
    if tip == 'klemens':
        kk = o.get('kesit')
        w = next((x[2] for x in KLEMENS_TABLO if kk and x[1] >= kk), None)
        if w: o['modul'] = round(w / 17.5, 3)
    o.pop('_n', None)
    return o


_DF = None
def df():
    global _DF
    if _DF is None:
        import pandas as pd
        if not BIRLESIK: raise SystemExit('HATA: birlesik.csv bulunamadı — aranan: ' + ', '.join(_yollar('birlesik.csv')))
        _DF = pd.read_csv(BIRLESIK, encoding='utf-8-sig', dtype=str); _DF['ARTIKUL'] = _DF['ARTIKUL'].str.strip()
    return _DF


_EB = None
def ebat():
    """ebat.csv → {(MARKA, ARTIKUL): {...}} ve {ARTIKUL: {...}} (marka bağımsız yedek). Dosya yoksa boş."""
    global _EB
    if _EB is None:
        _EB = {'yol': None, 'mk': {}, 'art': {}, 'cakisma': {}, 'diger': []}
        var_ = [y for y in EBAT if os.path.exists(y)]
        for yol in var_[:1]:
            for x in csv.DictReader(open(yol, encoding='utf-8-sig')):
                a = str(x.get('ARTIKUL') or '').strip()
                if not a: continue
                d = {'id': x.get('EBAT_ID'), 'seri': x.get('SERI')}
                for k, c in (('g', 'G'), ('y', 'Y'), ('d', 'D'), ('modul', 'MODUL'), ('kap', 'KAP')):
                    d[k] = num(x.get(c))
                key = (str(x.get('MARKA') or '').upper(), a)
                o = _EB['mk'].get(key)
                if o is not None:
                    if any(o.get(k) != d.get(k) for k in ('g', 'y', 'd', 'modul', 'kap')):
                        _EB['cakisma'][a] = f"{o.get('id')} ↔ {d.get('id')}"
                    continue
                _EB['mk'][key] = d
                _EB['art'].setdefault(a, d)
            _EB['yol'] = yol
        if len(var_) > 1:
            ilk = open(var_[0], 'rb').read()
            _EB['diger'] = [y for y in var_[1:] if open(y, 'rb').read() != ilk]
    return _EB


def ebat_bul(marka, art):
    e = ebat(); a = str(art).strip()
    return e['mk'].get((str(marka or '').upper(), a)) or e['art'].get(a)


def csv_bul(marka, art):
    d = df(); m = d[d.ARTIKUL == str(art).strip()]
    if marka:
        m2 = m[m.MARKA.str.upper() == marka.upper()]
        if len(m2): m = m2
    return None if m.empty else m.iloc[0]


def cnt_le(tbl, x, i=1):
    v = None
    for row in tbl:
        if row[0] is not None and x >= row[0]: v = row[i]
    return v


def _bara_ok(rows):
    return bool(rows) and all(len(r) >= 4 and all(num(v) is not None for v in r[1:4]) for r in rows)


def tablolar(js):
    t = js.get('tablolar') or {}
    P = {k: v for k, v, _ in PARAM}
    for k, v in (t.get('param') or {}).items():
        if v is None and P.get(k) is not None: continue
        P[k] = v
    bcu = [tuple(x) for x in t.get('bara_cu') or []]; bal = [tuple(x) for x in t.get('bara_al') or []]
    return dict(marka=[tuple(x) for x in t.get('marka') or MARKA], tips=[tuple(x) for x in t.get('tips') or TIPS],
                korpus_norm=[tuple(x) for x in t.get('korpus_norm') or KORPUS_NORM], param=P,
                param_not={k: n for k, _, n in PARAM}, kablo=[tuple(x) for x in t.get('kablo') or KABLO_TABLO],
                klemens=[tuple(x) for x in t.get('klemens') or KLEMENS_TABLO], sarf=[tuple(x) for x in t.get('sarf') or SARF_TABLO],
                sarf_tek={**SARF_TEK, **{k: v for k, v in (t.get('sarf_tek') or {}).items() if v}}, sni=[tuple(x) for x in t.get('sni') or SNI],
                bara_cu=bcu if _bara_ok(bcu) else BARA_CU, bara_al=bal if _bara_ok(bal) else BARA_AL,
                buk=[tuple(x) for x in t.get('buk') or BUK_TABLO], icerik=[tuple(x) for x in t.get('icerik') or ICERIK],
                kullanici=[tuple(x) for x in t.get('kullanici') or []], ek_kural=list(t.get('ek_kural') or []),
                yas=t.get('yas') or {'YAS_ESKI_AY': 12, 'YAS_YASAK_AY': 18},
                sira_mm=[tuple(x) for x in t.get('sira_mm') or SIRA_MM])


def bara_sec(IN, BT, P):
    """python aynası (Excel formülleri KURALLAR › PANO Z–AD): (şerit n, faz artikülü, PE şerit, PE artikülü, aşım)."""
    d = (IN or 0) / (num(P.get('BARA_DERATING')) or 1); p3 = num(P.get('BARA_PAR3')) or 2.25
    a1 = [num(x[2]) for x in BT]; a2 = [num(x[3]) for x in BT]; a3 = [v * p3 for v in a1]; mm = [num(x[1]) for x in BT]
    asim = False
    if d <= max(a1): n, col = 1, a1
    elif d <= max(a2): n, col = 2, a2
    elif d <= max(a3): n, col = 3, a3
    else: n, col, asim = 3, a3, True
    i = min(len(BT), sum(1 for v in col if v < d) + 1) - 1
    tgt = n * mm[i] / 2; npe = 2 if tgt > max(mm) else 1
    j = min(len(BT), sum(1 for v in mm if v < tgt / npe) + 1) - 1
    return n, str(BT[i][0]), npe, str(BT[j][0]), asim


def _int(x):
    try:
        return int(round(float(str(x).replace(',', '.'))))
    except Exception:
        return None


def _arts(m):
    a = m.get('artikul')
    if isinstance(a, list): return [str(x).strip().rstrip('_') for x in a if str(x).strip()]
    return [x.strip().rstrip('_') for x in str(a or '').split(',') if x.strip()]


def mutabakat_dogrula(js, PN, HATA, kt):
    """Talimat §5.1 / §7: şema · cetvel · BOM mutabakatı ve bağımsız ikinci sayım. Betik şemayı okumaz; beyanın tutarlılığını zorlar."""
    MUT = js.get('mutabakat') or []
    adlar = {p['ad'] for p in PN}
    for m in MUT:
        if m.get('pano') not in adlar: HATA.append(f"mutabakat: pano '{m.get('pano')}' panolar listesinde yok")
    for p in PN:
        ad = p['ad']
        if p.get('mutabakat_muaf'):
            kt(ad, '', 'KAYNAK: mutabakat muaf (dönüştürülmüş / eski dosya) — şema sayımı ve ikinci sayım yapılmadı', 'muaf'); continue
        ek = lambda s: s.get('kaynak') == 'EK' or 'projede yok' in str(s.get('not') or '')
        for s in p['satirlar']:
            if s.get('kaynak') == 'EK' and 'projede yok' not in str(s.get('not') or ''):
                HATA.append(f"{ad} · {s['artikul']}: kaynak EK ama NOT'ta 'projede yok — <neden> için eklendi' gerekçesi yok")
        cihaz = [s for s in p['satirlar'] if s['tip'] != 'korpus' and not ek(s)]
        if not cihaz: continue
        mg = [m for m in MUT if m.get('pano') == ad]
        if not mg:
            HATA.append(f"{ad}: mutabakat yok — §5.1 ŞEMA · CETVEL · BOM, COUNT_A / COUNT_B ve lokatör zorunlu"); continue
        kap = {}
        for m in mg:
            g = m.get('grup') or '?'; arts = _arts(m); aj = ','.join(arts)
            sema, bom, ca, cb = (_int(m.get(k)) for k in ('sema', 'bom', 'count_a', 'count_b'))
            elle = str(m.get('kaynak') or '').upper() == 'ELLE'
            lok = str(m.get('lokator') or '').strip()
            if not arts: HATA.append(f'{ad} · {g}: mutabakat artikül listesi boş')
            for a in arts:
                if a in kap: HATA.append(f'{ad} · {a}: iki mutabakat grubunda ({kap[a]} / {g})')
                kap[a] = g
            if not lok: HATA.append(f'{ad} · {g}: lokatör boş (şema · bara bölümü · hat/cetvel kodu)')
            if sema is None or bom is None:
                HATA.append(f'{ad} · {g}: sema / bom sayı değil'); continue
            if not elle and (ca is None or cb is None):
                HATA.append(f'{ad} · {g}: count_a ve count_b zorunlu (bağımsız ikinci sayım)'); continue
            if ca is not None and cb is not None:
                if sema != max(ca, cb): HATA.append(f'{ad} · {g}: sema {sema} ≠ max(count_a {ca}, count_b {cb})')
                if ca != cb:
                    kt(ad, aj, f"SAYIM: {g} COUNT_A {ca} / COUNT_B {cb} — iki sayım da çizili sembol, büyük olan ({max(ca, cb)}) kullanıldı · {lok}"
                       + (f" · {m['not']}" if m.get('not') else ''))
            toplam = sum(_int(s.get('adet')) or 0 for s in cihaz if s['artikul'] in arts)
            if toplam != bom:
                if elle: kt(ad, aj, f'SAYIM: {g} satır toplamı {toplam} ≠ mutabakat BOM {bom} (elle değişiklik) · {lok}')
                else: HATA.append(f'{ad} · {g}: satır adetleri toplamı {toplam} ≠ mutabakat bom {bom}')
            if sema != bom:
                if elle: kt(ad, aj, f'SAYIM: {g} şema {sema} / BOM {bom} — BOM elle değiştirildi (kullanıcı kararı) · {lok}')
                else: HATA.append(f'{ad} · {g}: şema {sema} ≠ BOM {bom} — §5.1 sayım düzeltilmeli')
            cet = _int(m.get('cetvel'))
            if cet is not None and cet != sema:
                kt(ad, aj, f"SAYIM: {g} şema {sema} / cetvel {cet} / BOM {bom} — {m.get('not') or 'şema esas alındı'} · {lok}")
        for s in cihaz:
            if s['artikul'] not in kap:
                HATA.append(f"{ad} · {s['artikul']} ({s['tip']}, {s.get('adet')} ad.): hiçbir mutabakat grubunda yok")


# =====================================================================================
def kur(js, out):
    P = js['proje']; FY = js['fiyat']; PN = js['panolar']; KT = []
    HATA = []; _gor = set()

    def kt(pano, art, konu, anahtar=None, durum='KONTROL'):
        pano = pano or ''; art = str(art or '')
        k1 = (pano, art, konu); k2 = (pano, art, anahtar) if anahtar else None
        if k1 in _gor or (k2 and k2 in _gor): return
        _gor.add(k1)
        if k2: _gor.add(k2)
        KT.append(dict(pano=pano, artikul=art, konu=konu, durum=durum))

    for x in js.get('kontrol', []):
        if not kodlu(x.get('konu')):
            HATA.append(f"kodsuz KONTROL ({x.get('pano') or '-'}): {str(x.get('konu'))[:70]} — {' / '.join(KODLAR)} kodlarından biriyle başlamalı")
        else:
            kt(x.get('pano'), x.get('artikul'), x.get('konu'), durum=x.get('durum') or 'KONTROL')
    TB = tablolar(js); PR_ = TB['param']
    TIPD = {t[0]: t for t in TB['tips']}
    SARFD = {row[0]: row for row in TB['sarf']}
    fmap = {str(f['artikul']).strip().rstrip('_'): f for f in FY}
    for f in FY:
        f['artikul'] = str(f['artikul']).strip().rstrip('_')
    # sarf artikülleri FIYAT'a (varsayılan)
    sarf_kod = {}
    for row in TB['sarf']:
        for i, tp in ((1, 'kablo'), (2, 'yuksuk'), (3, 'pabuc'), (4, 'klemens'), (5, 'klemens'), (6, 'aksesuar')):
            if row[i]: sarf_kod[str(row[i])] = tp
    for k, v in TB['sarf_tek'].items():
        if v: sarf_kod[str(v)] = SARF_TEK_TIP.get(k, 'aksesuar')
    for v, _ in TB['sni']: sarf_kod[str(v)] = 'nbara'
    for x in TB['bara_cu'] + TB['bara_al']: sarf_kod[str(x[0])] = 'bara'
    for p in PN:
        kn = (p.get('pano') or {}).get('kanal')
        if kn: sarf_kod[str(kn)] = 'kanal'
    for kod, tp in sarf_kod.items():
        if kod not in fmap:
            f = dict(artikul=kod, tip=tp, _sarf=True); FY.append(f); fmap[kod] = f
    ids = {str(f.get('id')) for f in FY if f.get('id')}; nid = 1
    for f in FY:
        if not f.get('id'):
            while f'F{nid:03d}' in ids: nid += 1
            f['id'] = f'F{nid:03d}'; ids.add(f['id'])
        if f.get('ad') in (None, '') or f.get('liste') in (None, '') or not f.get('marka'):
            cm_ = f.get('csv_marka') or {v: k for k, v in CSV_MARKA.items()}.get(str(f.get('marka') or '').lower())
            r = csv_bul(cm_, f['artikul']) if (cm_ or f.get('_sarf') or not f.get('ad')) else None
            if r is not None:
                f.setdefault('ad', r.ACIKLAMA); f.setdefault('birim', r.BIRIM); f.setdefault('kaynak', r.KAYNAK)
                if f.get('liste') in (None, ''): f['liste'] = num(r.LISTE)
                if not f.get('marka'): f['marka'] = CSV_MARKA.get(str(r.MARKA).upper(), str(r.MARKA).lower())
                f['_csv_modul'] = num(r.MODUL)
            elif f.get('_sarf'):
                f.setdefault('ad', f"{f['artikul']} (listede yok)"); f.setdefault('marka', 'ekt'); f.setdefault('kaynak', 'ELLE')
        f['birim'] = birim_norm(f.get('birim'))
        elle = {k for k in ('g', 'y', 'd', 'modul', 'kap') if f.get(k) not in (None, '')}
        for k, v in parse_ozellik(f.get('ad'), f.get('tip')).items():
            if f.get(k) in (None, ''): f[k] = v
        eb = ebat_bul(f.get('csv_marka') or '', f['artikul'])
        if eb:
            f['_ebat'] = eb.get('id')
            for k in ('g', 'y', 'd', 'modul', 'kap'):
                if k in elle or eb.get(k) in (None, ''): continue
                if k == 'modul' and f.get('modul') not in (None, '') and abs(float(f['modul']) - eb['modul']) > 0.01:
                    f['_modfark'] = (f['modul'], eb['modul'])
                f[k] = eb[k]
            if not f.get('not') and eb.get('id'): f['not'] = f"EBAT {eb['id']}" + (f" · {eb['seri']}" if eb.get('seri') and eb['id'] != 'K' else '')
    kul = set()
    for p in PN:
        for s in p['satirlar']:
            s['artikul'] = str(s['artikul']).strip().rstrip('_')
            if s['tip'] == 'din':
                HATA.append(f"{p['ad']} · DIN ray eklenmez (korpusla gelir)"); continue
            if s['artikul'] not in fmap:
                HATA.append(f"{p['ad']} · {s['artikul']} FIYAT listesinde yok"); continue
            if s['tip'] not in TIPD:
                HATA.append(f"{p['ad']} · {s['artikul']}: tanımsız tip '{s['tip']}' — geçerli: "
                            + ', '.join(t for t in TIPD if TIPD[t][1] != 'SARF')); continue
            if TIPD[s['tip']][1] == 'SARF':
                kt(p['ad'], s['artikul'], f"TEKNIK: sarf tipi '{s['tip']}' json'a elle yazılmış — betik sarfı kendisi üretir; mükerrer sarf olmadığını teyit et")
            if s.get('lok'):
                s['not'] = f"{s['lok']} · {s['not']}" if s.get('not') else s['lok']
            kul.add(s['artikul'])
    # v2.10: yasak desenler (talimat §6, §11) — HATA
    for a in sorted(kul):
        f = fmap[a]; ad = str(f.get('ad') or ''); mk = str(f.get('marka') or f.get('csv_marka') or '').lower(); tp = f.get('tip')
        neden = None
        if mk == 'chint' and tp == 'mccb' and re.search(r'\bNM1\b', ad): neden = 'CHINT NM1 yasak — NXM kullanılır'
        elif mk == 'schneider' and a.upper().startswith('A9F74'): neden = 'Schneider A9F74 yasak — A9F79 kullanılır'
        elif mk.startswith('abb') and re.search(r'\bBMS', ad, re.I): neden = 'ABB BMS yasak — SH200 / S200 kullanılır'
        elif tp == 'korpus' and mk == 'chint' and re.search(r'\bNEX(?!US)', ad, re.I): neden = 'CHINT NEX plastik korpus yasak'
        elif tp == 'kanal' and re.search(r'лоток|металл', ad, re.I): neden = 'metal lоток yasak — IEK ИМПАКТ plastik perforeli kanal'
        elif re.search(r'ВА88-32', ad) and (num(f.get('in')) or 0) > 125: neden = 'ВА88-32 en çok 125 A'
        if neden: HATA.append(f'{a}: {neden}')
        elif tp == 'korpus' and mk == 'chint' and re.search(r'\bNX30\b', ad):
            kt('', a, 'KORPUS: CHINT NX30 korpus — malzemesi (NEX plastik yasağı kapsamı) teyit edilmeli', 'nx30')
    # otomatik KONTROL: EBAT / modül (aynı artikülde tek modül satırı; EBAT farkı öncelikli)
    if ebat()['yol'] is None:
        kt('', '', 'KORPUS: EBAT dosyası (ebat.csv) bulunamadı — ölçüler yalnız açıklamadan ayrıştırıldı · aranan: ' + ', '.join(EBAT))
    else:
        for y in ebat()['diger']:
            kt('', '', f"KORPUS: iki farklı ebat.csv var — {ebat()['yol']} kullanıldı, {y} yok sayıldı")
        for a in sorted(kul):
            f = fmap[a]
            if f.get('_modfark'):
                kt('', a, f"KORPUS: modül: ayrıştırma {f['_modfark'][0]:g} ≠ EBAT {f['_modfark'][1]:g} ({f.get('_ebat')}) — EBAT kullanıldı", 'modul')
            if not f.get('_ebat') and f.get('tip') in ('oto', 'dif', 'uzo', 'bh', 'mccb', 'rubilnik', 'ats', 'contactor', 'ppn', 'korpus') and not f.get('g'):
                kt('', a, 'KORPUS: EBAT ölçüsü yok — G/Y/D FIYAT S–AA\'ya elle girilmeli', 'ebat_yok')
            if a in ebat()['cakisma']:
                kt('', a, f"KORPUS: EBAT çakışması ({ebat()['cakisma'][a]}) — ilk kayıt kullanıldı, ebat.csv teyit", 'ebat_cak')
    for a in sorted(kul):
        f = fmap[a]
        cm, pm = f.get('_csv_modul'), f.get('modul')
        if cm and pm and abs(cm - pm) > 0.01 and f.get('tip') in ('oto', 'dif', 'uzo', 'bh'):
            kt('', a, f"KORPUS: modül: LISTE {cm:g} ≠ kullanılan {pm:g} — FIYAT U {pm:g} kullanıldı, katalogla teyit", 'modul')
    # v2.10: faz lambası / lamba koruması (talimat §6.4)
    for p in PN:
        tipler = {s['tip'] for s in p['satirlar']}
        if GIRIS_PANO.search(str(p['ad'])) and not tipler & {'lamb', 'indicator'}:
            kt(p['ad'], '', 'TEKNIK: giriş panosunda faz sinyal lambası yok — §6.4 giriş başına L1/L2/L3 lamba + 3P C6 MCB', 'faz_lamba')
        if tipler & {'lamb', 'indicator'} and not any(s['tip'] == 'oto' and s.get('rol') == 'KUMANDA' for s in p['satirlar']):
            kt(p['ad'], '', 'TEKNIK: lamba var ama KUMANDA rolünde MCB yok — kaynak tarafında 1P MCB varsa teyit, yoksa §6.4 C6 MCB', 'lamba_mcb')
    # v2.10: sayım mutabakatı (talimat §5.1, §7)
    mutabakat_dogrula(js, PN, HATA, kt)
    if HATA:
        raise SystemExit('HATA — Excel kurulmadı, json düzeltilmeli:\n  ' + '\n  '.join(HATA[:40]) + (f'\n  … +{len(HATA) - 40}' if len(HATA) > 40 else ''))

    wb = openpyxl.Workbook(); wb.remove(wb.active)
    B = wb.create_sheet('BOM'); FS = wb.create_sheet('FIYAT'); K = wb.create_sheet('KURALLAR'); T = wb.create_sheet('TEKLIF')
    OZ = wb.create_sheet('OZET'); LS = wb.create_sheet('LISTE')
    for ws in (B, FS, K, T, OZ, LS):
        ws.sheet_view.showGridLines = True

    def nm(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)

    def put(ws, c, v, font=F10, fill=None, fmt=None, border=False):
        x = ws[c]; x.value = temiz(v) if isinstance(v, str) else v; x.font = font
        if fill: x.fill = fill
        if fmt: x.number_format = fmt
        if border: x.border = KEN
        return x

    # ---------------------------------------------------------------- LISTE (birlesik.csv tamamı + ham özellikler)
    LH = ['MARKA', 'ARTIKUL', 'ACIKLAMA', 'BIRIM', 'PARA', 'LISTE', 'KAYNAK', 'TARIH', 'SINIF', 'TIP', 'MARKA_KOD',
          'KUTUP', 'IN', 'MODUL', 'PAKET', 'KESIT', 'G', 'Y', 'D', 'KAP', 'MODUL (liste)']
    for i, h in enumerate(LH):
        put(LS, f'{CL(i+1)}1', h, F10, GRI)
    d = df(); LN = len(d) + 1
    for i, r in enumerate(d.itertuples(index=False), start=2):
        sinif = str(r.SINIF) if str(r.SINIF) != 'nan' else ''
        tip = SINIF_TIP.get(sinif, sinif)
        oz = parse_ozellik(r.ACIKLAMA, tip)
        eb = ebat()['mk'].get((str(r.MARKA).upper(), str(r.ARTIKUL).strip()))
        if eb:
            for k in ('g', 'y', 'd', 'modul', 'kap'):
                if eb.get(k) not in (None, ''): oz[k] = eb[k]
        LS.append([r.MARKA, temiz(r.ARTIKUL), temiz(r.ACIKLAMA), birim_norm(r.BIRIM), r.PARA if str(r.PARA) != 'nan' else None, num(r.LISTE),
                   r.KAYNAK if str(r.KAYNAK) != 'nan' else None, r.TARIH if str(r.TARIH) != 'nan' else None, sinif, tip,
                   CSV_MARKA.get(str(r.MARKA).upper(), str(r.MARKA).lower()), oz.get('kutup'), oz.get('in'), oz.get('modul'),
                   oz.get('paket'), oz.get('kesit'), oz.get('g'), oz.get('y'), oz.get('d'), oz.get('kap'), num(r.MODUL)])
    for c, n in zip('BCDFGIJKLMNOPQRST', ['L_ART', 'L_AD', 'L_BIRIM', 'L_LISTE', 'L_KAYNAK', 'L_SINIF', 'L_TIP', 'L_MARKA',
                                          'L_KUTUP', 'L_IN', 'L_MODUL', 'L_PAKET', 'L_KESIT', 'L_G', 'L_Y', 'L_D', 'L_KAP']):
        nm(n, f'LISTE!${c}$2:${c}${LN}')
    LS.freeze_panes = 'C2'; LS.auto_filter.ref = f'A1:U{LN}'
    for c, w in zip('ABCDEFGHIJK', [10, 16, 70, 7, 6, 11, 12, 11, 10, 10, 10]):
        LS.column_dimensions[c].width = w

    # ---------------------------------------------------------------- planlama (python aynası; formüller Excel'de)
    fattr = lambda a, k, dflt=None: fmap[str(a)].get(k, dflt)
    grup = lambda tip: TIPD.get(tip, (tip, 'KUMANDA'))[1]

    def bk_py(a, tip):
        g = grup(tip); IN = fattr(a, 'in') or 0
        if g == 'GUC': return 0 if (not IN or IN > PR_['KABLO_MAX_IN']) else cnt_le(TB['kablo'], IN)
        if g == 'KUMANDA': return PR_['KUMANDA_KESIT']
        if g == 'OLCU': return PR_['OLCU_KESIT']
        return 0

    def tel_py(a, tip, tar=False):
        t = TIPD.get(tip)
        if not t or bk_py(a, tip) == 0: return 0
        if tar: return 1 if tip in ('dif', 'uzo') else 0
        k = fattr(a, 'kutup') or 0
        return t[2] if t[2] is not None else 2 * k + (1 if k in (1, 3) else 0)

    def uc_py(a, tip):
        if fattr(a, 'uc'): return str(fattr(a, 'uc')).upper()
        t = TIPD.get(tip)
        if not t or t[6] is None: return 'YUKSUK'
        return 'PABUC' if (fattr(a, 'in') or 0) >= t[6] else 'YUKSUK'

    def kk_py(a, tip):
        b = bk_py(a, tip)
        if grup(tip) != 'GUC' or b == 0 or b > PR_['KLEMENS_MAX_KESIT']: return 0
        return cnt_le(TB['klemens'], b)

    _bar0 = {str((p.get('pano') or {}).get('bara')).upper() for p in PN if (p.get('pano') or {}).get('bara')}
    _bdef = str(P.get('bara_malzeme') or (next(iter(_bar0)) if len(_bar0) == 1 else 'CU')).upper()
    isnum = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
    SIRA_ORD = {t[0]: (num(t[5]) if num(t[5]) is not None else 10) for t in TB['tips']}
    PLAN = []; r = 2; BHATA = []
    for p in PN:
        pp = p.setdefault('pano', {}); sarf = p.get('sarf', True)
        blk = dict(p=p, h=r); r += 1
        rows = sorted(p['satirlar'], key=lambda x: SIRA_ORD.get(x['tip'], 10))
        blk['dev'] = [(r + i, x) for i, x in enumerate(rows)]; r += len(rows)
        blk['s'] = blk['h'] + 1; blk['e'] = r - 1
        S = []; blk['kes'] = []; blk['yset'] = set(); blk['pset'] = set(); blk['kset'] = []; blk['kanal'] = False; blk['tarak'] = []
        IN_ANA = max([fattr(x['artikul'], 'in') or 0 for x in rows if x.get('rol') == 'GIRIS' and x['tip'] in KORUMA_TIP] or [0])
        blk['in_ana'] = IN_ANA
        # v2.11 bara kararı
        bt = str(pp.get('bara_tip') or '').upper()
        neden = []
        if IN_ANA >= PR_['BARA_ESIK_A']: neden.append(f'IN_ANA {IN_ANA:g} ≥ {PR_["BARA_ESIK_A"]:g}')
        if re.search(r'D[İI]K', str(pp.get('montaj') or '').upper()): neden.append('DİKİLİ')
        if BARA_PANO.search(str(p['ad'])): neden.append('pano adı')
        if any(x.get('rol') == 'CIKIS' and x['tip'] in ('mccb', 'ppn', 'rubilnik') for x in rows): neden.append('MCCB/ППН çıkış')
        barali = (bt == 'FAZ') if bt in ('FAZ', 'SNI') else bool(neden)
        pp['bara_tip'] = 'FAZ' if barali else 'SNI'
        blk['bara_py'] = barali; blk['bara_neden'] = ', '.join(neden) or ('elle ' + bt if bt else '')
        kor_x = next((x for x in rows if x['tip'] == 'korpus'), None)
        if sarf and barali:
            if not kor_x or not num(fattr(kor_x['artikul'], 'g')) or not num(fattr(kor_x['artikul'], 'y')):
                BHATA.append(f"{p['ad']}: faz barası var ama korpus G/Y yok — bara boyu hesaplanamaz (korpus satırı / json fiyat g, y)")
            if IN_ANA == 0:
                kt(p['ad'], '', 'TEKNIK: faz barası var ama IN_ANA 0 (GIRIS koruma cihazı / IN yok) — bara en küçük seçildi, teyit', 'in0')
        # v2.11 tarak: sıra başına ≥ 2 adet oto/dif/uzo
        sg = {}
        for x in rows:
            if x['tip'] in TARAK_TIP and isnum(x.get('sira')):
                sg[x['sira']] = sg.get(x['sira'], 0) + (num(x.get('adet')) or 0)
        tsira = sorted(k for k, v in sg.items() if v >= 2)
        tar_ids = {id(x) for x in rows if x['tip'] in TARAK_TIP and isnum(x.get('sira')) and x['sira'] in tsira and x.get('rol') != 'GIRIS'}
        _tm = sorted({str(fattr(x['artikul'], 'marka') or '').lower() for x in rows if id(x) in tar_ids} - {'chint', ''})
        if sarf and _tm:
            kt(p['ad'], '', f"TEKNIK: tarak CHINT PIN — {', '.join(_tm)} cihazlarla uyumu (pin aralığı) teyit; gerekirse KURALLAR › SARF_TARAK_* değiştir", 'tarak_marka')
        blk['tar_ids'] = tar_ids
        if sarf and any(x['tip'] in TARAK_TIP and not isnum(x.get('sira')) and x.get('rol') != 'YEDEK' for x in rows):
            kt(p['ad'], '', 'TEKNIK: modüler cihazların bir kısmında sıra numarası (sira) yok — o cihazlar için tarak hesaplanmadı', 'tarak_sira')
        kutuplar = [fattr(x['artikul'], 'kutup') or 0 for x in rows if x.get('rol') == 'GIRIS' and x['tip'] in KORUMA_TIP] \
            or [fattr(x['artikul'], 'kutup') or 0 for x in rows if x['tip'] in KORUMA_TIP]
        blk['tarak_key'] = ('SARF_TARAK_3P63' if IN_ANA and IN_ANA <= 63 else 'SARF_TARAK_3P100') if max(kutuplar or [0]) >= 3 else 'SARF_TARAK_1P'
        if sarf:
            dv = [x for x in rows if x.get('rol', 'CIKIS') != 'YEDEK']
            modcik = any(x.get('rol') == 'CIKIS' and x['tip'] in ('oto', 'dif', 'uzo', 'bh') for x in rows)
            if barali:
                S += [('BARA_F', None), ('BARA_N', None), ('BARA_PE', None), ('IZOL', None)]
                if modcik: S += [('SNI_N', None), ('SNI_PE', None)]
            else:
                S += [('SNI_N', None), ('SNI_PE', None)]
            for sv in tsira:
                S.append(('TARAK', sv)); blk['tarak'].append(sv)
            if pp.get('kanal'):
                S.append(('KANAL', None)); blk['kanal'] = True
            kes = sorted({bk_py(x['artikul'], x['tip']) for x in dv if tel_py(x['artikul'], x['tip'], id(x) in tar_ids) > 0} - {0}); blk['kes'] = kes
            for k in kes:
                S.append(('KABLO', k))
            if pp.get('kapi', 1): S.append(('KAPI_KABLO', 6))
            if pp.get('klemens', False):
                kk = sorted({kk_py(x['artikul'], x['tip']) for x in dv if x.get('rol', 'CIKIS') == 'CIKIS'} - {0})
                blk['kset'] = kk
                for k in kk:
                    S.append(('KLEM_GRI', k)); S.append(('KLEM_PE', k))
                    if SARFD.get(k) and SARFD[k][6]: S.append(('ZAGL', k))
                if kk: S.append(('STOPOR', None))
            for k in kes:
                uc = {uc_py(x['artikul'], x['tip']) for x in dv if bk_py(x['artikul'], x['tip']) == k and tel_py(x['artikul'], x['tip'], id(x) in tar_ids) > 0}
                ykod = SARFD.get(k) and SARFD[k][2]
                if 'YUKSUK' in uc and ykod: S.append(('YUKSUK', k)); blk['yset'].add(k)
                if 'PABUC' in uc or ('YUKSUK' in uc and not ykod): S.append(('PABUC', k)); blk['pset'].add(k)
            if pp.get('kapi', 1): S.append(('KAPI_PABUC', 6))
            if kes: S.append(('BAG', None))
        blk['sarf'] = [(r + i, s) for i, s in enumerate(S)]; r += len(S)
        blk['mm'] = r; r += 1
        blk['isc'] = r; r += 1; blk['top'] = r; r += 1
        blk['end'] = blk['top'] - 1
        PLAN.append(blk)
    if BHATA:
        raise SystemExit('HATA — Excel kurulmadı, json düzeltilmeli:\n  ' + '\n  '.join(BHATA))

    # ---------------------------------------------------------------- v2.9: kullanılan artikül kümesi (python aynası)
    kullan = set(kul)

    def _sni(need):
        dl = [d for _, d in TB['sni']]
        if not dl: return None
        return TB['sni'][-1][0] if need > max(dl) else TB['sni'][min(len(dl) - 1, sum(1 for d in dl if d < need))][0]

    def ndelik(x):
        if x.get('rol') != 'CIKIS' or x['tip'] not in ('oto', 'dif', 'uzo', 'bh'): return 0
        k = fattr(x['artikul'], 'kutup') or 0
        return (x.get('adet') or 0) * (1 if (x['tip'] in ('dif', 'uzo') or k in (1, 2, 4)) else 0)

    def pedelik(x):
        if x.get('rol') != 'CIKIS' or grup(x['tip']) != 'GUC' or x['tip'] in ('mccb', 'ppn', 'rubilnik', 'condencator'): return 0
        return x.get('adet') or 0

    tarak_var = False
    for blk in PLAN:
        p = blk['p']; pp = p['pano']; rows_ = [x for _, x in blk['dev']]
        sy = lambda rol: sum((x.get('adet') or 0) for x in rows_ if x.get('rol') == rol)
        BT = TB['bara_al'] if str(pp.get('bara') or _bdef).upper() == 'AL' else TB['bara_cu']
        if blk['bara_py']:
            n_, fa_, npe_, pa_, asim_ = bara_sec(blk['in_ana'], BT, PR_)
            blk['bara_sec'] = (n_, fa_, npe_, pa_)
            if asim_:
                kt(p['ad'], fa_, f"TEKNIK: IN_ANA {blk['in_ana']:g} A bara tablosunu aşıyor (3 paralel en büyük şerit yetmiyor) — daha büyük bara PYS ile eklenmeli", 'bara_asim')
        for _, (kind, k) in blk['sarf']:
            kod = None
            if kind == 'SNI_N': kod = _sni(sum(ndelik(x) for x in rows_) + sy('KUMANDA') + 1)
            elif kind == 'SNI_PE': kod = _sni(sum(pedelik(x) for x in rows_) + 2 + (pp.get('kapi', 1) or 0))
            elif kind in ('BARA_F', 'BARA_N'): kod = blk['bara_sec'][1]
            elif kind == 'BARA_PE': kod = blk['bara_sec'][3]
            elif kind == 'IZOL': kod = TB['sarf_tek'].get('SARF_IZOLATOR')
            elif kind == 'KANAL': kod = pp.get('kanal')
            elif kind == 'STOPOR': kod = TB['sarf_tek'].get('SARF_STOPOR')
            elif kind == 'BAG': kod = TB['sarf_tek'].get('SARF_BAG')
            elif kind == 'TARAK': kod = TB['sarf_tek'].get(blk['tarak_key']); tarak_var = True
            else:
                col = {'KABLO': 1, 'KAPI_KABLO': 1, 'YUKSUK': 2, 'PABUC': 3, 'KAPI_PABUC': 3, 'KLEM_GRI': 4, 'KLEM_PE': 5, 'ZAGL': 6}.get(kind)
                if col and SARFD.get(k): kod = SARFD[k][col]
            if kod: kullan.add(str(kod))
    if tarak_var:
        kt('', '', 'TEKNIK: tarak tipi (CHINT PIN 1P / 3P 63 A / 3P 100 A) ve sıra akımı teyit — KURALLAR › SARF_TARAK_*; '
                   'baralı panoda tarak sırasının baraya bağlantısı (besleme teli / MCB) şemadan teyit', 'tarak_tip')
        if TIPD.get('tarak') and TIPD['tarak'][4] in (None, ''):
            kt('', '', 'TEKNIK: tarak işçilik normu boş — KURALLAR › TİP TABLOSU tarak NORM dk/adet girilmeli', 'tarak_norm')
    FY.sort(key=lambda f: 0 if f['artikul'] in kullan else 1)
    NKUL = sum(1 for f in FY if f['artikul'] in kullan)
    for f in FY[NKUL:]:
        if not f.get('not'): f['not'] = 'YEDEK SARF — bu projede kullanılmıyor (KURALLAR sarf tablosu)'

    # ---------------------------------------------------------------- FIYAT
    FL = len(FY) + 1; FMAX = FL + OTO_SATIR
    FH = ['ID', 'tip', 'Артикул', 'Номенклатура', 'LİSTE FİYATI', 'NET, тг', 'TOPLAM, тг', 'TOPLAM ADET', 'Ед.изм', 'marka', 'PARA',
          'KATSAYI', 'KUR', 'KAYNAK', 'LİSTE TARİHİ', 'YAŞ, ay', 'DURUM', 'NOT',
          'KUTUP', 'IN, A', 'MODÜL/adet', 'PAKET adet', 'KESİT mm²', 'GENİŞLİK mm', 'YÜKSEKLİK mm', 'DERİNLİK mm', 'KAPASİTE modül',
          'GRUP', 'BAĞLANTI KESİTİ', 'TEL/adet', 'TEL m/adet', 'KLEMENS KESİTİ', 'L+N KLEMENS/adet', 'PE KLEMENS/adet', 'IN sayı', 'KORUMA IN',
          'UÇ TİPİ', 'LİSTE satırı', 'BOM satırı']
    for i, h in enumerate(FH):
        put(FS, f'{CL(i+1)}1', h, F10, SARI if i < 18 else GRI, border=True)
    OZC = ['S', 'T', 'U', 'V', 'W', 'X', 'Y', 'Z', 'AA']; OZK = ['kutup', 'in', 'modul', 'paket', 'kesit', 'g', 'y', 'd', 'kap']
    LNM = ['L_KUTUP', 'L_IN', 'L_MODUL', 'L_PAKET', 'L_KESIT', 'L_G', 'L_Y', 'L_D', 'L_KAP']

    def turet(r, bos):
        mk = f'MATCH($J{r},MARKA_AD,0)'; tp = f'MATCH($B{r},TIP_AD,0)'; g = (lambda x: f'IF($C{r}="","",{x})') if bos else (lambda x: x)
        return {
            'F': '=' + g(f'IF(OR($E{r}="",$L{r}="",$M{r}="",$Q{r}="YASAK"),"",$E{r}*$L{r}*$M{r})'),
            'G': '=' + g(f'IF($F{r}="","",$F{r}*$H{r})'),
            'H': '=' + g(f'SUMIFS(BOM!$AL:$AL,BOM!$C:$C,$C{r})'),
            'K': '=' + g(f'IF(ISNA({mk}),"",INDEX(MARKA_PARA,{mk}))'), 'L': '=' + g(f'IF(ISNA({mk}),"",INDEX(MARKA_KAT,{mk}))'),
            'M': '=' + g(f'IF(OR($K{r}="",ISNA(MATCH($K{r},KUR_AD,0))),"",IF(INDEX(KUR_DEG,MATCH($K{r},KUR_AD,0))="","",INDEX(KUR_DEG,MATCH($K{r},KUR_AD,0))))'),
            'O': '=' + g(f'IF(ISNA(MATCH($N{r},KAYNAK_KOD,0)),"",IF(INDEX(KAYNAK_TARIH,MATCH($N{r},KAYNAK_KOD,0))="","",INDEX(KAYNAK_TARIH,MATCH($N{r},KAYNAK_KOD,0))))'),
            'P': '=' + g(f'IF($O{r}="","",ROUND((TARIH-$O{r})/30.4375,1))'),
            'Q': '=' + g(f'IF(ISNA({mk}),"MARKA YOK",IF(INDEX(MARKA_IZIN,{mk})="HAYIR","YASAK",IF($E{r}="","FİYAT YOK",IF($P{r}="","TARİH YOK",IF($P{r}>YAS_YASAK_AY,"YASAK",IF($P{r}>YAS_ESKI_AY,"ESKİ","OK"))))))'),
            'AB': '=' + g(f'IF(ISNA({tp}),"",INDEX(TIP_GRUP,{tp}))'),
            'AC': '=' + g(f'IF($AB{r}="GUC",IF(OR($AI{r}=0,$AI{r}>KABLO_MAX_IN),0,INDEX(KABLO_KESIT,COUNTIF(KABLO_IN_ALT,"<="&$AI{r}))),'
                          f'IF($AB{r}="KUMANDA",KUMANDA_KESIT,IF($AB{r}="OLCU",OLCU_KESIT,0)))'),
            'AD': '=' + g(f'IF(OR(N($AC{r})=0,ISNA({tp})),0,IF(INDEX(TIP_TEL,{tp})="",2*N($S{r})+IF(OR(N($S{r})=1,N($S{r})=3),1,0),INDEX(TIP_TEL,{tp})))'),
            'AE': '=' + g(f'IF(OR(N($AD{r})=0,ISNA({tp})),0,$AD{r}*INDEX(TIP_TELM,{tp}))'),
            'AF': '=' + g(f'IF(OR($AB{r}<>"GUC",N($AC{r})=0,N($AC{r})>KLEMENS_MAX_KESIT),0,INDEX(KLEMENS_KES,COUNTIF(KLEMENS_ALT,"<="&$AC{r})))'),
            'AG': '=' + g(f'IF(N($AF{r})=0,0,N($S{r})+IF(OR(N($S{r})=1,N($S{r})=3),1,0))'),
            'AH': '=' + g(f'IF(N($AF{r})=0,0,1)'), 'AI': '=' + g(f'N($T{r})'),
            'AJ': '=' + g(f'IF(OR($B{r}="mccb",$B{r}="oto",$B{r}="bh",$B{r}="rubilnik",$B{r}="ats",$B{r}="dif",$B{r}="uzo",$B{r}="ppn"),N($T{r}),0)'),
            'AK': '=' + g(f'IF(ISNA({tp}),"YUKSUK",IF(INDEX(TIP_PABUC,{tp})="","YUKSUK",IF(N($T{r})>=INDEX(TIP_PABUC,{tp}),"PABUC","YUKSUK")))'),
            'AL': '=' + g(f'IF(ISNA(MATCH($C{r},L_ART,0)),"LİSTEDE YOK",MATCH($C{r},L_ART,0))'),
            'AM': '=' + g(f'IF(ISNA(MATCH($C{r},BOM!$C:$C,0)),"",MATCH($C{r},BOM!$C:$C,0))')}

    SARIC = ('E',) + tuple(OZC) + ('AK',)
    for i, f in enumerate(FY):
        r = i + 2
        v = {'A': f['id'], 'B': f.get('tip'), 'C': f['artikul'], 'D': f.get('ad'), 'E': f.get('liste'), 'I': f.get('birim'),
             'J': f.get('marka'), 'N': f.get('kaynak'), 'R': f.get('not')}
        v.update({c: f.get(k) for c, k in zip(OZC, OZK)})
        v.update(turet(r, False))
        if f.get('uc'): v['AK'] = str(f['uc']).upper()
        for c, val in v.items():
            if val is None or val == '': continue
            put(FS, f'{c}{r}', val, F10n, SARI if c in SARIC else None, ACC if c in ('E', 'F', 'G') else ('dd.mm.yyyy' if c == 'O' else None), border=True)
        for c in SARIC:
            FS[f'{c}{r}'].fill = SARI; FS[f'{c}{r}'].border = KEN
    # otomatik satırlar (BOM'a yazılan, FIYAT'ta olmayan artikül)
    for k in range(1, OTO_SATIR + 1):
        r = FL + k
        LX = lambda n: f'=IF(OR($C{r}="",NOT(ISNUMBER($AL{r}))),"",IF(INDEX({n},$AL{r})="","",INDEX({n},$AL{r})))'
        v = {'A': f'Y{k:02d}', 'B': LX('L_TIP'), 'C': f'=IFERROR(INDEX(BOM!$C:$C,MATCH({k},BOM!$S:$S,0)),"")',
             'D': f'=IF($C{r}="","",IF(ISNUMBER($AL{r}),INDEX(L_AD,$AL{r}),"LİSTEDE YOK — fiyatı E sütununa elle girin"))',
             'E': LX('L_LISTE'), 'I': LX('L_BIRIM'), 'J': LX('L_MARKA'), 'N': LX('L_KAYNAK'), 'R': f'=IF($C{r}="","","YENİ — BOM\'dan otomatik")'}
        v.update({c: LX(n) for c, n in zip(OZC, LNM)})
        v.update(turet(r, True))
        for c, val in v.items():
            put(FS, f'{c}{r}', val, F10n, SARI if c in SARIC else None, ACC if c in ('E', 'F', 'G') else ('dd.mm.yyyy' if c == 'O' else None), border=True)
    FA_ANA_END = FL
    for c, n in [('A', 'FIYAT_ID'), ('C', 'FA'), ('D', 'FIYAT_AD'), ('E', 'FIYAT_LISTE'), ('F', 'FIYAT_NET'), ('I', 'FIYAT_BIRIM'), ('J', 'FIYAT_MARKA'),
                 ('Q', 'FIYAT_DURUM'), ('S', 'F_KUTUP'), ('U', 'F_MODUL'), ('V', 'F_PAKET'), ('X', 'F_G'), ('Y', 'F_Y'), ('Z', 'F_D'), ('AA', 'F_KAP'),
                 ('AC', 'F_BK'), ('AD', 'F_TEL'), ('AE', 'F_TELM'), ('AF', 'F_KK'), ('AG', 'F_KLN'), ('AH', 'F_KPE'), ('AI', 'F_IN'),
                 ('AJ', 'F_INKOR'), ('AK', 'F_UC')]:
        nm(n, f'FIYAT!${c}$2:${c}${FMAX}')
    nm('FA_ANA', f'FIYAT!$C$2:$C${FA_ANA_END}')
    put(FS, f'F{FMAX+1}', 'TOPLAM', F10); put(FS, f'G{FMAX+1}', f'=SUM(G2:G{FMAX})', F10, None, ACC)
    FS.freeze_panes = 'D2'; FS.auto_filter.ref = f'A1:AM{FMAX}'
    for i, w in enumerate([6, 10, 16, 58, 11, 11, 13, 8, 7, 10, 6, 8, 6, 12, 11, 7, 10, 28] + [8] * 21):
        FS.column_dimensions[CL(i+1)].width = w
    FS.column_dimensions.group('S', 'AM', hidden=True, outline_level=1)
    if NKUL + 1 < FL:
        FS.sheet_properties.outlinePr.summaryBelow = False
        FS.row_dimensions.group(NKUL + 2, FL, hidden=True, outline_level=1)
        for r_ in range(NKUL + 2, FL + 1):
            for c_ in ('A', 'B', 'C', 'D'):
                FS[f'{c_}{r_}'].fill = GRI
        put(FS, f'AN{NKUL + 1}', f'altında {FL - NKUL - 1} yedek sarf artikülü (gruplu, satır "+" ile açılır)', F10n)

    # ---------------------------------------------------------------- KURALLAR
    kr = [1]

    def row(vals, fills=(), bold=False, fill=None):
        rr = kr[0]
        for i, val in enumerate(vals):
            if val is None: continue
            if isinstance(val, tuple) and val[0] == 'ARR': val = ArrayFormula(f'{CL(i+1)}{rr}', val[1])
            put(K, f'{CL(i+1)}{rr}', val, F10 if bold else F10n, SARI if i in fills else fill, border=True)
        for i in fills:
            K[f'{CL(i+1)}{rr}'].fill = SARI; K[f'{CL(i+1)}{rr}'].border = KEN
        kr[0] += 1; return rr

    def baslik(t): kr[0] += 1; row([t], bold=True, fill=GRI)

    def tablo(heads, rows, names, fills=(), bos=0, key=None):
        row(heads, bold=True); r0 = kr[0]
        for x in rows: row(list(x), fills=fills)
        for _ in range(bos): row([None] * len(heads), fills=fills)
        for ci, n in names.items(): nm(n, f'KURALLAR!${CL(ci+1)}${r0}:${CL(ci+1)}${kr[0]-1}')
        if key: nm(key, f'KURALLAR!$A${r0-1}:${CL(len(heads))}${kr[0]-1}')
        return r0

    def tek(key, lab, val, fmt=None):
        rr = row([key, lab, val], fills=(2,)); nm(key, f'KURALLAR!$C${rr}')
        if fmt: K[f'C{rr}'].number_format = fmt
        return rr

    def bara_tablo(heads, rows, pre):
        row(heads, bold=True); r0 = kr[0]
        for x in list(rows) + [(None, None, None, None)] * 3:
            rr = kr[0]
            row([x[0], x[1], x[2], x[3], f'=IF(C{rr}="","",C{rr}*BARA_PAR3)'], fills=(0, 1, 2, 3))
        for ci, n in enumerate(('KOD', 'MM2', 'A', 'A2', 'A3')):
            nm(f'{pre}_{n}', f'KURALLAR!${CL(ci+1)}${r0}:${CL(ci+1)}${kr[0]-1}')

    # teklif notu varsayılanları: şalt markası BOM güç cihazlarından, bara malzemesi projeden/panolardan
    from collections import Counter
    _say = Counter()
    for p in PN:
        for x in p['satirlar']:
            if x['tip'] != 'korpus' and grup(x['tip']) == 'GUC':
                mk = str(fmap[x['artikul']].get('marka') or '').lower()
                if mk and mk != 'pys': _say[MARKA_GOSTER.get(mk, mk.upper())] += (x.get('adet') or 0) * (p.get('adet') or 1)
    salt_def = P.get('salt_marka') or ', '.join(k for k, _ in _say.most_common())
    _bar = {str((p.get('pano') or {}).get('bara')).upper() for p in PN if (p.get('pano') or {}).get('bara')}
    bara_def = str(P.get('bara_malzeme') or (next(iter(_bar)) if len(_bar) == 1 else 'CU')).upper()
    if not P.get('bara_malzeme') and any(b.get('bara_py') for b in PLAN):
        kt('', '', f'TEKNIK: TEKLIF notu bara malzemesi girilmedi — {bara_def} yazıldı' + (f' (panolarda karışık: {sorted(_bar)})' if len(_bar) > 1 else '') + ' · KURALLAR › BARA_MALZEME')
    _fark = [p['ad'] for p in PN if (p.get('pano') or {}).get('bara') and str(p['pano']['bara']).upper() != bara_def]
    if _fark:
        kt(', '.join(_fark[:5]) + (' …' if len(_fark) > 5 else ''), '',
           f'TEKNIK: pano barası TEKLIF notundaki {bara_def} ile farklı — KURALLAR › PANO I veya BARA_MALZEME teyit')
    row(['KURALLAR — sarı hücreler elle değiştirilir; BOM, FIYAT, TEKLIF ve OZET buna göre hesaplanır'], bold=True)
    baslik('PROJE / TEKLİF')
    tarih = datetime.date.fromisoformat(str(P.get('tarih') or datetime.date.today().isoformat())[:10])
    for key, lab, val in [('KOD', 'Proje kodu', P.get('kod')), ('TEKLIF_NO', '№ КП', P.get('teklif_no')), ('TARIH', 'Дата', tarih),
                          ('MUSTERI', 'Заказчик', P.get('musteri')), ('PROJE', 'Объект', P.get('proje')), ('SEHIR', 'Город', P.get('sehir')),
                          ('FIRMA', 'Поставщик', P.get('firma', 'TEKLED electric')), ('URETICI', 'pano üreticisi (BOM J)', P.get('uretici', 'TEKLED electric')),
                          ('PANO_EK', 'pano adı eki', P.get('pano_ek')), ('GECERLILIK', 'Срок действия, дн.', P.get('gecerlilik')),
                          ('TESLIM', 'Срок изготовления', P.get('teslim')), ('ODEME', 'Условия оплаты', P.get('odeme')),
                          ('KAR_ORANI', 'Kâr oranı', P.get('kar_orani')), ('ISKONTO', 'İskonto', P.get('iskonto', 0)), ('NDS_ORANI', 'НДС', P.get('nds', 0.16)),
                          ('FIRMA_ADRES', 'Адрес поставщика', P.get('firma_adres')), ('FIRMA_TEL', 'Телефон поставщика', P.get('firma_tel')),
                          ('FIRMA_EMAIL', 'Email поставщика', P.get('firma_email', 'yusuf.kaplan@tekled.com')), ('FIRMA_BIN', 'БИН поставщика', P.get('firma_bin')),
                          ('TESLIM_YERI', 'Место поставки', P.get('teslim_yeri')), ('GARANTI', 'Гарантия', P.get('garanti')),
                          ('IMZALAYAN', 'Подписант (ФИО)', P.get('imzalayan')), ('PROJE_NOT', 'Примечание к КП', P.get('not')),
                          ('FIRMA_TAM', 'Юр. наименование поставщика (КП, правый верх)', P.get('firma_tam')),
                          ('MUSTERI_ADRES', 'Адрес заказчика', P.get('musteri_adres')), ('MUSTERI_KISI', 'Контактное лицо', P.get('musteri_kisi')),
                          ('MUSTERI_TEL', 'Рабочий тел. заказчика', P.get('musteri_tel')), ('MUSTERI_MOBIL', 'Мобильный заказчика', P.get('musteri_mobil')),
                          ('MUSTERI_EMAIL', 'Эл. почта заказчика', P.get('musteri_email')),
                          ('BARA_MALZEME', 'Материал шин CU/AL (примечание КП + pano barası varsayılanı)', bara_def),
                          ('SALT_MARKA', 'Марка коммутационной аппаратуры (примечание КП)', salt_def),
                          ('SERTIFIKAT', 'Примечание: сертификаты', P.get('sertifikat', SERTIFIKAT_VARS)),
                          ('KAPANIS', 'Заключительная строка КП', P.get('kapanis')), ('IMZA_UNVAN', 'Должность подписанта', P.get('imza_unvan')),
                          ('HAZIRLAYAN', 'Исполнитель', P.get('hazirlayan')), ('BANKA', 'Название банка', P.get('banka')),
                          ('IIK', 'ИИК', P.get('iik')), ('BIK', 'БИК', P.get('bik'))]:
        tek(key, lab, val, 'dd.mm.yyyy' if key == 'TARIH' else ('0.0%' if key in ('KAR_ORANI', 'ISKONTO', 'NDS_ORANI') else None))
    baslik('KUR (1 birim = KZT)')
    tablo(['PARA', 'KZT'], [(k, 1 if k == 'KZT' else (js.get('kur') or {}).get(k)) for k in KUR], {0: 'KUR_AD', 1: 'KUR_DEG'}, fills=(1,))
    baslik('MARKA (FIYAT J değeri)')
    tablo(['MARKA', 'PARA', 'KATSAYI', 'FİYAT İZNİ', 'NOT'], TB['marka'], {0: 'MARKA_AD', 1: 'MARKA_PARA', 2: 'MARKA_KAT', 3: 'MARKA_IZIN', 4: 'MARKA_NOT'},
          fills=(0, 1, 2, 3), bos=5)
    baslik('FİYAT LİSTESİ KAYNAKLARI')
    kays = js.get('kaynaklar')
    if not kays and os.path.exists(KAYNAKLAR):
        kays = [dict(kod=x['KOD'], dosya=x['DOSYA'], tarih=x['TARIH']) for x in csv.DictReader(open(KAYNAKLAR, encoding='utf-8-sig'))]
    kays = [x for x in (kays or []) if x.get('kod') not in ('PYS', 'ELLE')] + [dict(kod='PYS', dosya='piyasa / elle', tarih=tarih.isoformat()),
                                                                              dict(kod='ELLE', dosya='elle girilen fiyat', tarih=tarih.isoformat())]
    if ebat()['yol']:
        _et = datetime.date.fromtimestamp(os.path.getmtime(ebat()['yol'])).isoformat()
        kays = [x for x in kays if x.get('kod') != 'EBAT'] + [dict(kod='EBAT', dosya=f"ebat.csv ({len(ebat()['mk'])} artikül ölçüsü)", tarih=_et)]
    krows = []
    for x in kays:
        tv = str(x.get('tarih') or '')[:10]; tv = datetime.date.fromisoformat(tv) if re.match(r'\d{4}-\d\d-\d\d', tv) else None
        krows.append((x['kod'], x.get('dosya'), tv))
    r0 = tablo(['KOD', 'DOSYA', 'TARİH'], krows, {0: 'KAYNAK_KOD', 1: 'KAYNAK_DOSYA', 2: 'KAYNAK_TARIH'}, fills=(0, 1, 2), bos=3)
    for rr in range(r0, kr[0]): K[f'C{rr}'].number_format = 'dd.mm.yyyy'
    for k2 in ('YAS_ESKI_AY', 'YAS_YASAK_AY'):
        tek(k2, 'ay', TB['yas'].get(k2))
    baslik('TİP TABLOSU · norm: 260901 Samgau 113 pano, korpus/lamba/buton 260927 · tel boyları başlangıç değeri — KONTROL, atölye ölçümüyle kalibre et · '
           'PABUÇ EŞİĞİ: IN ≥ eşik → PABUÇ, boş → YÜKSÜK · SIRA KATEGORİSİ: YERLEŞİM içerik kontrolü · YAN BOŞLUK: üretici ısı boşluğu mm')
    tablo(['TİP', 'GRUP', 'TEL/adet (boş=2×kutup+N)', 'TEL m', 'NORM dk/adet', 'SIRA', 'PABUÇ EŞİĞİ A', 'SIRA KATEGORİSİ', 'YAN BOŞLUK mm'],
          [tuple(x) for x in TB['tips']] + [(a, 'KORPUS', 0, 0, dd, 1, None, '', None) for a, dd, _ in TB['korpus_norm']],
          {0: 'TIP_AD', 1: 'TIP_GRUP', 2: 'TIP_TEL', 3: 'TIP_TELM', 4: 'NORM_DK', 5: 'TIP_SIRA', 6: 'TIP_PABUC', 7: 'TIP_KAT', 8: 'TIP_BOSLUK'},
          fills=tuple(range(9)), bos=5)
    baslik('PARAMETRELER (işçilik dahil)')
    for key, val in TB['param'].items():
        tek(key, TB['param_not'].get(key, ''), val)
    baslik('KULLANICI PARAMETRELERİ — formülde INDEX(USER_DEG,MATCH("KOD",USER_KOD,0))')
    tablo(['KOD', 'AÇIKLAMA', 'DEĞER'], [(a, b, c) for a, b, c in TB['kullanici']], {0: 'USER_KOD', 1: 'USER_ACK', 2: 'USER_DEG'}, fills=(0, 1, 2), bos=10)
    baslik('EK KURALLAR (metin)')
    tablo(['KURAL'], [(x,) for x in TB['ek_kural']], {0: 'EK_KURAL'}, fills=(0,), bos=10)
    baslik('KABLO KESİT TABLOSU (ПВ3, pano içi; artan sıra)')
    tablo(['IN alt sınır A', 'KESİT mm²'], TB['kablo'], {0: 'KABLO_IN_ALT', 1: 'KABLO_KESIT'}, fills=(0, 1), bos=3)
    baslik('KLEMENS TABLOSU (artan sıra)')
    tablo(['bağlantı kesiti alt', 'KLEMENS mm²', 'genişlik mm'], TB['klemens'], {0: 'KLEMENS_ALT', 1: 'KLEMENS_KES', 2: 'KLEMENS_GEN'}, fills=(0, 1, 2), bos=3)
    baslik('SARF ARTİKÜL TABLOSU (kesit başına; boş = listede yok)')
    tablo(['KESİT mm²', 'KABLO ПВ3', 'НШВИ (yüksük)', 'JG (pabuç)', 'KLEMENS L/N', 'KLEMENS PE', 'UÇ KAPAĞI'], TB['sarf'],
          {0: 'SARF_KESIT', 1: 'SARF_KABLO', 2: 'SARF_YUKSUK', 3: 'SARF_PABUC', 4: 'SARF_KLN', 5: 'SARF_KPE', 6: 'SARF_ZAGL'}, fills=tuple(range(7)), bos=3)
    for key, val in TB['sarf_tek'].items():
        tek(key, SARF_TEK_AD.get(key, ''), val)
    baslik('N / PE BARASI (ШНИ, delik sayısına göre artan)')
    tablo(['ARTİKÜL', 'DELİK'], TB['sni'], {0: 'SNI_KOD', 1: 'SNI_DELIK'}, fills=(0, 1), bos=3)
    baslik('BARA — BAKIR · kullanıcı taşıma kapasitesi tablosu (260930, 25 °C) · tek şerit akımına göre artan sıra korunur · 3 paralel = tek × BARA_PAR3')
    bara_tablo(['ARTİKÜL', 'mm²', 'A (1 şerit)', 'A (2 paralel)', 'A (3 paralel)'], TB['bara_cu'], 'CU')
    baslik('BARA — ALÜMİNYUM · kullanıcı taşıma kapasitesi tablosu (260930, 25 °C)')
    bara_tablo(['ARTİKÜL', 'mm²', 'A (1 şerit)', 'A (2 paralel)', 'A (3 paralel)'], TB['bara_al'], 'AL')
    baslik('KABLO ALANI / BÜKÜLME TABLOSU (YERLEŞİM) — kablo dış çapı ve cihaz üstü/altı kablo alanı; üretici/kablo katalogundan girilir, boşsa KONTROL')
    tablo(['KESİT mm²', 'KABLO DIŞ ÇAP mm', 'KABLO ALANI mm (üst ve alt, her biri)'], TB['buk'], {0: 'BUK_KESIT', 1: 'BUK_CAP', 2: 'BUK_ALAN'}, fills=(0, 1, 2), bos=3)
    baslik('SIRA İÇERİK TABLOSU (YERLEŞİM) — sıraya konabilecek cihaz kategorileri (virgülle)')
    tablo(['İÇERİK', 'İZİN VERİLEN KATEGORİ', 'AÇIKLAMA'], TB['icerik'], {0: 'ICERIK_AD', 1: 'ICERIK_IZIN'}, fills=(0, 1, 2), bos=3)
    baslik('SIRA YÜKSEKLİĞİ REFERANSI (talimat §6.2) — kategori · tip (boş = genel) · IN alt sınır · mm; SIRA KONTROLÜ plan mm bununla karşılaştırılır')
    tablo(['KATEGORİ', 'TİP (boş=genel)', 'IN alt A', 'SIRA mm'], TB['sira_mm'], {0: 'SIRA_KAT', 1: 'SIRA_TIP', 2: 'SIRA_INALT', 3: 'SIRA_MM'},
          fills=(0, 1, 2, 3), bos=3)

    # ---------------------------------------------------------------- PANO PARAMETRELERİ
    baslik('PANO PARAMETRELERİ (her pano bir satır) · KONTROL ≠ OK ise korpus büyütülür veya YENİLE · BARA TİPİ FAZ/SNI (sarı) değişirse YENİLE')
    PH = ['PANO', 'MONTAJ', 'IP', 'KORPUS KAYNAĞI', 'SIRA ADEDİ', 'MODÜLER SIRA', 'SIRA YÜKSEKLİKLERİ mm', 'TOPLAM SIRA mm', 'BARA (CU/AL)',
          'KAPI', 'KLEMENS', 'IN_ANA A', 'İÇ GENİŞLİK mm', 'İÇ YÜKSEKLİK mm', 'MODÜL İHTİYAÇ', 'MODÜL KAPASİTE', 'KONTROL',
          'TEL/UÇ KAPSAMI FARKI', 'KANAL ARTİKÜL', 'KANAL YATAY adet (boş=sıra+1)', 'KAPI YÖNÜ', 'PERDE adet', 'YERLEŞİM', 'PANO #',
          'BARA TİPİ (FAZ/SNI)', 'FAZ ŞERİT adet', 'BARA BOYU m (faz başına)', 'FAZ/N BARA ARTİKÜL', 'PE ŞERİT adet', 'PE BARA ARTİKÜL', 'BARA KARAR NEDENİ']
    row(PH, bold=True)
    for bi, blk in enumerate(PLAN):
        p = blk['p']; pp = p['pano']; rr = kr[0]; blk['pr'] = rr
        s, e, h, last = blk['s'], blk['e'], blk['h'], blk['end']
        kor = next((rw for rw, x in blk['dev'] if x['tip'] == 'korpus'), None); blk['kor'] = kor
        KX = lambda nme: f'IF(ISNA(MATCH(BOM!$C${kor},FA,0)),0,N(INDEX({nme},MATCH(BOM!$C${kor},FA,0))))' if kor else '0'
        KAD = f'MAX(1,N(BOM!$H${kor}))' if kor else '1'
        BR = lambda c: f'BOM!${c}${s}:${c}${e}'
        SK = lambda k: f'INDEX(SARF_YUKSUK,MATCH({k},SARF_KESIT,0))&""'
        tel_total = f'SUM({BR("AE")})'
        kab = '+'.join(f'SUMIFS({BR("AE")},{BR("AD")},{k})' for k in blk['kes']) or '0'
        ucc = [f'SUMIFS({BR("AE")},{BR("AD")},{k},{BR("AG")},"YUKSUK")*({SK(k)}<>"")' for k in sorted(blk['yset'])]
        ucc += [f'(SUMIFS({BR("AE")},{BR("AD")},{k},{BR("AG")},"PABUC")+SUMIFS({BR("AE")},{BR("AD")},{k},{BR("AG")},"YUKSUK")*({SK(k)}=""))' for k in sorted(blk['pset'])]
        klm = '+'.join(f'SUMIFS({BR("AI")},{BR("AH")},{k})' for k in blk['kset']) or '0'
        fark = (f'=IFERROR(ROUND(({tel_total}-({kab}))+({tel_total}-({"+".join(ucc) or "0"}))'
                f'+IF(K{rr}="EVET",SUM({BR("AI")})-({klm}),0)+IF(AND(S{rr}<>"",{"FALSE()" if blk["kanal"] else "TRUE()"}),1,0),3),-1)')
        plan = pp.get('sira_plan') or [['', int(x)] for x in str(pp.get('sira_yuk') or '').split('+') if x.strip().isdigit()]
        pp['sira_plan'] = plan
        _top = sum(int(num(x[1]) or 0) for x in plan)
        if pp.get('sira_yuk_mm') not in (None, '') and plan and num(pp['sira_yuk_mm']) != _top:
            kt(p['ad'], '', f"KORPUS: sira_yuk_mm {pp['sira_yuk_mm']} ≠ sıra planı toplamı {_top} — plan toplamı kullanıldı", 'siramm')
        # v2.11 bara formülleri (Z–AD)
        dd = f'L{rr}/BARA_DERATING'
        nf = lambda M: f'IF({dd}<=MAX({M}_A),1,IF({dd}<=MAX({M}_A2),2,IF({dd}<=MAX({M}_A3),3,"AŞIM")))'
        nn = f'IF(ISNUMBER(Z{rr}),Z{rr},3)'
        ix = lambda M: (f'MIN(COUNT({M}_A),IF({nn}=1,COUNTIF({M}_A,"<"&{dd}),IF({nn}=2,COUNTIF({M}_A2,"<"&{dd}),COUNTIF({M}_A3,"<"&{dd})))+1)')
        tg = lambda M: f'{nn}*INDEX({M}_MM2,MATCH(AB{rr},{M}_KOD,0))/2'
        pix = lambda M: f'MIN(COUNT({M}_MM2),COUNTIF({M}_MM2,"<"&{tg(M)}/AC{rr})+1)'
        vals = [f'=BOM!$D${h}', pp.get('montaj'), pp.get('ip'), pp.get('korpus_kaynak'), pp.get('sira', len(plan) or None), pp.get('sira_modul'),
                '+'.join(str(x[1]) for x in plan) or pp.get('sira_yuk'), None, pp.get('bara') or '=BARA_MALZEME', pp.get('kapi', 1),
                'EVET' if pp.get('klemens', False) else 'HAYIR', f'=MAX({BR("X")})',
                f'=MAX(0,{KX("F_G")}-KORPUS_IC_PAY_MM)*{KAD}', f'=MAX(0,{KX("F_Y")}-KORPUS_IC_PAY_MM)',
                f'=ROUND(SUM(BOM!$AK${s}:$AK${last}),1)', f'=IF({KX("F_KAP")}>0,{KX("F_KAP")},N(F{rr})*INT(M{rr}/MODUL_MM))',
                (f'=IF(AND(M{rr}=0,P{rr}=0),"KORPUS ÖLÇÜSÜ YOK",IF(O{rr}>P{rr}*(1-KORPUS_YEDEK),"MODÜL YETERSİZ",IF(AND(N{rr}>0,N(H{rr})>N{rr}),"SIRA YÜKSEKLİĞİ AŞIYOR",'
                 f'IF(OR((Y{rr}="FAZ")<>{"TRUE()" if blk.get("bara_py") else "FALSE()"},AND(L{rr}>=BARA_ESIK_A,Y{rr}<>"FAZ"),N(R{rr})<>0),"SARF SATIRLARI GÜNCEL DEĞİL",'
                 f'IF(Z{rr}="AŞIM","BARA TABLOSU AŞILDI",IF(AND(Y{rr}="FAZ",N(AA{rr})=0),"BARA BOYU YOK",'
                 f'IF(AND(I{rr}<>"CU",I{rr}<>"AL"),"BARA CU/AL OLMALI","OK")))))))') if p.get('sarf', True) else
                f'=IF(AND(M{rr}=0,P{rr}=0),"KORPUS ÖLÇÜSÜ YOK",IF(O{rr}>P{rr}*(1-KORPUS_YEDEK),"MODÜL YETERSİZ","OK"))',
                fark if p.get('sarf', True) else 0, pp.get('kanal'), pp.get('kanal_yatay'), pp.get('kapi_yon'), pp.get('perde'), None, bi + 1,
                pp.get('bara_tip'),
                f'=IF(Y{rr}<>"FAZ","",IF(I{rr}="AL",{nf("AL")},{nf("CU")}))',
                f'=IF(Y{rr}<>"FAZ",0,ROUND(M{rr}/1000+{KX("F_Y")}*BARA_DIKEY_ORAN/1000*{KAD},3))',
                f'=IF(Y{rr}<>"FAZ","",IF(I{rr}="AL",INDEX(AL_KOD,{ix("AL")}),INDEX(CU_KOD,{ix("CU")}))&"")',
                f'=IF(Y{rr}<>"FAZ","",IF(I{rr}="AL",IF({tg("AL")}>MAX(AL_MM2),2,1),IF({tg("CU")}>MAX(CU_MM2),2,1)))',
                f'=IF(Y{rr}<>"FAZ","",IF(I{rr}="AL",INDEX(AL_KOD,{pix("AL")}),INDEX(CU_KOD,{pix("CU")}))&"")',
                blk.get('bara_neden')]
        row(vals, fills=(1, 2, 3, 4, 5, 6, 8, 9, 10, 18, 19, 20, 21, 24))
    PR0 = PLAN[0]['pr'] if PLAN else kr[0]; PR1 = kr[0] - 1

    # ---------------------------------------------------------------- SIRA YERLEŞİMİ
    baslik('YERLEŞİM › SIRA KONTROLÜ — cihaz sırası BOM T sütunundadır (sayı = sıra, K = kapı, G = gövde); içerik ve plan mm sarı')
    row(['PANO', 'SIRA №', 'İÇERİK', 'PLAN mm', 'CİHAZ ADEDİ', 'CİHAZ GENİŞLİĞİ mm', 'VERİ EKSİK', 'KANAL PAYI mm', 'GEREKEN GENİŞLİK mm',
         'İÇ GENİŞLİK mm', 'EN YÜKSEK CİHAZ mm', 'KABLO ALANI mm', 'GEREKEN YÜKSEKLİK mm', 'İÇERİĞE AYKIRI', 'KONTROL', 'PANO #', 'REF mm (§6.2)'], bold=True)
    SR0 = kr[0]
    for bi, blk in enumerate(PLAN):
        p = blk['p']; pp = p['pano']; s, e, pr = blk['s'], blk['e'], blk['pr']
        BR = lambda c: f'BOM!${c}${s}:${c}${e}'
        blk['sr'] = (kr[0], kr[0] + len(pp['sira_plan']) - 1)
        for ni, (ic, mm) in enumerate(pp['sira_plan'], start=1):
            rr = kr[0]
            IZ = f'IFERROR(INDEX(ICERIK_IZIN,MATCH($C{rr},ICERIK_AD,0))&"","")'
            ayk = '+'.join(f'COUNTIFS({BR("T")},$B{rr},{BR("Y")},"{c}")*ISERROR(SEARCH("{c}",{IZ}))' for c in KATEGORI)
            row([f'=BOM!$D${blk["h"]}', ni, ic, mm, f'=SUMIFS({BR("H")},{BR("T")},$B{rr})', f'=SUMIFS({BR("AA")},{BR("T")},$B{rr})',
                 f'=SUMIFS({BR("AB")},{BR("T")},$B{rr})',
                 f'=IF(KURALLAR!$S${pr}="",0,2*IFERROR(N(INDEX(F_G,MATCH(KURALLAR!$S${pr},FA,0))),0))', f'=F{rr}+H{rr}', f'=KURALLAR!$M${pr}',
                 ('ARR', f'=MAX(IF({BR("T")}=$B{rr},{BR("V")},0))'), ('ARR', f'=MAX(IF({BR("T")}=$B{rr},{BR("Z")},0))'),
                 f'=K{rr}+2*L{rr}', f'={ayk}',
                 f'=IF(OR($C{rr}="",ISNA(MATCH($C{rr},ICERIK_AD,0))),"İÇERİK TANIMSIZ",IF(N{rr}>0,"İÇERİĞE AYKIRI CİHAZ",IF(AND(E{rr}=0,OR($C{rr}="KLEMENS",$C{rr}="BARA")),"OK",'
                 f'IF(E{rr}=0,"BOŞ SIRA",IF(AND(N(Q{rr})>0,N(D{rr})<N(Q{rr})),"PLAN mm REF ALTINDA",IF(J{rr}=0,"ÖLÇÜ YOK — yalnız modül",IF(G{rr}>0,"BOYUT VERİSİ YOK",IF(I{rr}>J{rr},"GENİŞLİK AŞIYOR",IF(M{rr}>D{rr},"YÜKSEKLİK AŞIYOR","OK")))))))))',
                 bi + 1, ('ARR', f'=MAX(MAX(IF({BR("T")}=$B{rr},{BR("AO")},0)),IF(OR($C{rr}="KLEMENS",$C{rr}="BARA"),SUMPRODUCT(MAX((SIRA_KAT=$C{rr})*SIRA_MM)),0))')],
                fills=(2, 3))
        K[f'H{pr}'].value = f'=SUM(D{blk["sr"][0]}:D{blk["sr"][1]})' if pp['sira_plan'] else 0
        K[f'H{pr}'].border = KEN; K[f'H{pr}'].font = F10n

    # ---------------------------------------------------------------- YERLEŞİM (pano)
    baslik('YERLEŞİM › PANO KONTROLÜ — boyutlar FIYAT X–Z (sarı), paylar PARAMETRELER, kablo alanı BÜKÜLME TABLOSU; boş veri KONTROL')
    row(['PANO', 'KORPUS D mm', 'PLAN TOPLAM mm', 'GEREKEN YÜKSEKLİK mm', 'İÇ YÜKSEKLİK mm', 'KAPI CİHAZI adet', 'EN DERİN KAPI CİHAZI mm',
         'KAPI CİHAZI SINIRI mm', 'EN DERİN SIRA CİHAZI mm', 'CİHAZ DERİNLİK SINIRI mm', 'KANAL DOLULUK', 'ÖLÇÜ BÖLÜMÜ', 'PERDE', 'BARA BÖLGESİ',
         'KLEMENS SIRASI', 'KAPI YÖNÜ', 'SIRASIZ CİHAZ', 'SIRA HATASI', 'EKSİK VERİ / PARAMETRE', 'SONUÇ'], bold=True)
    YR0 = kr[0]
    for blk in PLAN:
        p = blk['p']; pp = p['pano']; s, e, pr = blk['s'], blk['e'], blk['pr']; rr = kr[0]; blk['yr'] = rr
        BR = lambda c: f'BOM!${c}${s}:${c}${e}'
        a, b = blk['sr']; SC = f'$C${a}:$C${b}' if pp['sira_plan'] else f'$C${rr}'; SO = f'$O${a}:$O${b}' if pp['sira_plan'] else None
        kor = blk['kor']
        KD = f'IF(ISNA(MATCH(BOM!$C${kor},FA,0)),0,N(INDEX(F_D,MATCH(BOM!$C${kor},FA,0))))' if kor else '0'
        kk = [k for k in blk['kes'] if k < 35]
        kmt = f'MATCH(KURALLAR!$S${pr},FA,0)'
        if kk:
            dol = '+'.join(f'SUMIFS({BR("AE")},{BR("AD")},{k})*INDEX(BUK_CAP,MATCH({k},BUK_KESIT,0))^2' for k in kk)
            bos = ','.join(f'INDEX(BUK_CAP,MATCH({k},BUK_KESIT,0))&""=""' for k in kk)
            kanal = (f'=IF(KURALLAR!$S${pr}="","GEREKMEZ",IFERROR(IF(OR({bos},N(INDEX(F_G,{kmt}))=0,N(INDEX(F_Y,{kmt}))=0),"VERİ YOK",'
                     f'ROUND(({dol})*PI()/4/(INDEX(F_G,{kmt})*INDEX(F_Y,{kmt})),3)),"VERİ YOK"))')
        else:
            kanal = f'=IF(KURALLAR!$S${pr}="","GEREKMEZ",0)'
        hata = f'COUNTIF({SO},"<>OK")-COUNTIF({SO},"ÖLÇÜ YOK — yalnız modül")' if SO else '1'
        row([f'=BOM!$D${blk["h"]}', f'={KD}', f'=KURALLAR!$H${pr}',
             f'=C{rr}+MAX(0,COUNT(KURALLAR!$B${a}:$B${b})-1)*N(BAKIM_PAY_MM)+N(UST_ALT_PAY_MM)' if pp['sira_plan'] else f'=C{rr}',
             f'=KURALLAR!$N${pr}', f'=SUMIFS({BR("H")},{BR("T")},"K")', ('ARR', f'=MAX(IF({BR("T")}="K",{BR("W")},0))'),
             f'=IF(B{rr}=0,"",B{rr}-N(MONTAJ_PLAKA_PAY_MM)-N(KAPI_ARKA_PAY_MM))', ('ARR', f'=MAX(IF(ISNUMBER({BR("T")}),{BR("W")},0))'),
             f'=IF(B{rr}=0,"",B{rr}-N(MONTAJ_PLAKA_PAY_MM))', kanal,
             f'=IF(SUMIFS({BR("H")},{BR("Y")},"OLCU")=0,"GEREKMEZ",IF(COUNTIF({SC},"OLCU")>0,"VAR","YOK"))',
             f'=IF(L{rr}<>"VAR","GEREKMEZ",IF(N(KURALLAR!$V${pr})>0,"VAR","YOK"))',
             f'=IF(KURALLAR!$Y${pr}="FAZ",IF(COUNTIF({SC},"BARA")>0,"VAR","YOK"),"GEREKMEZ")',
             f'=IF(KURALLAR!$K${pr}="EVET",IF(COUNTIF({SC},"KLEMENS")>0,"VAR","YOK"),"GEREKMEZ")',
             f'=IF(F{rr}=0,"GEREKMEZ",IF(KURALLAR!$U${pr}="","YOK",KURALLAR!$U${pr}))',
             f'=COUNTIFS({BR("T")},"",{BR("N")},"<>KORPUS")', f'={hata}',
             f'=SUMIFS({BR("AB")},{BR("T")},"K")+(BAKIM_PAY_MM="")+(UST_ALT_PAY_MM="")+IF(F{rr}>0,(MONTAJ_PLAKA_PAY_MM="")+(KAPI_ARKA_PAY_MM=""),0)'
             f'+IF(KURALLAR!$S${pr}<>"",(KANAL_DOLULUK_MAX="")+(K{rr}="VERİ YOK"),0)+IF(AND(B{rr}=0,KURALLAR!$M${pr}>0),1,0)',
             f'=IF(Q{rr}>0,"SIRASI ATANMAMIŞ CİHAZ: "&Q{rr},IF(L{rr}="YOK","ÖLÇÜ BÖLÜMÜ YOK",IF(M{rr}="YOK","PERDE YOK",IF(N{rr}="YOK","BARA BÖLGESİ YOK",'
             f'IF(O{rr}="YOK","KLEMENS SIRASI YOK",IF(KURALLAR!$M${pr}=0,"YALNIZ MODÜL — korpus ölçüsü yok",IF(D{rr}>E{rr},"YÜKSEKLİK AŞIYOR",'
             f'IF(AND(ISNUMBER(H{rr}),G{rr}>H{rr}),"KAPI CİHAZI SIĞMIYOR",IF(AND(ISNUMBER(J{rr}),I{rr}>J{rr}),"CİHAZ DERİNLİĞİ AŞIYOR",'
             f'IF(AND(ISNUMBER(K{rr}),N(KANAL_DOLULUK_MAX)>0,N(K{rr})>N(KANAL_DOLULUK_MAX)),"KANAL DOLU",IF(R{rr}>0,"SIRA KONTROLÜ: "&R{rr},'
             f'IF(P{rr}="YOK","KAPI YÖNÜ YOK",IF(S{rr}>0,"VERİ/PARAMETRE EKSİK: "&S{rr},"OK")))))))))))))'])
        K[f'W{pr}'].value = f'=T{rr}'; K[f'W{pr}'].border = KEN; K[f'W{pr}'].font = F10n
    YR1 = kr[0] - 1

    # ---------------------------------------------------------------- v2.10: SAYIM MUTABAKATI
    baslik('SAYIM MUTABAKATI (§5.1) — ŞEMA · CETVEL · BOM · COUNT_A/B · lokatör; BOM CANLI = BOM sayfasındaki güncel adet · YENİLE geri okur')
    row(['PANO', 'GRUP', 'ARTİKÜL', 'ŞEMA', 'CETVEL', 'BOM', 'COUNT_A', 'COUNT_B', 'LOKATÖR', 'KAYNAK', 'BOM CANLI', 'DURUM', 'NOT'], bold=True)
    MR0 = kr[0]; bmap = {}
    for blk in PLAN: bmap.setdefault(blk['p']['ad'], blk)
    for m in js.get('mutabakat') or []:
        blk = bmap.get(m.get('pano'))
        if not blk: continue
        arts = _arts(m); s, e = blk['s'], blk['e']; rr = kr[0]
        canli = '+'.join(f'SUMIFS(BOM!$H${s}:$H${e},BOM!$C${s}:$C${e},"{a}",BOM!$N${s}:$N${e},"<>SARF")' for a in arts) or '0'
        row([m.get('pano'), m.get('grup'), ', '.join(arts), _int(m.get('sema')), m.get('cetvel'), _int(m.get('bom')), _int(m.get('count_a')),
             _int(m.get('count_b')), m.get('lokator'), m.get('kaynak'), f'={canli}',
             f'=IF(K{rr}<>D{rr},"ŞEMA ≠ BOM",IF(AND(ISNUMBER(E{rr}),E{rr}<>D{rr}),"ŞEMA ≠ CETVEL","OK"))', m.get('not')])
    MR1 = kr[0] - 1

    baslik('OTOMATİK KONTROL')
    for lab, f, key in [('Markası KURALLAR\'da olmayan kalem (kullanılan)', '=COUNTIFS(FIYAT!$Q:$Q,"MARKA YOK",FIYAT!$H:$H,">0")', None),
                        ('Fiyatı olmayan / yasak kalem (kullanılan)', '=COUNTIFS(FIYAT!$Q:$Q,"FİYAT YOK",FIYAT!$H:$H,">0")+COUNTIFS(FIYAT!$Q:$Q,"YASAK",FIYAT!$H:$H,">0")+COUNTIFS(FIYAT!$Q:$Q,"MARKA YOK",FIYAT!$H:$H,">0")', 'OK_FIYATSIZ'),
                        ('Eski fiyat (kullanılan)', '=COUNTIFS(FIYAT!$Q:$Q,"ESKİ",FIYAT!$H:$H,">0")', None),
                        ('LİSTEDE YOK (kullanılan)', '=COUNTIFS(FIYAT!$AL:$AL,"LİSTEDE YOK",FIYAT!$H:$H,">0")', None),
                        ('BOM artikülü FIYAT\'ta yok', '=COUNTIF(BOM!$P:$P,"ARTİKÜL YOK")', None),
                        ('Normu boş / tanımsız tip kullanımı', '=COUNTIF(BOM!$AM:$AM,"BOŞ")+COUNTIF(BOM!$AM:$AM,"YOK")', None),
                        ('Kur boş para birimi kullanımı', '=COUNTIFS(FIYAT!$H:$H,">0",FIYAT!$K:$K,"?*",FIYAT!$M:$M,"")', None),
                        ('Pano kontrolü OK olmayan', f'=COUNTIF(Q{PR0}:Q{PR1},"<>OK")', 'OK_PANO'),
                        ('TEL/UÇ KAPSAMI FARKI toplamı', f'=SUM(R{PR0}:R{PR1})', None),
                        ('Yerleşim kontrolü OK olmayan', f'=COUNTIF(T{YR0}:T{YR1},"<>OK")', 'OK_YERLESIM'),
                        ('Otomatik satır kapasitesi', f'=IF(MAX(BOM!$S:$S)>{OTO_SATIR},"DOLU — YENİLE",MAX(BOM!$S:$S)&" / {OTO_SATIR}")', 'OK_OTO'),
                        ('Mutabakat: ŞEMA ≠ BOM (canlı)', f'=COUNTIF(L{MR0}:L{MR1},"ŞEMA ≠ BOM")' if MR1 >= MR0 else 0, 'OK_MUTABAKAT'),
                        ('Kâr oranı', '=IF(KAR_ORANI="","GİRİLMEDİ","OK")', None)]:
        rr = row([lab, None, f])
        if key: nm(key, f'KURALLAR!$C${rr}')
    for _g in ('tekled', 'teksan', 'kz'):
        try:
            import io as _io, base64 as _b64
            from openpyxl.drawing.image import Image as _XI
            _XI(_io.BytesIO(open(LOGO, 'rb').read() if (_g == 'tekled' and os.path.exists(LOGO)) else _b64.b64decode(gorsel_paket()[_g])))
        except Exception as _e:
            kt('TEKLIF', '', f'TEKLIF: görsel eklenemedi: {_g} ({_e})')
    baslik('KONTROL LİSTESİ (mühendis kararı gereken kalemler)')
    row(['PANO', 'ARTİKÜL', 'KONU', 'DURUM'], bold=True); k0 = kr[0]
    for x in KT:
        row([x.get('pano'), x.get('artikul'), x.get('konu'), x.get('durum', 'KONTROL')], fills=(3,))
    for _ in range(10): row([None] * 4, fills=(0, 1, 2, 3))
    nm('KT_DURUM', f'KURALLAR!$D${k0}:$D${kr[0]-1}')
    for i, w in enumerate([34, 30, 16, 16, 14, 14, 18, 14, 14, 12, 12, 12, 14, 14, 14, 14, 26, 14, 16, 12, 12, 10, 26, 8, 12, 10, 12, 14, 10, 14, 34]):
        K.column_dimensions[CL(i+1)].width = w

    # ---------------------------------------------------------------- BOM
    BH = ['Наименование', 'malzeme', 'Артикул', 'Номенклатура', 'Цена,тг', 'Цена,тг', 'СУММ.', 'Кол- во', 'Ед.изм', 'завод из готовитель',
          'Источник', 'malz', 'PANO ADET', 'ROL', 'NORM', 'DURUM', 'NOT', 'FID', 'YENİ', 'SIRA',
          'G mm/adet', 'Y mm', 'D mm', 'KORUMA IN', 'KATEGORİ', 'KABLO ALANI mm', 'SIRA GENİŞLİĞİ mm', 'VERİ EKSİK', 'İŞ ADEDİ', 'KESİT',
          'TEL adet', 'TEL m', 'UÇ', 'KLEMENS KESİT', 'L+N KLEMENS', 'PE KLEMENS', 'MODÜL', 'TOPLAM ADET', 'NORM dk', 'İŞ dk', 'SIRA REF mm',
          'N DELİK (ШНИ)', 'PE DELİK (ШНИ)', 'TARAK']
    for i, h in enumerate(BH):
        put(B, f'{CL(i+1)}1', h, F10, SARI if i < 20 else GRI, border=True)
    MT = lambda rr: f'MATCH($C{rr},FA,0)'
    CBLK = [2, 2]

    def fx(col, rr):
        return f'=IF(ISNA({MT(rr)}),"",IF(INDEX({col},{MT(rr)})="","",INDEX({col},{MT(rr)})))'

    def cell(rr, c, v, fmt=None, font=F10):
        return put(B, f'{c}{rr}', v, font, None, fmt, border=c in ('A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K'))

    def yardimci(rr, sirali):
        m = MT(rr); na = f'ISNA({m})'; tm = f'MATCH($B{rr},TIP_AD,0)'
        bs, be = CBLK
        TB_ = lambda c: f'${c}${bs}:${c}${be}'
        tar_ad = '+'.join(f'SUMIFS({TB_("H")},{TB_("T")},$T{rr},{TB_("B")},"{t}")' for t in TARAK_TIP)
        tipok = 'OR(' + ','.join(f'$B{rr}="{t}"' for t in TARAK_TIP) + ')'
        H = {'AC': f'=N($H{rr})*IF($I{rr}="упак",MAX(1,IF({na},1,N(INDEX(F_PAKET,{m})))),1)',
             'AD': f'=IF({na},0,N(INDEX(F_BK,{m})))',
             'AE': f'=IF(OR({na},$N{rr}="YEDEK",$N{rr}="SARF"),0,IF($AR{rr}=1,IF(N(INDEX(F_BK,{m}))=0,0,N($H{rr})*IF(OR($B{rr}="dif",$B{rr}="uzo"),1,0)),'
                   f'N($H{rr})*N(INDEX(F_TEL,{m}))))',
             'AF': f'=IF(OR({na},$N{rr}="YEDEK",$N{rr}="SARF"),0,IF($AR{rr}=1,$AE{rr}*IF(ISNA({tm}),0,N(INDEX(TIP_TELM,{tm}))),N($H{rr})*N(INDEX(F_TELM,{m}))))',
             'AG': f'=IF({na},"",INDEX(F_UC,{m})&"")',
             'AH': f'=IF(OR({na},$N{rr}<>"CIKIS"),0,N(INDEX(F_KK,{m})))',
             'AI': f'=IF($AH{rr}=0,0,N($H{rr})*N(INDEX(F_KLN,{m})))', 'AJ': f'=IF($AH{rr}=0,0,N($H{rr})*N(INDEX(F_KPE,{m})))',
             'AK': f'=IF({na},0,N($H{rr})*N(INDEX(F_MODUL,{m})))', 'AL': f'=N($H{rr})*N($M{rr})',
             'AM': f'=IF(ISNA(MATCH($O{rr},TIP_AD,0)),"YOK",IF(INDEX(NORM_DK,MATCH($O{rr},TIP_AD,0))="","BOŞ",INDEX(NORM_DK,MATCH($O{rr},TIP_AD,0))))',
             'AN': f'=IF(ISNUMBER($AM{rr}),$AC{rr}*$AM{rr},0)',
             'X': f'=IF({na},0,IF($N{rr}="GIRIS",N(INDEX(F_INKOR,{m})),0))',
             'AP': f'=IF(AND($N{rr}="CIKIS",OR($B{rr}="oto",$B{rr}="dif",$B{rr}="uzo",$B{rr}="bh")),N($H{rr})*IF(OR($B{rr}="dif",$B{rr}="uzo"),1,'
                   f'IF(OR(IF({na},0,N(INDEX(F_KUTUP,{m})))=1,IF({na},0,N(INDEX(F_KUTUP,{m})))=2,IF({na},0,N(INDEX(F_KUTUP,{m})))=4),1,0)),0)',
             'AQ': f'=IF(AND($N{rr}="CIKIS",NOT(ISNA({tm})),$B{rr}<>"mccb",$B{rr}<>"ppn",$B{rr}<>"rubilnik",$B{rr}<>"condencator"),'
                   f'IF(INDEX(TIP_GRUP,{tm})="GUC",N($H{rr}),0),0)',
             'AR': f'=IF(AND(ISNUMBER($T{rr}),{tipok},$N{rr}<>"GIRIS"),IF(({tar_ad})>=2,1,0),0)'}
        if sirali:
            bos = f'IF(ISNA({tm}),0,N(INDEX(TIP_BOSLUK,{tm})))'
            H.update({'U': f'=IF(OR($T{rr}="",$T{rr}="G",{na}),"",IF(N(INDEX(F_G,{m}))>0,INDEX(F_G,{m}),IF(N(INDEX(F_MODUL,{m}))>0,INDEX(F_MODUL,{m})*MODUL_MM,"")))',
                      'V': f'=IF(OR($T{rr}="",$T{rr}="G",{na}),"",IF(N(INDEX(F_Y,{m}))>0,INDEX(F_Y,{m}),""))',
                      'W': f'=IF(OR($T{rr}="",$T{rr}="G",{na}),"",IF(N(INDEX(F_D,{m}))>0,INDEX(F_D,{m}),""))',
                      'Y': f'=IF(ISNA({tm}),"",INDEX(TIP_KAT,{tm})&"")',
                      'Z': f'=IF(NOT(ISNUMBER($T{rr})),"",IF($AD{rr}=0,0,IF(ISNA(MATCH($AD{rr},BUK_KESIT,0)),"",IF(INDEX(BUK_ALAN,MATCH($AD{rr},BUK_KESIT,0))&""="","",INDEX(BUK_ALAN,MATCH($AD{rr},BUK_KESIT,0))))))',
                      'AA': f'=IF(OR(NOT(ISNUMBER($T{rr})),$U{rr}=""),"",N($H{rr})*($U{rr}+{bos}))',
                      'AB': f'=IF(ISNUMBER($T{rr}),($U{rr}="")+($V{rr}="")+($Z{rr}="")+IF(AND($Y{rr}="KONTAKTOR",NOT(ISNA({tm}))),IF(INDEX(TIP_BOSLUK,{tm})&""="",1,0),0),IF($T{rr}="K",IF($W{rr}="",1,0),0))',
                      'AO': f'=IF(OR(NOT(ISNUMBER($T{rr})),{na}),"",SUMPRODUCT(MAX((SIRA_KAT=$Y{rr})*(((COUNTIF(SIRA_TIP,$B{rr})>0)*(SIRA_TIP=$B{rr})'
                            f'+(COUNTIF(SIRA_TIP,$B{rr})=0)*(SIRA_TIP=""))>0)*(SIRA_INALT<=N(INDEX(F_IN,{m})))*SIRA_MM)))'})
        for c, v in H.items():
            put(B, f'{c}{rr}', v, F10n)

    def satir(rr, tip, C, H, rol, norm, notu, fid=None, sira=None):
        cell(rr, 'A', f'=A{rr-1}'); cell(rr, 'B', tip)
        cell(rr, 'C', f'=IFERROR(INDEX(FA,MATCH($R{rr},FIYAT_ID,0)),"FID YOK")' if fid else C)
        cell(rr, 'D', f'=IF(ISNA({MT(rr)}),"ARTİKÜL FIYAT SAYFASINDA YOK",INDEX(FIYAT_AD,{MT(rr)}))')
        cell(rr, 'E', fx('FIYAT_LISTE', rr), ACC); cell(rr, 'F', fx('FIYAT_NET', rr), ACC)
        cell(rr, 'G', f'=IF(OR(F{rr}="",H{rr}=""),"",H{rr}*F{rr})', ACC); cell(rr, 'H', H)
        cell(rr, 'I', fx('FIYAT_BIRIM', rr)); cell(rr, 'J', fx('FIYAT_MARKA', rr)); cell(rr, 'K', f'=K{rr-1}')
        cell(rr, 'L', f'=G{rr}', ACC); cell(rr, 'M', f'=M{rr-1}'); cell(rr, 'N', rol); cell(rr, 'O', norm)
        cell(rr, 'P', f'=IF(ISNA({MT(rr)}),"ARTİKÜL YOK",INDEX(FIYAT_DURUM,{MT(rr)}))', font=F10n)
        if notu: cell(rr, 'Q', notu, font=F10n)
        if fid: cell(rr, 'R', fid, font=F10n)
        cell(rr, 'S', f'=IF(OR($C{rr}="",$C{rr}="FID YOK"),"",IF(ISNUMBER(MATCH($C{rr},FA_ANA,0)),"",IF(COUNTIF($C$1:$C{rr-1},$C{rr})>0,"",MAX($S$1:$S{rr-1})+1)))', font=F10n)
        if sira is not None:
            x = put(B, f'T{rr}', sira, F10n, SARI)
        yardimci(rr, sira is not None or fid is not None)

    TOPLAMLAR = []
    ek = P.get('pano_ek')
    for blk in PLAN:
        p = blk['p']; pp = p['pano']; h, s, e, last, pr = blk['h'], blk['s'], blk['e'], blk['end'], blk['pr']
        CBLK[0], CBLK[1] = s, e
        KP = lambda col: f'KURALLAR!${col}${pr}'
        for c, v in {'A': f'=D{h}', 'D': p['ad'] + (f' ({ek})' if ek else ''), 'H': p.get('adet', 1), 'I': 'комплект',
                     'J': p.get('korpus_marka') or P.get('uretici') or 'TEKLED electric', 'K': p.get('kaynak', P.get('kod')), 'M': f'=H{h}'}.items():
            put(B, f'{c}{h}', v, F10, PANOF, border=c in 'ABCDEFGHIJK')
        for c in 'BCEFG':
            B[f'{c}{h}'].fill = PANOF; B[f'{c}{h}'].border = KEN
        for rr, x in blk['dev']:
            tip = x['tip']; a = x['artikul']
            y = fattr(a, 'y') or 0
            norm = x.get('norm') or (('KORPUS_AYAKLI' if y >= 1500 else 'KORPUS_ORTA' if y >= 700 else 'KORPUS_KUCUK') if tip == 'korpus' else tip)
            sira = x.get('sira')
            if sira is None and tip != 'korpus' and len(TIPD.get(tip, ())) > 7 and TIPD[tip][7] == 'KAPI': sira = 'K'
            if tip == 'korpus': sira = None
            satir(rr, tip, None, x['adet'], x.get('rol') or ('CIKIS' if grup(tip) == 'GUC' else grup(tip)), norm, x.get('not'),
                  fid=fmap[a]['id'], sira=sira if tip != 'korpus' else None)
            if tip == 'korpus':
                for c in ('U', 'V', 'W', 'Y', 'Z', 'AA', 'AB', 'AO'): B[f'{c}{rr}'].value = None
        BR = lambda c: f'${c}${s}:${c}${e}'
        paket = lambda rr: f'MAX(1,IF(ISNA({MT(rr)}),1,N(INDEX(F_PAKET,{MT(rr)}))))'
        ST = lambda col, k: f'=IF(INDEX({col},MATCH({k},SARF_KESIT,0))&""="","",INDEX({col},MATCH({k},SARF_KESIT,0)))'
        kablo_rows, gri_rows = [], []
        needN = f'(SUM({BR("AP")})+SUMIFS({BR("H")},{BR("N")},"KUMANDA")+1)'
        needPE = f'(SUM({BR("AQ")})+2+N({KP("J")}))'
        Wic = f'{KP("M")}/1000'
        NN = f'IF(ISNUMBER({KP("Z")}),{KP("Z")},3)'
        for rr, (kind, k) in blk['sarf']:
            if kind in ('SNI_N', 'SNI_PE'):
                need = needN if kind == 'SNI_N' else needPE
                C = f'=IF({need}>MAX(SNI_DELIK),INDEX(SNI_KOD,COUNT(SNI_DELIK)),INDEX(SNI_KOD,COUNTIF(SNI_DELIK,"<"&{need})+1))'
                H = f'=IF({need}>MAX(SNI_DELIK),ROUNDUP({need}/MAX(SNI_DELIK),0),1)'
                satir(rr, 'nbara', C, H, 'SARF', 'nbara', ('N barası ШНИ (otomatik)' if kind == 'SNI_N' else 'PE barası ШНИ (otomatik)')
                      + (' · modüler çıkışlar' if blk['bara_py'] else ''))
            elif kind == 'BARA_F':
                satir(rr, 'bara', f'={KP("AB")}', f'=ROUND(3*{NN}*{KP("AA")}*BARA_FIRE,2)', 'SARF', 'bara',
                      'faz barası L1-L2-L3 · şerit adedi, boy ve artikül KURALLAR › PANO Z–AB')
            elif kind == 'BARA_N':
                satir(rr, 'bara', f'={KP("AB")}', f'=ROUND({NN}*{KP("AA")}*BARA_FIRE,2)', 'SARF', 'bara', 'N barası (faz ile aynı)')
            elif kind == 'BARA_PE':
                satir(rr, 'bara', f'={KP("AD")}', f'=ROUND(N({KP("AC")})*{KP("AA")}*BARA_FIRE,2)', 'SARF', 'bara',
                      'PE barası (faz toplam kesitinin yarısı) · KURALLAR › PANO AC–AD')
            elif kind == 'IZOL':
                satir(rr, 'izolator', '=SARF_IZOLATOR', f'=4*(ROUNDUP({KP("AA")}/IZOLATOR_ARALIK_M,0)+2)', 'SARF', 'izolator',
                      'bara izolatörü (L1-L2-L3-N; PE gövdeye)')
            elif kind == 'TARAK':
                mod = '+'.join(f'SUMIFS({BR("AK")},{BR("T")},{k},{BR("B")},"{t}")' for t in TARAK_TIP)
                satir(rr, 'tarak', f'={blk["tarak_key"]}', f'=ROUNDUP(({mod})*MODUL_MM/TARAK_BOY_MM,0)', 'SARF', 'tarak',
                      f'tarak (PIN) · sıra {k:g} · modül × 17,5 mm / parça boyu')
            elif kind == 'KANAL':
                satir(rr, 'kanal', f'=IF({KP("S")}="","",{KP("S")})',
                      f'=IF({KP("S")}="",0,ROUND(((IF({KP("T")}="",N({KP("E")})+1,N({KP("T")})))*{Wic}+2*{KP("N")}/1000)*KANAL_FIRE,1))',
                      'SARF', 'kanal', 'plastik perforeli kanal · artikül ve yatay adet KURALLAR › PANO')
            elif kind == 'KABLO':
                kablo_rows.append(rr)
                satir(rr, 'kablo', ST('SARF_KABLO', k), f'=ROUND(SUMIFS({BR("AF")},{BR("AD")},{k})*KABLO_FIRE,1)', 'SARF', 'kablo', f'ПВ3 {k:g} mm² pano içi')
            elif kind == 'KAPI_KABLO':
                kablo_rows.append(rr)
                satir(rr, 'kablo', ST('SARF_KABLO', 6), f'=N({KP("J")})*KAPI_TOPRAK_M', 'SARF', 'kablo', 'kapak topraklama iletkeni 6 mm²')
            elif kind == 'KLEM_GRI':
                gri_rows.append(rr)
                satir(rr, 'klemens', ST('SARF_KLN', k), f'=IF({KP("K")}="EVET",ROUNDUP(SUMIFS({BR("AI")},{BR("AH")},{k})*KLEMENS_YEDEK,0),0)', 'SARF', 'klemens', f'L/N klemens {k:g} mm²')
            elif kind == 'KLEM_PE':
                satir(rr, 'klemens', ST('SARF_KPE', k), f'=IF({KP("K")}="EVET",ROUNDUP(SUMIFS({BR("AJ")},{BR("AH")},{k})*KLEMENS_YEDEK,0),0)', 'SARF', 'klemens', f'PE klemens {k:g} mm²')
            elif kind == 'ZAGL':
                satir(rr, 'aksesuar', ST('SARF_ZAGL', k), f'=IF(H{gri_rows[-1]}>0,1,0)', 'SARF', 'aksesuar', f'klemens uç kapağı {k:g} mm²')
            elif kind == 'STOPOR':
                satir(rr, 'aksesuar', '=SARF_STOPOR', '=2*(' + '+'.join(f'(H{g}>0)' for g in gri_rows) + ')', 'SARF', 'aksesuar', 'klemens grubu başına 2 stoper')
            elif kind == 'YUKSUK':
                satir(rr, 'yuksuk', ST('SARF_YUKSUK', k), f'=IF(C{rr}="",0,ROUNDUP(SUMIFS({BR("AE")},{BR("AD")},{k},{BR("AG")},"YUKSUK")*2*YUKSUK_FIRE/{paket(rr)},0))',
                      'SARF', 'yuksuk', f'yüksük {k:g} mm² · UÇ TİPİ YÜKSÜK teller × 2 uç')
            elif kind == 'PABUC':
                satir(rr, 'pabuc', ST('SARF_PABUC', k),
                      f'=ROUNDUP((SUMIFS({BR("AE")},{BR("AD")},{k},{BR("AG")},"PABUC")+IF(INDEX(SARF_YUKSUK,MATCH({k},SARF_KESIT,0))&""="",SUMIFS({BR("AE")},{BR("AD")},{k},{BR("AG")},"YUKSUK"),0))*2*PABUC_FIRE,0)',
                      'SARF', 'pabuc', f'kablo pabucu {k:g} mm² · UÇ TİPİ PABUÇ teller (+ НШВИ yoksa kafes) × 2 uç')
            elif kind == 'KAPI_PABUC':
                satir(rr, 'pabuc', ST('SARF_PABUC', 6), f'=2*N({KP("J")})', 'SARF', 'pabuc', 'kapak topraklama pabucu')
            elif kind == 'BAG':
                satir(rr, 'bag', '=SARF_BAG', '=ROUNDUP((' + '+'.join(f'N(H{x})' for x in kablo_rows) + f')*BAG_ADET_M/{paket(rr)},0)', 'SARF', 'bag', 'kablo bağı')
        rr = blk['mm']; isr = blk['isc']
        for c, v in {'A': f'=A{rr-1}', 'B': 'монтажный материал', 'C': 'монтажный материал', 'D': 'монтажный материал', 'E': f'=G{isr}*MONTAJ_MALZ_KAT',
                     'F': f'=E{rr}', 'G': f'=H{rr}*F{rr}', 'H': 1, 'I': 'шт', 'J': 'монтажный материал', 'K': f'=K{rr-1}', 'L': f'=G{rr}', 'M': f'=M{rr-1}'}.items():
            cell(rr, c, v, ACC if c in 'EFGL' else None)
        rr = blk['isc']; lastmat = rr - 2
        saat = p.get('saat')
        E = saat if saat not in (None, '') else f'=ROUND(SUM(AN{s}:AN{lastmat})*ISCILIK_KAT/60,2)'
        for c, v in {'A': f'=A{rr-1}', 'B': 'монтажная\xa0работа', 'C': 'монтажная\xa0работа', 'D': 'монтажная\xa0работа', 'E': E, 'F': '=DK_UCRET',
                     'G': f'=H{rr}*F{rr}', 'H': f'=60*E{rr}', 'I': 'мин', 'J': 'монтажная\xa0работа', 'K': f'=K{rr-1}', 'L': f'=G{rr}', 'M': f'=M{rr-1}'}.items():
            cell(rr, c, v, ACC if c in 'GL' else None)
        rr = blk['top']
        cell(rr, 'A', f'=A{rr-1}')
        for c in 'BCDEFHIJ':
            B[f'{c}{rr}'].border = KEN
        cell(rr, 'G', f'=SUM(G{h+1}:G{rr-1})', '#,##0', TOTF); cell(rr, 'K', f'=K{rr-1}')
        cell(rr, 'L', f'=SUM(L{h+1}:L{rr-1})', '#,##0', TOTF); cell(rr, 'M', f'=M{rr-1}')
        cell(rr, 'P', f'=IF(COUNTIF(P{s}:P{rr-1},"FİYAT YOK")+COUNTIF(P{s}:P{rr-1},"YASAK")+COUNTIF(P{s}:P{rr-1},"MARKA YOK")+COUNTIF(P{s}:P{rr-1},"ARTİKÜL YOK")>0,"EKSİK FİYAT","")', font=F10)
        cell(rr, 'Q', f'=TRIM(IF(KURALLAR!$Q${pr}="OK","","PANO: "&KURALLAR!$Q${pr})&" "&IF(KURALLAR!$W${pr}="OK","","YERLEŞİM: "&KURALLAR!$W${pr}))', font=F10)
        TOPLAMLAR.append(blk)
    B.freeze_panes = 'M2'
    for i, w in enumerate([36.5, 11, 14, 72.5, 10, 10, 15.7, 9.3, 7, 12, 22, 12, 7, 9, 13, 11, 40, 7, 6, 6] + [9] * 24):
        B.column_dimensions[CL(i+1)].width = w
    B.column_dimensions.group('N', 'P', hidden=True, outline_level=1)
    B.column_dimensions.group('R', 'S', hidden=True, outline_level=1)
    B.column_dimensions.group('U', 'AR', hidden=True, outline_level=1)

    # ---------------------------------------------------------------- TEKLIF (v2.7 — TEKLED КП şablonu)
    YESIL = '1E8449'
    CAM = 'Cambria'
    ALTIN_F = PatternFill('solid', fgColor='FFC000')
    f_yesil = lambda sz=10, b=False, i=False: Font(name=CAM, size=sz, bold=b, italic=i, color=YESIL)
    f_siyah = lambda sz=10, b=False, i=False: Font(name=CAM, size=sz, bold=b, italic=i, color='000000')
    sola = Alignment(horizontal='left', vertical='center', wrap_text=True)
    sola_ust = Alignment(horizontal='left', vertical='top', wrap_text=True)
    orta = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sag = Alignment(horizontal='right', vertical='center')

    T.sheet_view.showGridLines = False
    GEN = dict(zip('ABCDEFGHIJK', [5, 20, 40, 10, 9, 15, 17, 15, 40, 14, 14]))
    for c, w in GEN.items():
        T.column_dimensions[c].width = w

    def kutu(rng, val, font, fill=None, align=None, cerceve=False):
        if ':' in rng: T.merge_cells(rng)
        hucre = put(T, rng.split(':')[0], val, font, fill)
        for satir_ in T[rng] if ':' in rng else ((T[rng],),):
            for hh in satir_:
                hh.font = font
                if fill: hh.fill = fill
                if cerceve: hh.border = KEN
        if align: hucre.alignment = align
        return hucre

    def resim(ad, hucre_col, hucre_row, yuk_px, col_off=0, row_off=0):
        """ad: tekled / teksan / kz. Kaynak: LOGO (yalnız tekled) → GORSEL paketi. Yoksa False."""
        try:
            import io, base64
            from openpyxl.drawing.image import Image as XLImage
            from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
            from openpyxl.drawing.xdr import XDRPositiveSize2D
            from openpyxl.utils.units import pixels_to_EMU
            veri = None
            if ad == 'tekled' and os.path.exists(LOGO):
                veri = open(LOGO, 'rb').read()
            else:
                gp = gorsel_paket()
                if gp.get(ad): veri = base64.b64decode(gp[ad])
            if not veri:
                _eksik.add(f'{ad}: görsel verisi yok'); return False
            im = XLImage(io.BytesIO(veri))
            w = int(im.width * yuk_px / im.height); im.width, im.height = w, yuk_px
            im.anchor = OneCellAnchor(_from=AnchorMarker(col=hucre_col, colOff=pixels_to_EMU(col_off), row=hucre_row, rowOff=pixels_to_EMU(row_off)),
                                      ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(yuk_px)))
            T.add_image(im)
            return True
        except Exception as ex:
            _eksik.add(f'{ad}: {ex}')
            return False

    _eksik = set()
    for r_, h_ in ((1, 30), (2, 30), (3, 14)):
        T.row_dimensions[r_].height = h_
    if not resim('tekled', 0, 0, 62, 4, 6):
        kutu('A1:C2', '=FIRMA', Font(name=CAM, size=22, bold=True, color='C8102E'), None, sola)
    resim('teksan', 3, 0, 70, 20, 2)
    kutu('G1:H3', '=IF(FIRMA_TAM="",FIRMA,FIRMA_TAM)&IF(FIRMA_BIN="","",CHAR(10)&"БИН:"&FIRMA_BIN)', f_yesil(9), None, sola_ust)

    T.row_dimensions[5].height = 28
    kutu('A5:H5', 'КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ', f_yesil(16, True), None, orta)

    kutu('A7:B7', 'Исх.№', f_yesil(10)); kutu('C7', '=IF(TEKLIF_NO="","",TEKLIF_NO)', f_yesil(10), None, sola)
    kutu('A8:B8', 'Дата:', f_yesil(10)); x_ = kutu('C8', '=TARIH', f_yesil(10), None, Alignment(horizontal='left', vertical='center'))
    x_.number_format = 'dd.mm.yyyy'
    musteri = [('Кому', '=MUSTERI', 16), ('Адрес', '=IF(MUSTERI_ADRES="","",MUSTERI_ADRES)', 28), ('Объект', '=IF(PROJE="","",PROJE)', 28),
               ('Контактное лицо', '=IF(MUSTERI_KISI="","",MUSTERI_KISI)', 14), ('Рабочий', '=IF(MUSTERI_TEL="","",MUSTERI_TEL)', 14),
               ('Мобильный', '=IF(MUSTERI_MOBIL="","",MUSTERI_MOBIL)', 14), ('Эл. почта', '=IF(MUSTERI_EMAIL="","",MUSTERI_EMAIL)', 14)]
    for i, (lab, f, hgt) in enumerate(musteri):
        r_ = 7 + i
        T.row_dimensions[r_].height = max(T.row_dimensions[r_].height or 15, hgt)
        kutu(f'D{r_}:E{r_}', lab + ' :', f_siyah(10, True), None, Alignment(horizontal='left', vertical='top'))
        kutu(f'F{r_}:H{r_}', f, f_siyah(10), None, sola_ust)

    hr = 15
    T.row_dimensions[hr].height = 40
    for i, hh in enumerate(['№', 'Изображение', 'Наименование товара', 'Кол-во', 'Ед.', 'Цена с НДС', 'Сумма', 'Срок поставки товара']):
        c = put(T, f'{"ABCDEFGH"[i]}{hr}', hh, f_siyah(10, True), ALTIN_F, border=True); c.alignment = orta
    put(T, f'I{hr}', 'iç kontrol (yazdırılmaz)', F10n)
    put(T, f'J{hr}', 'birim без НДС', F10n); put(T, f'K{hr}', 'tutar без НДС', F10n)
    gorsel_satir = P.get('teklif_gorsel', True)
    for i, blk in enumerate(TOPLAMLAR):
        rr = hr + 1 + i; blk['tr'] = rr; h, tr = blk['h'], blk['top']
        T.row_dimensions[rr].height = 42 if gorsel_satir else 20
        put(T, f'J{rr}', f'=IF(KAR_ORANI="","",ROUND(BOM!G{tr}*(1+KAR_ORANI)*(1-ISKONTO),0))', F10n, None, '#,##0')
        put(T, f'K{rr}', f'=IF(J{rr}="","",D{rr}*J{rr})', F10n, None, '#,##0')
        for c, v, fnt, al, fmt in [('A', i + 1, f_yesil(), orta, None), ('B', None, f_siyah(), orta, None),
                                   ('C', f'=BOM!D{h}', f_siyah(), sola, None), ('D', f'=BOM!H{h}', f_yesil(), orta, '#,##0'),
                                   ('E', 'компл.', f_yesil(), orta, None),
                                   ('F', f'=IF(J{rr}="","",ROUND(J{rr}*(1+NDS_ORANI),2))', f_yesil(), orta, '#,##0.00'),
                                   ('G', f'=IF(F{rr}="","",D{rr}*F{rr})', f_yesil(), orta, '#,##0.00'),
                                   ('H', '=IF(TESLIM="","",TESLIM)', f_yesil(), orta, None)]:
            cc = put(T, f'{c}{rr}', v, fnt, None, fmt, border=True); cc.alignment = al
        if gorsel_satir:
            resim('tekled', 1, rr - 1, 26, 14, 15)
        put(T, f'I{rr}', f'=TRIM(BOM!P{tr}&" "&BOM!Q{tr})', F10n)
    l0, l1 = hr + 1, hr + len(TOPLAMLAR)
    TT = l1 + 1
    T.row_dimensions[TT].height = 22
    for c in 'ABCDEFGH':
        put(T, f'{c}{TT}', None, f_siyah(10, True), ALTIN_F, border=True)
    x_ = put(T, f'B{TT}', 'Сумма:', f_siyah(10, True), ALTIN_F, border=True); x_.alignment = orta
    x_ = put(T, f'G{TT}', f'=SUM(G{l0}:G{l1})', f_siyah(10, True), ALTIN_F, '#,##0.00 "KZT"', border=True); x_.alignment = orta
    put(T, f'K{TT}', f'=SUM(K{l0}:K{l1})', F10n, None, '#,##0')
    put(T, f'I{TT}', '=IF(KAR_ORANI="","KÂR ORANI GİRİLMEDİ","")', F10n)
    blk_tt = TT

    notlar = [('GECERLILIK<>""', '"Срок действия ценового предложения "&GECERLILIK&" рабочих дней."'),
              ('TESLIM_YERI<>""', '"Место поставки: "&TESLIM_YERI&"."'),
              ('ODEME<>""', '"Условия оплаты: "&ODEME&"."'),
              ('GARANTI<>""', '"Гарантия: "&GARANTI&"."'),
              ('1=1', '"В щитах с системой сборных шин применяются "&IF(BARA_MALZEME="AL","алюминиевые (Al)","медные (Cu)")&" шины."'),
              ('SALT_MARKA<>""', '"Коммутационная аппаратура (автоматические выключатели, рубильники, контакторы и др.): "&SALT_MARKA&"."'),
              ('SERTIFIKAT<>""', 'SERTIFIKAT'),
              ('PROJE_NOT<>""', 'PROJE_NOT')]
    rr = TT + 2; n0, n1 = rr, rr + len(notlar) - 1
    for k, (kosul, metin) in enumerate(notlar):
        put(T, f'I{rr}', f'=IF({kosul},{metin},"")', F10n)
        put(T, f'J{rr}', f'=IF(I{rr}="","",COUNTIF(I${n0}:I{rr},"?*"))', F10n)
        T.row_dimensions[rr].height = 24
        kutu(f'A{rr}:H{rr}', f'=IFERROR({k + 1}&") "&INDEX(I${n0}:I${n1},MATCH({k + 1},J${n0}:J${n1},0)),"")', f_siyah(10, False, True), None, sola)
        rr += 1
    rr += 1
    T.row_dimensions[rr].height = 20
    kutu(f'A{rr}:H{rr}', '=IF(KAPANIS="","",KAPANIS)', f_yesil(12, True), None, orta); rr += 2

    kutu(f'A{rr}:C{rr}', 'С уважением', f_yesil(10), None, sola); rr += 1
    kutu(f'A{rr}:C{rr}', '=IF(IMZA_UNVAN="","",IMZA_UNVAN)', f_yesil(10), None, sola)
    kutu(f'F{rr}:H{rr}', '=IF(IMZALAYAN="","",IMZALAYAN)', f_yesil(10), None, sola); rr += 3

    alt = [('Адрес', 'FIRMA_ADRES'), ('Телефон', 'FIRMA_TEL'), ('Эл. почта', 'FIRMA_EMAIL'), ('Название банка', 'BANKA'), ('ИИК', 'IIK'), ('БИК', 'BIK')]
    a0 = rr
    for lab, key in alt:
        kutu(f'A{rr}:B{rr}', lab, f_siyah(8), None, Alignment(horizontal='left', vertical='center'))
        kutu(f'C{rr}', f'=":   "&{key}', f_siyah(8), None, Alignment(horizontal='left', vertical='center'))
        rr += 1
    kutu(f'E{a0}:F{a0}', 'Исполнитель:', f_siyah(10), None, Alignment(horizontal='left', vertical='center'))
    kutu(f'G{a0}:H{a0}', '=IF(HAZIRLAYAN="","",HAZIRLAYAN)', f_yesil(10, False, True), None, Alignment(horizontal='left', vertical='center'))
    T.row_dimensions[rr].height = 34
    resim('kz', 0, rr - 1, 40, 2, 2)
    rr += 1

    T.print_area = f'A1:H{rr}'
    T.page_setup.orientation = 'portrait'; T.page_setup.paperSize = 9
    T.page_setup.fitToWidth = 1; T.page_setup.fitToHeight = 0
    T.sheet_properties.pageSetUpPr.fitToPage = True
    T.page_margins.left = 0.4; T.page_margins.right = 0.4; T.page_margins.top = 0.4; T.page_margins.bottom = 0.5
    T.oddFooter.center.text = 'Страница &P'
    T.column_dimensions.group('I', 'K', hidden=True, outline_level=1)

    # ---------------------------------------------------------------- OZET
    OH = ['PANO', 'ADET', 'BİRİM TEKLİF, тг', 'TUTAR, тг', 'KORPUS KODU', 'KORPUS', 'MALZEME (birim), тг', 'İŞÇİLİK (birim), тг',
          'İŞÇİLİK dk (birim)', 'MALİYET (toplam), тг', 'KÂR, тг', 'PANO KONTROL', 'YERLEŞİM']
    for i, hh in enumerate(OH):
        put(OZ, f'{CL(i+1)}1', hh, F10, SARI, border=True)
    for i, blk in enumerate(TOPLAMLAR):
        rr = i + 2; h, tr, kor, isr = blk['h'], blk['top'], blk['kor'], blk['isc']
        for c, v in {'A': f'=BOM!D{h}', 'B': f'=BOM!H{h}', 'C': f'=TEKLIF!J{blk["tr"]}', 'D': f'=TEKLIF!K{blk["tr"]}',
                     'E': f'=BOM!C{kor}' if kor else '', 'F': f'=BOM!D{kor}' if kor else '', 'G': f'=BOM!G{tr}-BOM!G{isr}', 'H': f'=BOM!G{isr}',
                     'I': f'=BOM!H{isr}', 'J': f'=BOM!G{tr}*B{rr}', 'K': f'=IF(D{rr}="","",D{rr}-J{rr})',
                     'L': f'=KURALLAR!Q{blk["pr"]}', 'M': f'=KURALLAR!W{blk["pr"]}'}.items():
            put(OZ, f'{c}{rr}', v, F10n, None, '#,##0' if c in 'CDGHIJK' else None, border=True)
    o1 = len(TOPLAMLAR) + 1; rr = o1 + 1
    put(OZ, f'A{rr}', 'TOPLAM', TOTF, border=True); put(OZ, f'B{rr}', f'=SUM(B2:B{o1})', TOTF, border=True)
    put(OZ, f'D{rr}', f'=IF(COUNT(D2:D{o1})=0,"",SUM(D2:D{o1}))', TOTF, None, '#,##0', border=True)
    put(OZ, f'I{rr}', f'=SUMPRODUCT(B2:B{o1},I2:I{o1})', TOTF, None, '#,##0', border=True)
    put(OZ, f'J{rr}', f'=SUM(J2:J{o1})', TOTF, None, '#,##0', border=True); put(OZ, f'K{rr}', f'=IF(COUNT(K2:K{o1})=0,"",SUM(K2:K{o1}))', TOTF, None, '#,##0', border=True)
    rr += 2
    for lab, f, fmt, key in [('Итого без НДС', f'=TEKLIF!K{TT}', '#,##0', 'OZ_TEKLIF'), ('НДС', f'=TEKLIF!G{TT}-TEKLIF!K{TT}', '#,##0', None),
                             ('Итого с НДС', f'=TEKLIF!G{TT}', '#,##0', None), ('Toplam maliyet', f'=J{o1+1}', '#,##0', 'OZ_MALIYET'),
                             ('Kâr oranı', '=IF(KAR_ORANI="","GİRİLMEDİ",KAR_ORANI)', '0.0%', None), ('Kâr tutarı', f'=K{o1+1}', '#,##0', None),
                             ('İşçilik dk (toplam)', f'=I{o1+1}', '#,##0', None), ('İşçilik saat', f'=ROUND(I{o1+1}/60,1)', '#,##0.0', None),
                             ('İşçilik gün', f'=IF(N(GUN_SAAT)=0,"GUN_SAAT girilmedi",ROUND(I{o1+1}/60/GUN_SAAT,1))', '#,##0.0', None),
                             ('Fiyatsız kalem', '=OK_FIYATSIZ', None, None), ('Otomatik eklenen artikül', '=MAX(0,MAX(BOM!$S:$S))', None, None),
                             ('Pano kontrolü OK olmayan', '=OK_PANO', None, None), ('Yerleşim kontrolü OK olmayan', '=OK_YERLESIM', None, None),
                             ('Mutabakat: ŞEMA ≠ BOM', '=OK_MUTABAKAT', None, None),
                             ('KONTROL sayısı', '=COUNTIF(KT_DURUM,"KONTROL")', None, 'OZ_KONTROL')]:
        put(OZ, f'A{rr}', lab, F10, border=True); put(OZ, f'D{rr}', f, F10, None, fmt, border=True)
        if key: nm(key, f'OZET!$D${rr}')
        rr += 1
    for c, w in zip('ABCDEFGHIJKLM', [40, 7, 15, 15, 14, 40, 15, 15, 12, 16, 14, 26, 30]):
        OZ.column_dimensions[c].width = w

    wb._sheets = [B, FS, K, T, OZ, LS]
    wb.save(out)
    print('OK', out, 'pano', len(PN), 'fiyat', len(FY), 'BOM satır', PLAN[-1]['top'] if PLAN else 0, 'KONTROL', len(KT))


# =====================================================================================
def yenile(src, dst):
    wf = openpyxl.load_workbook(src); wv = openpyxl.load_workbook(src, data_only=True)
    uyari = []

    def dest(n):
        d = wf.defined_names.get(n)
        if d is None: return None
        return next(iter(d.destinations))

    def val(n):
        x = dest(n)
        if not x: return None
        v = wv[x[0]][x[1].replace('$', '')].value
        return v.date().isoformat() if isinstance(v, datetime.datetime) else v

    def col(n):
        x = dest(n)
        if not x: return []
        cells = wv[x[0]][x[1].replace('$', '')]
        return [c[0].value if isinstance(c, tuple) else c.value for c in cells]

    def tbl(names, key=0):
        cols = [col(n) if n else [] for n in names]; L = max(len(c) for c in cols)
        out = []
        for i in range(L):
            rowv = [c[i] if i < len(c) else None for c in cols]
            if rowv[key] in (None, ''): continue
            out.append([v.date().isoformat() if isinstance(v, datetime.datetime) else v for v in rowv])
        return out

    isf = lambda ws, c: isinstance(wf[ws][c].value, str) and str(wf[ws][c].value).startswith('=')
    P = {k: val(k) for k in ('KOD', 'TEKLIF_NO', 'TARIH', 'MUSTERI', 'PROJE', 'SEHIR', 'FIRMA', 'URETICI', 'PANO_EK', 'GECERLILIK', 'TESLIM', 'ODEME',
                             'KAR_ORANI', 'ISKONTO', 'FIRMA_ADRES', 'FIRMA_TEL', 'FIRMA_EMAIL', 'FIRMA_BIN', 'TESLIM_YERI', 'GARANTI',
                             'IMZALAYAN', 'PROJE_NOT', 'FIRMA_TAM', 'MUSTERI_ADRES', 'MUSTERI_KISI', 'MUSTERI_TEL', 'MUSTERI_MOBIL',
                             'MUSTERI_EMAIL', 'BARA_MALZEME', 'SALT_MARKA', 'SERTIFIKAT', 'KAPANIS', 'IMZA_UNVAN', 'HAZIRLAYAN', 'BANKA', 'IIK', 'BIK')}
    P['nds'] = val('NDS_ORANI'); P['kod'] = P.pop('KOD'); P = {k.lower(): v for k, v in P.items()}
    P['not'] = P.pop('proje_not', None)
    if P.get('nds') is None: P['nds'] = 0.16
    kontrol = []
    kur_ = {a: b for a, b in tbl(['KUR_AD', 'KUR_DEG']) if a != 'KZT'}
    tips_all = tbl(['TIP_AD', 'TIP_GRUP', 'TIP_TEL', 'TIP_TELM', 'NORM_DK', 'TIP_SIRA', 'TIP_PABUC', 'TIP_KAT', 'TIP_BOSLUK'])
    korn = [(r[0], r[4], '') for r in tips_all if str(r[0]).startswith('KORPUS_')]
    tips = [[None if v == '' else v for v in r] for r in tips_all if not str(r[0]).startswith('KORPUS_')]
    # v2.11 göç: v2.10 varsayılan normları yeni değerlere; eksik tipler varsayılandan
    goc = []
    yeni_norm = {**{a: n for a, n, _ in KORPUS_NORM}, **{t[0]: t[4] for t in TIPS}}
    for rw in korn + tips:
        pass
    korn = [(a, (yeni_norm[a] if (a in ESKI_NORM and num(n) == ESKI_NORM[a]) else n), x) for a, n, x in korn]
    for rw in tips:
        if rw[0] in ESKI_NORM and num(rw[4]) == ESKI_NORM[rw[0]]:
            rw[4] = yeni_norm[rw[0]]; goc.append(rw[0])
    if any(a in ESKI_NORM for a, _, _ in korn): goc.append('korpus normları')
    if tips and any(rw[5] in (None, '') or rw[1] in (None, '') for rw in tips):
        tips = []; goc.append('tip tablosu eski formatta (SIRA/PABUÇ sütunu yok) — v2.11 varsayılanı')
    var_t = {rw[0] for rw in tips}
    for t in TIPS:
        if t[0] not in var_t and tips: tips.append(list(t)); goc.append(f'tip {t[0]} eklendi')
    if not korn: korn = list(KORPUS_NORM)
    param = {k: val(k) for k, _, _ in PARAM}
    marka_ = tbl(['MARKA_AD', 'MARKA_PARA', 'MARKA_KAT', 'MARKA_IZIN', 'MARKA_NOT'])
    for rw in marka_:
        if rw[0] == 'legrand' and num(rw[2]) == 0.5 and 'teyitsiz' in str(rw[4] or ''):
            rw[2] = 1.0; rw[4] = 'onaylı 260927'; goc.append('legrand 1,00')
    TBL = dict(marka=marka_, tips=tips, korpus_norm=korn, param=param,
               kablo=tbl(['KABLO_IN_ALT', 'KABLO_KESIT']), klemens=tbl(['KLEMENS_ALT', 'KLEMENS_KES', 'KLEMENS_GEN']),
               sarf=tbl(['SARF_KESIT', 'SARF_KABLO', 'SARF_YUKSUK', 'SARF_PABUC', 'SARF_KLN', 'SARF_KPE', 'SARF_ZAGL']),
               sarf_tek={k: val(k) for k in SARF_TEK}, sni=tbl(['SNI_KOD', 'SNI_DELIK']),
               bara_cu=tbl(['CU_KOD', 'CU_MM2', 'CU_A', 'CU_A2']), bara_al=tbl(['AL_KOD', 'AL_MM2', 'AL_A', 'AL_A2']),
               buk=tbl(['BUK_KESIT', 'BUK_CAP', 'BUK_ALAN']), icerik=tbl(['ICERIK_AD', 'ICERIK_IZIN', None]),
               kullanici=tbl(['USER_KOD', 'USER_ACK', 'USER_DEG']), ek_kural=[r[0] for r in tbl(['EK_KURAL'])],
               yas={'YAS_ESKI_AY': val('YAS_ESKI_AY'), 'YAS_YASAK_AY': val('YAS_YASAK_AY')},
               sira_mm=tbl(['SIRA_KAT', 'SIRA_TIP', 'SIRA_INALT', 'SIRA_MM']))
    for k in ('bara_cu', 'bara_al'):
        if not _bara_ok(TBL[k]):
            TBL.pop(k); goc.append(f'{k} (eski ПУЭ tablosu → kullanıcı tablosu 260930)')
    if not TBL['sira_mm']: TBL.pop('sira_mm')
    if not tips: TBL.pop('tips'); TBL.pop('korpus_norm')
    for k in ('kablo', 'klemens', 'sarf', 'sni', 'buk', 'icerik', 'marka'):
        if not TBL.get(k): TBL.pop(k, None)
    if TBL.get('yas') and None in TBL['yas'].values(): TBL.pop('yas')
    if 'sarf' in TBL:
        TBL['sarf'] = [[str(x) if (i > 0 and x not in (None, '')) else (None if x == '' else x) for i, x in enumerate(r)] for r in TBL['sarf']]
    if goc:
        kontrol.append(dict(pano='', artikul='', durum='KONTROL', konu='KAYNAK: v2.11 göçü — ' + ', '.join(goc)))
    kays = [dict(kod=a, dosya=b, tarih=c) for a, b, c in tbl(['KAYNAK_KOD', 'KAYNAK_DOSYA', 'KAYNAK_TARIH'])]

    # FIYAT
    F = wf['FIYAT']; FV = wv['FIYAT']; fy = {}; OZC = ['S', 'T', 'U', 'V', 'W', 'X', 'Y', 'Z', 'AA']; OZK = ['kutup', 'in', 'modul', 'paket', 'kesit', 'g', 'y', 'd', 'kap']
    for r in range(2, F.max_row + 1):
        fid = F[f'A{r}'].value
        if fid in (None, ''): break
        art = FV[f'C{r}'].value
        if art in (None, ''): continue
        art = str(art).strip().rstrip('_'); oto = str(fid).startswith('Y')
        e = FV[f'E{r}'].value
        f = dict(id=None if oto else str(fid), tip=FV[f'B{r}'].value, artikul=art, ad=FV[f'D{r}'].value,
                 liste=e if isinstance(e, (int, float)) else None, birim=FV[f'I{r}'].value, marka=FV[f'J{r}'].value, kaynak=FV[f'N{r}'].value,
                 not_=FV[f'R{r}'].value)
        if str(f['not_'] or '').startswith('YEDEK SARF'): f['not_'] = None
        for c, k in zip(OZC, OZK):
            v = FV[f'{c}{r}'].value
            if v not in (None, ''): f[k] = v
        if not isf('FIYAT', f'AK{r}') and F[f'AK{r}'].value: f['uc'] = str(F[f'AK{r}'].value).upper()
        if oto:
            listede = isinstance(FV[f'AL{r}'].value, (int, float))
            if not listede:
                f['ad'] = f"{art} (listede yok)"; f['marka'] = f.get('marka') or 'pys'; f['kaynak'] = 'ELLE'
            f['not_'] = 'BOM\'a elle eklendi'
            kontrol.append(dict(pano='', artikul=art, konu='ARTIKUL: BOM\'a elle eklenen artikül FIYAT\'a alındı' + ('' if listede else ' — LİSTEDE YOK, fiyat elle'), durum='KONTROL'))
            if isf('FIYAT', f'E{r}') is False and f['liste'] is not None and listede:
                kontrol.append(dict(pano='', artikul=art, konu=f'FIYAT: otomatik satıra elle fiyat {f["liste"]} girildi — liste fiyatıyla karşılaştır', durum='KONTROL'))
        f['not'] = f.pop('not_')
        oz = parse_ozellik(f.get('ad'), f.get('tip'))
        fark = [k for k in ('kutup', 'in', 'modul', 'kap', 'g', 'y') if k in oz and f.get(k) not in (None, '') and num(f[k]) is not None and abs(float(f[k]) - float(oz[k])) > 0.01]
        if fark: uyari.append(f"özellik≠açıklama: {art} {', '.join(f'{k} {f[k]}→{oz[k]}' for k in fark)}")
        if art in fy:
            if fy[art].get('id') is None and f.get('id'): fy[art] = f
            continue
        fy[art] = f
    # BOM
    Bf = wf['BOM']; Bv = wv['BOM']; PN = []; cur = None; ek = P.get('pano_ek')
    for r in range(2, Bf.max_row + 1):
        if Bf[f'I{r}'].value == 'комплект':
            ad = str(Bf[f'D{r}'].value or '')
            if ek and ad.endswith(f' ({ek})'): ad = ad[:-len(f' ({ek})')]
            cur = dict(ad=ad, adet=Bv[f'H{r}'].value or 1, korpus_marka=Bf[f'J{r}'].value, kaynak=Bf[f'K{r}'].value, satirlar=[], pano={})
            PN.append(cur); continue
        if cur is None: continue
        b = str(Bf[f'B{r}'].value or '').replace('\xa0', ' ').strip()
        if not b: continue
        if b.startswith('монтаж'):
            if 'работ' in b and not isf('BOM', f'E{r}') and Bf[f'E{r}'].value not in (None, ''): cur['saat'] = Bf[f'E{r}'].value
            continue
        rol = Bf[f'N{r}'].value
        if rol in (None, '', 'SARF'): continue
        art = Bv[f'C{r}'].value
        if art in (None, '', 'FID YOK'):
            uyari.append(f'BOM satır {r}: artikül okunamadı ({art})'); continue
        art = str(art).strip().rstrip('_')
        if art not in fy:
            rc_ = csv_bul(None, art)
            if rc_ is not None:
                fy[art] = dict(artikul=art, tip=b, csv_marka=str(rc_.MARKA))
                kontrol.append(dict(pano=cur['ad'], artikul=art, konu='ARTIKUL: BOM artikülü FIYAT\'ta yoktu — birlesik.csv\'den alındı', durum='KONTROL'))
            else:
                fy[art] = dict(artikul=art, tip=b, ad=Bv[f'D{r}'].value, marka='pys', kaynak='ELLE', birim='шт')
                kontrol.append(dict(pano=cur['ad'], artikul=art, konu='ARTIKUL: BOM artikülü FIYAT\'ta yoktu — PYS olarak alındı', durum='KONTROL'))
        s = dict(tip=b, artikul=art, adet=Bv[f'H{r}'].value, rol=rol)
        if Bf[f'Q{r}'].value: s['not'] = Bf[f'Q{r}'].value
        o = Bf[f'O{r}'].value
        if o and o != b and not str(o).startswith('KORPUS_'): s['norm'] = o
        t = Bf[f'T{r}'].value
        if t not in (None, ''): s['sira'] = int(t) if isinstance(t, (int, float)) else str(t)
        for c in ('E', 'F'):
            v = Bf[f'{c}{r}'].value
            if not isf('BOM', f'{c}{r}') and isinstance(v, (int, float)):
                f = fy[art]; eski = f.get('liste')
                if c == 'E': f['liste'] = v
                else:
                    kat = FV_k = None
                    for rr in range(2, F.max_row + 1):
                        if str(FV[f'C{rr}'].value) == art: kat = FV[f'L{rr}'].value; FV_k = FV[f'M{rr}'].value; break
                    f['liste'] = v / (kat * FV_k) if isinstance(kat, (int, float)) and isinstance(FV_k, (int, float)) and kat * FV_k else v
                kontrol.append(dict(pano=cur['ad'], artikul=art, konu=f'FIYAT: BOM {c} sütununa elle fiyat {v} — FIYAT liste fiyatı yapıldı (önceki liste {eski})', durum='KONTROL'))
        cur['satirlar'].append(s)
    # PANO / SIRA / KONTROL tabloları
    K = wv['KURALLAR']
    hdr = lambda a, b: next((r for r in range(1, K.max_row + 1) if K[f'A{r}'].value == a and K[f'B{r}'].value == b), None)
    pr = hdr('PANO', 'MONTAJ')
    if pr:
        r = pr + 1
        while K[f'A{r}'].value not in (None, ''):
            i = K[f'X{r}'].value
            if isinstance(i, (int, float)) and 1 <= i <= len(PN):
                bt = str(K[f'Y{r}'].value or '').upper()
                PN[int(i) - 1]['pano'] = dict(montaj=K[f'B{r}'].value, ip=K[f'C{r}'].value, korpus_kaynak=K[f'D{r}'].value, sira=K[f'E{r}'].value,
                                              sira_modul=K[f'F{r}'].value, sira_yuk=K[f'G{r}'].value, bara=None if isf('KURALLAR', f'I{r}') else K[f'I{r}'].value, kapi=K[f'J{r}'].value,
                                              klemens=K[f'K{r}'].value == 'EVET', kanal=K[f'S{r}'].value, kanal_yatay=K[f'T{r}'].value,
                                              kapi_yon=K[f'U{r}'].value, perde=K[f'V{r}'].value, sira_plan=[],
                                              **({'bara_tip': bt} if bt in ('FAZ', 'SNI') else {}))
            r += 1
    sr = hdr('PANO', 'SIRA №')
    if sr:
        r = sr + 1
        while K[f'A{r}'].value not in (None, ''):
            i = K[f'P{r}'].value
            if isinstance(i, (int, float)) and 1 <= i <= len(PN):
                PN[int(i) - 1]['pano'].setdefault('sira_plan', []).append([K[f'C{r}'].value or '', K[f'D{r}'].value])
            r += 1
    kr_ = hdr('PANO', 'ARTİKÜL')
    if kr_:
        for r in range(kr_ + 1, K.max_row + 1):
            if K[f'C{r}'].value in (None, ''): continue
            kn = K[f'C{r}'].value; kn = kn if kodlu(kn) else f'KAYNAK: eski — {kn}'
            kontrol.append(dict(pano=K[f'A{r}'].value, artikul=K[f'B{r}'].value, konu=kn, durum=K[f'D{r}'].value or 'KONTROL'))
    # v2.10: SAYIM MUTABAKATI geri okunur; BOM'da elle değişen adet → kaynak ELLE; grupsuz yeni satır → elle grup
    MUT = []; mr = hdr('PANO', 'GRUP')
    if mr:
        r = mr + 1
        while K[f'A{r}'].value not in (None, ''):
            m = dict(pano=K[f'A{r}'].value, grup=K[f'B{r}'].value, artikul=[x.strip() for x in str(K[f'C{r}'].value or '').split(',') if x.strip()],
                     sema=_int(K[f'D{r}'].value), cetvel=K[f'E{r}'].value, bom=_int(K[f'F{r}'].value), count_a=_int(K[f'G{r}'].value),
                     count_b=_int(K[f'H{r}'].value), lokator=K[f'I{r}'].value, kaynak=K[f'J{r}'].value)
            if K[f'M{r}'].value: m['not'] = K[f'M{r}'].value
            canli = K[f'K{r}'].value
            if isinstance(canli, (int, float)) and _int(canli) != m['bom']:
                kontrol.append(dict(pano=m['pano'], artikul=', '.join(m['artikul']), durum='KONTROL',
                                    konu=f"SAYIM: {m['grup']} BOM elle {m['bom']} → {_int(canli)} — mutabakat ELLE işaretlendi (şema {m['sema']})"))
                m['bom'] = _int(canli); m['kaynak'] = 'ELLE'
            MUT.append({k: v for k, v in m.items() if v not in (None, '')}); r += 1
    for p in PN:
        mg = [m for m in MUT if m['pano'] == p['ad']]
        if not mg:
            p['mutabakat_muaf'] = True; continue
        kapsam = {a for m in mg for a in m['artikul']}
        yeni = {}
        for s in p['satirlar']:
            if s['tip'] == 'korpus' or 'projede yok' in str(s.get('not') or '') or s['artikul'] in kapsam: continue
            yeni[s['artikul']] = yeni.get(s['artikul'], 0) + (_int(s.get('adet')) or 0)
        for a, n in yeni.items():
            MUT.append(dict(pano=p['ad'], grup=f'elle eklendi {a}', artikul=[a], sema=n, bom=n, lokator='kullanıcı — BOM\'a elle eklendi', kaynak='ELLE'))
            kontrol.append(dict(pano=p['ad'], artikul=a, durum='KONTROL', konu=f'SAYIM: BOM\'a elle eklenen {a} ({n} ad.) şemada yok — kullanıcı kararı, mutabakata ELLE grup açıldı'))
    seen = set(); KT = []
    for x in kontrol:
        key = (x.get('pano'), str(x.get('artikul')), x.get('konu'))
        if key in seen: continue
        seen.add(key); KT.append(x)
    js = dict(proje=P, kur=kur_, kaynaklar=kays, fiyat=[{k: v for k, v in f.items() if v not in (None, '')} for f in fy.values()],
              panolar=PN, mutabakat=MUT, kontrol=KT, tablolar=TBL)
    json.dump(js, open(dst, 'w'), ensure_ascii=False, indent=1, default=str)
    print('OK json', dst, 'pano', len(PN), 'artikül', len(fy), 'kontrol', len(KT))
    for u in uyari[:30]: print(' !', u)


# =====================================================================================
def donustur(src, dst):
    wb = openpyxl.load_workbook(src); ws = wb.worksheets[0]; wv = openpyxl.load_workbook(src, data_only=True).worksheets[0]
    FY = {}; PN = []; cur = None; pys = 0
    for r in range(2, ws.max_row + 1):
        I = str(ws[f'I{r}'].value or ''); b = ws[f'B{r}'].value
        if 'комплект' in I:
            cur = dict(ad=ws[f'D{r}'].value, adet=ws[f'H{r}'].value or 1, korpus_marka=ws[f'J{r}'].value, kaynak=ws[f'K{r}'].value,
                       satirlar=[], saat=None, sarf=False, pano={}, mutabakat_muaf=True)
            PN.append(cur); continue
        if cur is None or not b: continue
        b = str(b).replace('\xa0', ' ').strip(); bl = b.lower()
        if 'монтаж' in bl and 'работ' in bl: cur['saat'] = wv[f'E{r}'].value; continue
        if 'монтаж' in bl: continue
        art = ws[f'C{r}'].value; ad = str(ws[f'D{r}'].value or '').strip(); mk = str(ws[f'J{r}'].value or '').strip()
        art = str(art).strip().rstrip('_') if art not in (None, '') else None
        if not art: pys += 1; art = f'PYS-{pys:03d}'
        if art in FY and FY[art]['ad'] != ad: pys += 1; art = f'{art}-{pys:03d}'
        FY.setdefault(art, dict(artikul=art, ad=ad, tip=b, marka=mk, birim=str(ws[f'I{r}'].value or 'шт'),
                                liste=wv[f'E{r}'].value if isinstance(wv[f'E{r}'].value, (int, float)) else None, kaynak='ELLE'))
        h = wv[f'H{r}'].value
        cur['satirlar'].append(dict(tip=b, artikul=art, adet=h if isinstance(h, (int, float)) else 0, rol=''))
    js = dict(proje=dict(kod=PN[0]['kaynak'] if PN else '', tarih=datetime.date.today().isoformat()), fiyat=list(FY.values()), panolar=PN, kontrol=[])
    json.dump(js, open(dst, 'w'), ensure_ascii=False, indent=1, default=str)
    print('OK json', dst, 'pano', len(PN), 'artikül', len(FY))


if __name__ == '__main__':
    if sys.argv[1] == '--donustur':
        donustur(sys.argv[2], sys.argv[3])
    elif sys.argv[1] == '--yenile':
        yenile(sys.argv[2], sys.argv[3])
    else:
        kur(json.load(open(sys.argv[1])), sys.argv[2])
