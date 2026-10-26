#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
bom_oku.py — YSF HICH / TEKLED BOM Excel (şablon "fiyat burda", kolon A–O) → uretim onay JSON iskeleti.

Deterministik: AI yok, tahmin yok. Okunamayan / sınıflanamayan her satır "kontrol" listesine düşer.

Kullanım:
    python3 bom_oku.py BOM.xlsx proje.json [--sayfa "fiyat burda"] [--proje-no 261006] [--sehir Атырау] [--musteri GULDANA]

Kolon şeması (bom-workflow):
  A=ROW  B=Наименование1  C=malzeme tipi  D=Артикул  E=Номенклатура  F=Цена  G=net  H=tutar
  I=Кол.  J=Ед.изм  K=завод(marka)  L=malz tutar  M=mal grup  N=srk euro  O=oguzhan
  Pano satırı: J içinde "компл" → E=pano adı, I=pano adedi.
  Fiyat kolonları (F,G,H,L,N,O) ÇIKTIYA ALINMAZ.

BOM robotu (bom_kur.py v2.x) düzeni otomatik tanınır (1. satırda "malzeme" başlığı); kolonlar başlık adından
bulunur (malzeme · Артикул · Номенклатура · Кол-во · Ед.изм · завод), fiyat kolonları silinmiş kopyalar da okunur.
  Standart: A=pano adı  B=malzeme tipi  C=Артикул  D=Номенклатура  E–G=fiyat  H=Кол-во  I=Ед.изм  J=marka
  Pano satırı: I içinde "компл" → D=pano adı, H=pano adedi.
--sayfa verilmezse sırayla "fiyat burda", "BOM", ilk sayfa okunur.
"""
import json, re, sys, argparse
from datetime import date
import openpyxl

# C kolonu kısaltması → iç kategori
C_MAP = {
    'korpus': 'korpus', 'корпус': 'korpus',
    'tmş': 'mccb', 'tms': 'mccb', 'tmş.': 'mccb',
    'oto': 'mcb', 'otomat': 'mcb',
    'dif': 'rccb', 'dİf': 'rccb', 'dıf': 'rccb', 'kaçak': 'rccb',
    'rcbo': 'rcbo', 'avdt': 'rcbo',
    'impuls': 'role', 'İmpuls': 'role', 'i̇mpuls': 'role',
    'röle': 'role', 'role': 'role',
    'röle altlığı': 'aksesuar', 'role altligi': 'aksesuar',
    'sinyal lambası': 'lamba', 'sinyal lambasi': 'lamba', 'lamb': 'lamba', 'lamba': 'lamba',
    'buton': 'buton', 'şalter': 'secici', 'salter': 'secici', 'pako': 'secici',
    'klemens': 'sarf', 'loobar': 'sarf', 'bara': 'sarf', 'din': 'sarf', 'kanal': 'sarf', 'kablo': 'sarf',
    'pabuç': 'sarf', 'yüksük': 'sarf',
    'rubilnik': 'rubilnik', 'tt': 'tt', 'sayac': 'sayac', 'sayaç': 'sayac',
    'opn': 'opn', 'parafudr': 'opn', 'sigorta': 'sigorta', 'buşon': 'sigorta',
    'kontaktör': 'kontaktor', 'kontaktor': 'kontaktor', 'termik': 'termik',
    'vfd': 'vfd', 'softstarter': 'vfd',
    'montajный материал': 'montaj', 'montaj': 'montaj', 'монтажный материал': 'montaj',
    'montajная работа': 'iscilik', 'işçilik': 'iscilik', 'iscilik': 'iscilik', 'монтажная работа': 'iscilik',
}

# Номенклатура (E) üzerinden yedek sınıflama — sıra önemlidir (RCBO, MCB'den önce)
E_REGEX = [
    ('korpus',    r'корпус|щит\s|щмп|щрн|вру\s*\d|бокс|nexus|titan|шкаф|kutu|pano gövde'),
    ('rcbo',      r'авдт|rcbo|nxble|ds20\d|дифф?\.?\s*автомат|диф\.?\s*авт|ad\d{2}'),
    ('rccb',      r'узо|rccb|nl1-|nxl-|f20\d|вд\d{1,3}|дифф?\.?\s*выключ'),
    ('mccb',      r'nxm|nm1|nm8|nsx|xt\d|tmax|ва\s?88|в литом|mccb|компакт'),
    ('rubilnik',  r'рубильник|выключатель.нагрузки|разъединитель|nh40|nh1|ot\d{2,}|вр-?32|load break'),
    ('mcb',       r'nxb|nb1|dz47|s20\d|ва\s?47|ic60|авт\.?\s*выкл|автоматический выключатель|mcb'),
    ('kontaktor', r'контактор|nxc|nc1|кми|af\d{2}|lc1|kontakt'),
    ('termik',    r'теплов\w* реле|nxr|nr2|рти|lrd|ta25|термик'),
    ('opn',       r'узип|опн|ограничит\w* перенапр|nu6|surge|ovr|parafudr'),
    ('sigorta',   r'предохранит|плавк|fuse|ппн|rt\d{2}'),
    ('sayac',     r'счетчик|счётчик|меркурий|сар4|отан|saiman|энергомера|ce\d{3}|meter'),
    ('tt',        r'трансформатор тока|т-0[,.]66|тшп|топ-0|current transf'),
    ('lamba',     r'индикатор|лампа|nd16|ad22|сигнальн'),
    ('buton',     r'кнопк|np8.*b[an]|np2|xb\d|грибок|аварийн\w* стоп|button'),
    ('secici',    r'переключат|np8.*x|кулачк|lw\d{2}|pako'),
    ('vfd',       r'преобразователь частоты|nvf|частотн|устройство плавного|vfd|ats\d'),
    ('role',      r'реле|таймер|nkg|nhp|фотореле|rm4|nxj|импульсн'),
    ('sarf',      r'шина|изолятор|клемм|din|дин-рейк|кабел\w*.канал|кабель|провод|пв[13]|наконечник|гильз|шнк|шни|стяжк|маркир'),
]
# BOM robotu (bom_kur.py TIPS) tip kodu → iç kategori; yeni düzende C_MAP'ten önce bakılır
# (dif = RCBO NXBLE, uzo = УЗО; ats / tanımsız tip Номенклатура regex'ine düşer)
TIP_MAP = {
    'korpus': 'korpus', 'mccb': 'mccb', 'oto': 'mcb', 'dif': 'rcbo', 'uzo': 'rccb',
    'bh': 'rubilnik', 'rubilnik': 'rubilnik', 'ppn': 'sigorta', 'contactor': 'kontaktor',
    'sayac': 'sayac', 'tt': 'tt', 'lamb': 'lamba', 'indicator': 'lamba', 'button': 'buton',
    'perekluchatel': 'secici', 'rele': 'role', 'timer': 'role', 'avr_rele': 'role', 'rn': 'role',
    'fotorele': 'role',
    'motor': 'aksesuar', 'shunt': 'aksesuar', 'kilit': 'aksesuar', 'kontak': 'aksesuar', 'krm': 'aksesuar',
    'aksesuar': 'aksesuar', 'rozetka': 'aksesuar', 'stable': 'aksesuar', 'condencator': 'aksesuar',
    'bara': 'sarf', 'nbara': 'sarf', 'izolator': 'sarf', 'kanal': 'sarf', 'kablo': 'sarf', 'klemens': 'sarf',
    'yuksuk': 'sarf', 'pabuc': 'sarf', 'bag': 'sarf', 'tarak': 'sarf',
}
# eski şablon kolonları: tip, artikül, ad, adet, birim, marka, ek ad (B); yeni düzen başlıktan bulunur (yeni_duzen)
DUZEN_ESKI = dict(tip='C', art='D', ad='E', adet='I', birim='J', marka='K', ad2='B')

DEVICE_CATS = {'mccb', 'mcb', 'rccb', 'rcbo', 'rubilnik', 'kontaktor', 'termik', 'opn', 'sigorta',
               'sayac', 'tt', 'lamba', 'buton', 'secici', 'vfd', 'role'}
DOOR_CATS = {'lamba', 'buton', 'secici'}
POZ_PREFIX = {'mccb': 'QF', 'mcb': 'QF', 'rcbo': 'QFD', 'rccb': 'QD', 'rubilnik': 'QS', 'kontaktor': 'KM',
              'termik': 'KK', 'opn': 'FV', 'sigorta': 'FU', 'sayac': 'PI', 'tt': 'TA', 'lamba': 'HL',
              'buton': 'SB', 'secici': 'SA', 'vfd': 'UZ', 'role': 'KL'}


# TEKLED sabit seçimleri (Yusuf onayı 08.10.2026) — docs/06_TEKLED_STANDARTLARI.md ile aynı tutulmalı
TEKLED_STANDART = {
    'bara_renk': 'ПУЭ',                 # L1 жёлтый / L2 зелёный / L3 красный (ПУЭ 1.1.30)
    'ic_faz_renk': 'IEC',               # iç kablo L1 коричневый / L2 чёрный / L3 серый
    'kumanda_renk': '~230 В: L чёрный, N синий; =24 В: + красный, − белый; ~24 В: L коричневый, N зелёный',
    'kesit_tablosu_onayli': True,
    'tork_tablosu_onayli': True,
    'sarf_goster': False,
}


def norm(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def kategori(c_val, e_val):
    c = norm(c_val).lower()
    if c in C_MAP:
        return C_MAP[c], 'C'
    for k, v in C_MAP.items():
        if c and c.startswith(k):
            return v, 'C'
    e = norm(e_val).lower()
    for cat, rx in E_REGEX:
        if re.search(rx, e):
            return cat, 'E'
    return None, None


def parametre_cek(text):
    """Номенклатура metninden In / kutup / eğri / Icu / IΔn çıkarır. Bulunamayan alan None kalır."""
    t = norm(text)
    p = {'In': None, 'kutup': None, 'egri': None, 'Icu_kA': None, 'rcd_mA': None, 'rcd_tip': None}
    m = re.search(r'(\d)\s*[PР]\s*\+\s*N', t, re.I)
    if m:
        p['kutup'] = f'{m.group(1)}P+N'
    else:
        m = re.search(r'\b(\d)\s*[PР]\b', t, re.I)
        if m:
            p['kutup'] = f'{m.group(1)}P'
    m = re.search(r'\b([BCD])\s?(\d{1,3})\b', t)          # C16, C 25
    if m:
        p['egri'] = m.group(1)
        p['In'] = float(m.group(2))
    m2 = re.search(r'х-ка\s*([BCD])', t, re.I)
    if m2:
        p['egri'] = m2.group(1)
    if p['In'] is None:
        m = re.search(r'(?<![\d/])(\d{1,4}(?:[.,]\d)?)\s*[AА](?![a-zа-я])', t)
        if m:
            p['In'] = float(m.group(1).replace(',', '.'))
    m = re.search(r'(\d{1,3}(?:[.,]\d)?)\s*к\s*[AА]', t, re.I) or re.search(r'(\d{1,3}(?:[.,]\d)?)\s*kA', t, re.I)
    if m:
        p['Icu_kA'] = float(m.group(1).replace(',', '.'))
    m = re.search(r'(\d{2,3})\s*m[AА]', t, re.I) or re.search(r'(\d{2,3})\s*мА', t, re.I)
    if m:
        p['rcd_mA'] = int(m.group(1))
    m = re.search(r'(?:type|тип[а]?)\s*(AC|A|B|F)\b', t, re.I)
    if m:
        p['rcd_tip'] = m.group(1).upper()
    for k in ('In',):
        if p[k] is not None and p[k] == int(p[k]):
            p[k] = int(p[k])
    return p


def yeni_duzen(ws):
    """BOM robotu (bom_kur.py) düzeni: 1. satırda "malzeme" başlığı. Kolonlar başlık adından bulunur
    (fiyat kolonları silinmiş "onay dosyası için" kopyaları da okunur). Değilse None."""
    bas = {norm(c.value).lower(): c.column_letter for c in ws[1] if norm(c.value)}
    if 'malzeme' not in bas:
        return None
    bul = lambda *on: next((h for b, h in bas.items() if b.startswith(on)), None)
    dz = dict(tip=bas['malzeme'], art=bul('артикул'), ad=bul('номенклатура'), adet=bul('кол'),
              birim=bul('ед'), marka=bul('завод'), ad2=None)
    return dz if all(dz[k] for k in ('art', 'ad', 'adet', 'birim')) else None


def oku(xlsx, sayfa=None):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if sayfa:
        ws = wb[sayfa]
    else:
        ad = next((s for s in ('fiyat burda', 'BOM') if s in wb.sheetnames), None)
        ws = wb[ad] if ad else wb.worksheets[0]
    dz = yeni_duzen(ws)
    yeni = dz is not None
    dz = dz or DUZEN_ESKI

    def kat(tip, ad):
        t = norm(tip).lower()
        if yeni and t in TIP_MAP:
            return TIP_MAP[t], 'C'
        return kategori(tip, ad)

    panolar, hatalar = [], []
    cur = None
    for r in range(1, ws.max_row + 1):
        g = lambda k: ws[f'{dz[k]}{r}'].value if dz[k] else None
        B, C, D, E, I, J, K = g('ad2'), g('tip'), g('art'), g('ad'), g('adet'), g('birim'), g('marka')
        if all(norm(x) == '' for x in (B, C, D, E, I, J)):
            continue
        txt = (norm(B) + ' ' + norm(E)).lower()
        if re.search(r'итого|toplam|total|всего', txt) and not norm(D):
            continue
        if r <= 2 and re.search(r'наименование|номенклатура|артикул|row', txt + norm(D).lower()):
            continue                                   # başlık satırı
        if 'компл' in norm(J).lower() and not norm(D) and kat(C, '')[0] is None \
                and not re.search(r'монтажн|montaj|işçilik|работа', txt):
            ad = norm(E) or norm(B)
            try:
                adet = int(float(I)) if I not in (None, '') else 1
            except Exception:
                adet = 1
            cur = {'kod': f'P-{len(panolar) + 1:02d}', 'ad': ad, 'aciklama': '', 'adet': adet,
                   'bom_satir': r, 'malzemeler': [], 'kontrol': []}
            panolar.append(cur)
            continue
        if cur is None:
            hatalar.append(f'Satır {r}: pano başlığından önce malzeme satırı — "{norm(E)}"')
            continue
        cat, kaynak = kat(C, E)
        try:
            q = float(I) if I not in (None, '') else None
        except Exception:
            q = None
        mal = {'satir': r, 'tip_C': norm(C), 'artikul': norm(D), 'ad': norm(E), 'adet': q,
               'birim': norm(J) or 'шт', 'marka': norm(K), 'kategori': cat or 'belirsiz', 'kategori_kaynak': kaynak}
        if cat is None:
            cur['kontrol'].append({'seviye': 'KONTROL', 'mesaj':
                f'BOM стр. {r}: категория не определена — «{norm(E)}» (C="{norm(C)}"). Отнести вручную.'})
        if q is None:
            cur['kontrol'].append({'seviye': 'STOP', 'mesaj': f'BOM стр. {r}: количество не читается — «{norm(E)}».'})
        cur['malzemeler'].append(mal)
    return panolar, hatalar


def iskelet(panolar):
    """Cihaz malzemelerini devre satırlarına açar (adet kadar). Yük/kablo/poz boş — çizimden doldurulacak."""
    for p in panolar:
        devreler, kapi = [], []
        sayac = {}
        korpus = next((m for m in p['malzemeler'] if m['kategori'] == 'korpus'), None)
        p['korpus'] = {'ad': korpus['ad'] if korpus else '', 'artikul': korpus['artikul'] if korpus else '',
                       'marka': korpus['marka'] if korpus else '', 'H': None, 'W': None, 'D': None,
                       'IP': None, 'montaj': None, 'RAL': None, 'din_sira': None, 'sira_modul': None,
                       'kapi': 'Одностворчатая; петли справа, открывание вправо', 'kilit': 'Замок на левой кромке двери',
                       'kablo_giris': None}
        if korpus:
            m = re.search(r'(\d{3,4})\s*[xх×]\s*(\d{3,4})\s*[xх×]\s*(\d{2,4})', korpus['ad'])
            if m:
                p['korpus'].update(H=int(m.group(1)), W=int(m.group(2)), D=int(m.group(3)))
            m = re.search(r'IP\s?(\d{2})', korpus['ad'], re.I)
            if m:
                p['korpus']['IP'] = f'IP{m.group(1)}'
        else:
            p['kontrol'].append({'seviye': 'STOP', 'mesaj': 'В BOM нет строки корпуса.'})
        for m in p['malzemeler']:
            if m['kategori'] not in DEVICE_CATS or not m['adet']:
                continue
            n = int(round(m['adet']))
            if abs(n - m['adet']) > 1e-6:
                p['kontrol'].append({'seviye': 'STOP', 'mesaj': f'BOM стр. {m["satir"]}: дробное кол-во аппарата ({m["adet"]}).'})
            par = parametre_cek(m['ad'])
            for _ in range(n):
                pre = POZ_PREFIX[m['kategori']]
                sayac[pre] = sayac.get(pre, 0) + 1
                row = {'poz': f'{pre}{sayac[pre]}', 'kategori': m['kategori'], 'cihaz': m['ad'],
                       'artikul': m['artikul'], 'marka': m['marka'], **par, 'modul': None,
                       'grup': 'cikis', 'yuk': '', 'kablo': '', 'P_kW': None, 'not': ''}
                if m['kategori'] in DOOR_CATS:
                    row.update(renk=None, yazi='')
                    kapi.append(row)
                else:
                    devreler.append(row)
        p['devreler'] = devreler
        p['kapi_elemanlari'] = kapi
        p['teknik'] = {k: None for k in ('Ue', 'Ui', 'f', 'sistem', 'fazlar', 'In_pano', 'Ik_proje_kA', 'Icw_kA',
                                         'kumanda', 'form', 'IK', 'klimatik', 'yer', 'ortam_C', 'giris', 'cikis',
                                         'bara', 'standart')}
        p['teknik'].update(f='50 Гц', standart='ГОСТ IEC 61439-1-2024, ГОСТ IEC 61439-2-2015')  # KONTROL: proje std
        p['notlar'] = []
    return panolar


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx'); ap.add_argument('json_out')
    ap.add_argument('--sayfa'); ap.add_argument('--proje-no', default=''); ap.add_argument('--sehir', default='')
    ap.add_argument('--musteri', default=''); ap.add_argument('--obje', default='')
    a = ap.parse_args()
    panolar, hatalar = oku(a.xlsx, a.sayfa)
    panolar = iskelet(panolar)
    data = {'proje': {'no': a.proje_no, 'sehir': a.sehir, 'musteri': a.musteri, 'obje': a.obje,
                      'rev': 'R0', 'tarih': date.today().strftime('%d.%m.%Y'), 'firma': 'TEKLED',
                      'kaynak_bom': a.xlsx.split('/')[-1], 'kaynak_sema': '', 'hazirlayan': '', 'kontrol_eden': '',
                      'rev_gecmisi': [{'rev': 'R0', 'tarih': date.today().strftime('%d.%m.%Y'), 'aciklama': 'Первый выпуск'}]},
            'genel': dict(TEKLED_STANDART),
            'panolar': panolar, 'okuma_hatalari': hatalar}
    with open(a.json_out, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    n_dev = sum(len(p['devreler']) + len(p['kapi_elemanlari']) for p in panolar)
    print(f'Pano: {len(panolar)} | cihaz satırı: {n_dev} | okuma hatası: {len(hatalar)}')
    for p in panolar:
        print(f"  {p['kod']} {p['ad']} ×{p['adet']}: {len(p['malzemeler'])} malzeme, "
              f"{len(p['devreler'])} devre, {len(p['kapi_elemanlari'])} kapı, {len(p['kontrol'])} kontrol")
    for h in hatalar:
        print('  HATA:', h)


if __name__ == '__main__':
    main()
